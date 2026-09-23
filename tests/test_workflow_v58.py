"""Customer workflow regressions. All numeric fixtures are synthetic test data."""
import io,json,threading,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from datetime import datetime,timedelta
from concurrent.futures import ThreadPoolExecutor
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e,bulk_refresh as b,v42_collectors as c,v42_main as main
from app import document_imports as imports


class SharedWorkflow(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.patches=[patch.object(e,'RUNTIME_DIR',self.root),patch.object(e,'ROUTE_CONFIG',{})]
        for p in self.patches:p.start()
        self.manager=b.BulkManager(self.root/'bulk')
    def tearDown(self):
        for t in (self.manager.worker,self.manager.export_worker):
            if t:t.join(120)
        for p in reversed(self.patches):p.stop()
        self.temp.cleanup()
    def save(self,company,o='Москва',d='Санкт-Петербург',value=734):
        aid=e.begin_live_attempt(company,o,d)
        stamp=(datetime.now().astimezone()-timedelta(days=2)).isoformat()
        e.save_live_update(company,o,d,{'w001':{'kind':'exact','price':value},'w100':{'kind':'exact','price':2200,'rate_per_kg':22}},
                          {'captured_at':stamp,'source_url':'https://example.invalid/v58-test-only'},aid)
        e.finish_live_attempt(company,o,d,aid,rows=2)
        return {'company':company,'ok':True,'rows':2,'message':'TEST FIXTURE'}
    def paused_plan(self,scope='reference',origins=None,destinations=None):
        with patch.object(self.manager,'_launch'):
            job=self.manager.start(scope,origins,destinations)
        with self.manager.db() as db:db.execute("UPDATE jobs SET status='paused' WHERE id=?",(job['job_id'],))
        return job['job_id']
    def test_full_254_export_uses_previous_runs_and_confirmed_files_even_before_any_new_check(self):
        old=self.paused_plan('origins',['Москва'],['Санкт-Петербург'])
        for n,company in enumerate(e.COMPANIES):
            result=self.save(company,value=700+n)
            self.manager._save_result(old,0,'Москва','Санкт-Петербург',result)
        # A later one-route result is newer than the old job's snapshot.
        self.save('Werner',value=890)
        csv='Компания;Откуда;Куда;Вес, кг;Цена, руб\nДЛ;Москва;Санкт-Петербург;1;987\n'.encode()
        preview=imports.preview(csv,'price.csv','ДЛ','Москва','Санкт-Петербург');imports.commit(preview['token'])
        ident=self.paused_plan();self.manager.prepare_export(ident);self.manager.export_worker.join(120)
        status=self.manager.status(ident)
        self.assertEqual(status['export_status'],'ready',status)
        self.assertEqual(status['total_routes'],254);self.assertEqual(status['completed_checks'],0)
        from app.excel_export import SHEETS
        workbook=load_workbook(self.manager.download(ident),data_only=True)
        expected=b.plan_routes()
        for n,company in enumerate(e.COMPANIES):
            ws=workbook[SHEETS.get(company,company)]
            rows=[r for r in ws.iter_rows(min_row=2,values_only=True) if r[1] and r[2]]
            self.assertEqual([(r[1],r[2]) for r in rows],expected,company)
            route=next(r for r in rows if r[1:3]==('Москва','Санкт-Петербург'))
            self.assertEqual(route[3],987 if company=='ДЛ' else 890 if company=='Werner' else 700+n)
            self.assertEqual(route[13],22)
            self.assertIsNone(route[4])  # no invented missing tariff
        self.assertEqual(workbook['Направления']['C255'].value,expected[-1][0]);workbook.close()
        # A new manager and cleared memory still see all stored prices.
        restored=b.BulkManager(self.root/'bulk');e._cached_json.cache_clear()
        self.assertEqual(e.quote('Werner','Москва','Санкт-Петербург','w001')['price'],890)
        self.assertEqual(e.quote('ДЛ','Москва','Санкт-Петербург','w001')['price'],987)
        self.assertFalse(e.quote('Werner','Москва','Санкт-Петербург','w001')['online'])
        self.save('Werner',value=901)
        self.assertTrue(restored.status(ident)['export_outdated'])
        with self.assertRaisesRegex(ValueError,'Данные изменились'):restored.download(ident)
    def test_pause_mid_route_persists_each_company_and_resumes_only_remaining(self):
        calls=[];pause_once=[True]
        def collect(companies,o,d,profile,on_progress,*,full_grid,should_stop):
            calls.append(companies.copy());out=[]
            for company in companies:
                if should_stop():break
                result=self.save(company,o,d);on_progress(result);out.append(result)
                if pause_once[0]:pause_once[0]=False;self.manager.pause(self.manager.active())
            return out
        self.manager.collector=collect
        with patch('app.excel_export.export_bytes',return_value=b'test'):
            job=self.manager.start('origins',['Казань'],['Уфа']);self.manager.worker.join(10)
            self.assertEqual(self.manager.status()['status'],'paused');self.assertEqual(self.manager.status()['completed_checks'],1)
            self.assertEqual(e.quote(e.COMPANIES[0],'Казань','Уфа','w001')['price'],734)
            self.manager.resume(job['job_id']);self.manager.worker.join(15);self.manager.export_worker.join(15)
            self.assertEqual(self.manager.status()['completed_checks'],17)
            self.assertEqual(calls[1],e.COMPANIES[1:])
    def test_saved_report_does_not_claim_new_online_check(self):
        self.save('ДЛ')
        aid=e.begin_live_attempt('ДЛ','Москва','Санкт-Петербург')
        e.save_live_update('ДЛ','Москва','Санкт-Петербург',{'w001':{'kind':'exact','price':777}},{'source_url':'https://example.invalid/fixture'},aid)
        e.finish_live_attempt('ДЛ','Москва','Санкт-Петербург',aid,rows=1)
        self.assertTrue(e.quote('ДЛ','Москва','Санкт-Петербург','w001')['online'])
        payload=b.capture_company('ДЛ','Москва','Санкт-Петербург',{'ok':True,'snapshot':True})
        self.assertTrue(all(not i['collected_online'] for i in payload['items']))

    def test_stop_route_api_and_history_preserve_data(self):
        self.save('ДЛ');ident=self.paused_plan()
        job={'status':'running','job_id':'fixture','origin':'Москва','destination':'Санкт-Петербург'}
        with patch.object(main,'BULK',self.manager),patch.dict(main.COLLECT_JOBS,{'Москва|Санкт-Петербург':job},clear=True):
            client=TestClient(main.app)
            response=client.post('/api/collect/stop',json={'origin':'Москва','destination':'Санкт-Петербург'})
            self.assertEqual(response.status_code,200);self.assertTrue(response.json()['stop_requested'])
            self.assertEqual(client.get('/api/bulk/history').json()['jobs'][0]['job_id'],ident)
            response=client.get('/api/compare',params={'origin':'Москва','destination':'Санкт-Петербург','profile':'w001','companies':'ДЛ'})
            self.assertEqual(response.json()['items'][0]['price'],734)


class Scheduling(unittest.TestCase):
    def test_stop_does_not_start_queued_carriers(self):
        from app import online_tariffs
        entered=threading.Event();release=threading.Event();stop=threading.Event();started=[];lock=threading.Lock()
        companies=['test-'+str(i) for i in range(17)]
        def collect(company,*args):
            with lock:
                started.append(company)
                if len(started)==7:entered.set()
            release.wait(5)
            return {'w100':{'kind':'exact','price':1200}},{}
        with patch.dict(online_tariffs.ADAPTERS,{c:None for c in companies}),patch.object(online_tariffs,'collect',side_effect=collect),patch.object(c,'begin_live_attempt',return_value='test'),patch.object(c,'save_live_update'),patch.object(c,'finish_live_attempt'),patch.object(c,'_log'):
            with ThreadPoolExecutor(1) as pool:
                future=pool.submit(c.collect_selected,companies,'Казань','Уфа',should_stop=stop.is_set)
                try:self.assertTrue(entered.wait(5));stop.set()
                finally:release.set()
                result=future.result(10)
        self.assertEqual(len(started),7);self.assertEqual(len(result),7)

if __name__=='__main__':unittest.main()
