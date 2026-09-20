import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app import v42_engine as engine, v42_main as main, document_imports as imports
from app import tariff_documents as docs, official_documents as official

BASE=Path(__file__).resolve().parent.parent
FIX=Path(__file__).parent/'fixtures'
ROUTE=('Казань','Екатеринбург')


def table(company='Werner',origin=ROUTE[0],destination=ROUTE[1],rows=((100,1234),(200,2345))):
    wb=Workbook();ws=wb.active;ws.append(docs.GENERIC_HEADER)
    for weight,price in rows:ws.append([company,origin,destination,weight,price])
    out=io.BytesIO();wb.save(out);return out.getvalue()


class ImportLifecycle(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.ctx=patch.object(engine,'RUNTIME_DIR',Path(self.tmp.name));self.ctx.start()
        self.client=TestClient(main.app)
    def tearDown(self):self.ctx.stop();self.tmp.cleanup()
    def preview(self,raw=None,name='tariff.xlsx',company='Werner',route=ROUTE):
        return self.client.post('/api/import/preview',data={'company':company,'origin':route[0],'destination':route[1]},files={'file':(name,raw if raw is not None else table())})
    def commit(self,p):return self.client.post('/api/import/commit',json={'token':p['token']})
    def quote(self,company='Werner',route=ROUTE,profile='w100'):return engine.quote(company,*route,profile)
    def test_preview_commit_route_isolation_persistence_and_remove(self):
        raw=table();p=self.preview(raw);self.assertEqual(p.status_code,200,p.text);data=p.json()
        self.assertEqual(len(data['rows']),2);self.assertIsNone(self.quote()['price'])
        self.assertEqual(self.commit(data).status_code,200)
        q=self.quote();self.assertEqual(q['price'],1234);self.assertTrue(q['uploaded']);self.assertFalse(q['online'])
        self.assertEqual(q['freshness'],'user_document');self.assertNotIn('LIVE:',q['message'])
        for company,route in [('Мейджик',ROUTE),('Werner',ROUTE[::-1]),('Werner',('Казань','Самара'))]:self.assertIsNone(self.quote(company,route)['price'])
        engine._cached_json.cache_clear();self.assertEqual(self.quote()['price'],1234)
        self.assertEqual(self.client.get('/api/import-file/'+q['source_file']).content,raw)
        # Automatic source selection prefers fresh online evidence; a pinned file does not.
        from app.price_library import select_document
        select_document('Werner',*ROUTE,None)
        aid=engine.begin_live_attempt('Werner',*ROUTE)
        engine.save_live_update('Werner',*ROUTE,{'w100':{'kind':'exact','price':2000}}, {'origin':ROUTE[0],'destination':ROUTE[1]},aid)
        engine.finish_live_attempt('Werner',*ROUTE,aid,rows=1)
        self.assertEqual(self.quote()['price'],2000);self.assertTrue(self.quote()['online'])
        result=self.client.get('/api/compare',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'Werner'}).json()
        self.assertEqual(result['online_count'],1);self.assertEqual(result['last_good_count'],0);self.assertEqual(result['imported_count'],0)
        self.assertEqual(imports.pack(*ROUTE)['profiles']['w100']['Werner']['price'],1234)
        self.assertEqual(self.client.delete('/api/import',params={'company':'Werner','origin':ROUTE[0],'destination':ROUTE[1]}).status_code,200)
        self.assertEqual(self.quote()['price'],2000);self.assertTrue(self.quote()['online'])
    def test_replacement_removes_previous_imported_weights(self):
        self.commit(self.preview().json());self.commit(self.preview(table(rows=((100,1500),))).json())
        self.assertEqual(self.quote()['price'],1500);self.assertIsNone(self.quote(profile='w200')['price'])
    def test_all_17_companies_accept_explicit_template(self):
        for company in engine.COMPANIES:
            with self.subTest(company=company):
                p=self.preview(table(company=company),company=company);self.assertEqual(p.status_code,200,p.text)
                self.assertEqual(self.commit(p.json()).status_code,200);self.assertEqual(self.quote(company)['price'],1234)
    def test_wrong_route_company_duplicate_invalid_values_rejected(self):
        for raw in [table(origin=ROUTE[1],destination=ROUTE[0]),table(company='ДЛ'),table(rows=((100,1),(100,2))),table(rows=((100,0),)),table(rows=((100,-1),)),table(rows=((100,'=1+2'),)),table(rows=((123,500),))]:
            self.assertEqual(self.preview(raw).status_code,400)
            self.assertIsNone(self.quote()['price'])
        p=self.preview(table(rows=((100,None),(200,456))))
        self.assertEqual(p.status_code,200);self.assertEqual(len(p.json()['rows']),1)
    def test_expired_reused_and_tampered_preview_rejected(self):
        data=self.preview().json();p=imports._token_path(data['token']);saved=json.loads(p.read_text());saved['meta']['created_at']='2000-01-01T00:00:00+00:00';p.write_text(json.dumps(saved))
        self.assertEqual(self.commit(data).status_code,400)
        data=self.preview().json();imports._token_path(data['token']).with_suffix('.xlsx').write_bytes(b'tampered')
        self.assertEqual(self.commit(data).status_code,400)
        data=self.preview().json();self.assertEqual(self.commit(data).status_code,200);self.assertEqual(self.commit(data).status_code,400)
        self.assertEqual(self.client.post('/api/import/commit',json={'token':'../../settings'}).status_code,400)
    def test_empty_wrong_extension_oversized_zip_and_scan(self):
        for raw,name in [(b'', 'x.pdf'),(b'<html>401</html>','x.pdf'),(b'MZ','x.exe'),(b'PKbad','x.xlsx')]:self.assertEqual(self.preview(raw,name).status_code,400)
        with patch.object(docs,'MAX_BYTES',5):self.assertEqual(self.preview(b'%PDF-too-big','x.pdf').status_code,400)
        z=io.BytesIO()
        with zipfile.ZipFile(z,'w') as out:out.writestr('../outside.xls',b'bad')
        self.assertEqual(self.preview(z.getvalue(),'x.zip','Возовоз').status_code,400)
        from pypdf import PdfWriter
        out=io.BytesIO();pdf=PdfWriter();pdf.add_blank_page(width=200,height=200);pdf.write(out)
        response=self.preview(out.getvalue(),'scan.pdf');self.assertEqual(response.status_code,400);self.assertIn('текстового слоя',response.text)
    def test_export_matches_filters_provenance_and_keeps_filename_as_text(self):
        self.commit(self.preview(name='=SUM(1).xlsx').json())
        for include,expected in [(True,1234),(False,None)]:
            response=self.client.get('/api/export/excel',params={'layout':'matrix','origin':ROUTE[0],'destination':ROUTE[1],'companies':'Werner','live_only':True,'include_imports':include})
            wb=load_workbook(io.BytesIO(response.content),data_only=False)
            prices=[r for r in wb['Тарифы'].values if r[0]=='до 100 кг'];self.assertEqual(prices[0][3],expected)
            audit=[r for r in wb['Источники'].values if r[1]=='до 100 кг'][0]
            self.assertEqual(audit[2],'Файл пользователя');self.assertEqual(audit[7],'=SUM(1).xlsx')
            self.assertTrue(all(c.data_type!='f' for row in wb['Источники'] for c in row));wb.close()
    def test_template_has_no_sample_prices(self):
        response=self.client.get('/api/import/template',params={'company':'ДЛ','origin':ROUTE[0],'destination':ROUTE[1]})
        w=load_workbook(io.BytesIO(response.content),data_only=True)
        self.assertEqual(w.active.max_row,30);self.assertTrue(all(r[4] is None for r in list(w.active.values)[1:]));w.close()


class NativeDocuments(unittest.TestCase):
    def test_dellin_first_table_origin_and_boundaries(self):
        raw=(BASE/'tests/fixtures/dellin_original_example.pdf').read_bytes()
        vals,meta=docs.parse_document(raw,'dl.pdf','ДЛ','Москва','Санкт-Петербург')
        self.assertEqual(len(vals),28);self.assertEqual(vals['w100']['price'],1480);self.assertEqual(meta['source_page'],4)
        self.assertNotIn('w20000',vals)
        regional,proof=docs.parse_dellin(raw,'Москва','Петропавловск-Камчатский')
        self.assertEqual(regional['w100']['price'],6010);self.assertEqual(proof['source_page'],4)
        for company,origin,dest in [('ДЛ','Санкт-Петербург','Москва'),('Возовоз','Москва','Санкт-Петербург'),('ДЛ','Москва','Несуществующий город')]:
            with self.assertRaises(ValueError):docs.parse_document(raw,'dl.pdf',company,origin,dest)
    def test_vozovoz_regular_sheet_units_fixed_bounds_rounding(self):
        raw=(BASE/'tests/fixtures/vozovoz_original_example.xls').read_bytes()
        vals,meta=docs.parse_document(raw,'voz.xls','Возовоз','Москва','Санкт-Петербург')
        self.assertEqual(len(vals),29);self.assertEqual(vals['w100']['price'],1860);self.assertEqual(vals['w003']['price'],720)
        self.assertEqual(vals['w040']['price'],1100);self.assertEqual(vals['min']['price'],1080)
        self.assertEqual(meta['document_date'],'2026-09-01')
        with self.assertRaises(ValueError):docs.parse_vozovoz(raw,'Санкт-Петербург','Москва')
        for count in (1,2):
            z=io.BytesIO()
            with zipfile.ZipFile(z,'w') as out:
                for n in range(count):out.writestr(f'Возовоз_{n}.xls',raw)
                out.writestr('readme.txt','ignored')
            if count==1:self.assertEqual(docs.parse_document(z.getvalue(),'voz.zip','Возовоз','Москва','Санкт-Петербург')[0],vals)
            else:
                with self.assertRaises(ValueError):docs.parse_vozovoz_zip(z.getvalue(),'Москва','Санкт-Петербург')
    def test_other_native_workbooks_and_kit_pdf(self):
        for company,file,expected in [('ПЭК','pek_cities.xlsx',2070),('Фортуна','fortuna_cities.xlsx',4260),('ЭкспедицияПлюс','expedition_cities.xlsx',3778.12),('КИТ','kit_kazan_excerpt.pdf',1990)]:
            with self.subTest(company=company):
                vals,meta=docs.parse_document((FIX/file).read_bytes(),file,company,*ROUTE)
                self.assertEqual(vals['w100']['price'],expected);self.assertTrue(meta['route_verified'])
    def test_generic_text_pdf_checks_company_route_and_decimal_price(self):
        raw=(FIX/'generic_import.pdf').read_bytes()
        vals,meta=docs.parse_document(raw,'table.pdf','Werner',*ROUTE)
        self.assertEqual(vals['w100']['price'],1234.5);self.assertEqual(vals['min']['price'],700)
        self.assertEqual(meta['source_pages'],[1]);self.assertNotIn('w300',vals)
        for company,route in [('ДЛ',ROUTE),('Werner',ROUTE[::-1])]:
            with self.assertRaises(ValueError):docs.parse_document(raw,'table.pdf',company,*route)
    def test_future_date_does_not_become_current(self):
        with tempfile.TemporaryDirectory() as tmp,patch.object(engine,'RUNTIME_DIR',Path(tmp)),patch.object(imports,'parse_document',return_value=({'w100':{'kind':'exact','price':500}},{'document_date':'2099-01-01'})):
            with self.assertRaisesRegex(ValueError,'будущая дата'):imports.preview(b'example','file.xlsx','ДЛ',*ROUTE)
    def test_vozovoz_city_ids_resolve_catalog_not_hardcoded_corridor(self):
        data=[{'name':1,'guid':2},'Казань','f2a303ff-0124-11e5-80c7-00155d903d03',{'name':4,'guid':5},'Екатеринбург','17e84adf-0128-11e5-80c7-00155d903d03']
        html=('<script id="__NUXT_DATA__" type="application/json">'+json.dumps(data)+'</script>').encode()
        with patch.object(official.carrier_catalogs,'catalog_bytes',return_value=html):
            self.assertEqual(official.vozovoz_city_ids(*ROUTE),(data[2],data[5]))
            with self.assertRaises(ValueError):official.vozovoz_city_ids('Самара',ROUTE[1])
    def test_dellin_online_download_uses_current_link_and_validates_origin(self):
        raw=(BASE/'tests/fixtures/dellin_original_example.pdf').read_bytes()
        with patch.object(official,'fetch',return_value=(raw,{})) as fetch:
            vals,meta=official.collect('ДЛ','Москва','Санкт-Петербург')
            self.assertEqual(len(vals),28);self.assertIn('city=3&is_region=0&is_future=0',fetch.call_args.args[3])
            with self.assertRaises(ValueError):official.collect('ДЛ','Санкт-Петербург','Москва')


if __name__=='__main__':unittest.main()
