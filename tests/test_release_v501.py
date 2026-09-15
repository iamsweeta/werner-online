"""Regressions for standalone exports and the public Vozovoz document flow."""
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e, v42_main as main, official_documents as docs
from app.network import Download, RejectedResponse

ROUTE=('Москва','Санкт-Петербург')
RAW=(Path(__file__).parent/'fixtures/vozovoz_original_example.xls').read_bytes()
URL='https://files.vozovoz.ru/tariff/regular.xls'


class VozovozRecovery(unittest.TestCase):
    def collect(self, downloads, urls=None):
        with patch.object(docs,'vozovoz_city_ids',return_value=('origin-id','destination-id')), \
             patch.object(docs,'vozovoz_generate',return_value=urls or [URL,'https://files.vozovoz.ru/tariff/extra.xls']), \
             patch.object(docs,'fetch',side_effect=downloads):
            return docs.collect('Возовоз',*ROUTE)

    def test_failed_extra_download_preserves_verified_route_in_either_order(self):
        for values in [(OSError('extra unavailable'),(RAW,{'source_url':URL})),
                       ((RAW,{'source_url':URL}),OSError('extra unavailable'))]:
            result,meta=self.collect(values)
            self.assertEqual(len(result),29)
            self.assertEqual(result['w100']['price'],1860)
            self.assertTrue(meta['route_verified'])
            self.assertEqual(meta['source_url'],URL)
            self.assertEqual(meta['additional_document_errors'],['extra unavailable'])

    def test_duplicate_identical_files_do_not_hide_valid_prices(self):
        values,_=self.collect([(RAW,{}),(RAW,{})])
        self.assertEqual(len(values),29)

    def test_conflicting_files_are_rejected(self):
        with patch.object(docs,'parse_document',side_effect=[({'w100':{'price':1000}},{}),({'w100':{'price':2000}},{})]):
            with self.assertRaisesRegex(ValueError,'разные тарифы'):
                self.collect([(RAW,{}),(RAW,{})])

    def test_all_failed_or_future_documents_do_not_create_prices(self):
        with self.assertRaisesRegex(ValueError,'не получен прайс'):
            self.collect([OSError('offline'),ValueError('bad format')])
        with patch.object(docs,'parse_document',return_value=({'w100':{'price':1}},{'document_date':'2099-01-01'})):
            with self.assertRaisesRegex(ValueError,'будущие тарифы'):
                self.collect([(RAW,{})],[URL])

    def test_generator_tls_fallback_keeps_route_and_deduplicates_links(self):
        reply=Download(json.dumps(['/tariff/a.xls','/tariff/a.xls']).encode(),URL,'requests','application/json',200)
        with patch.object(docs,'_vozovoz_generate_reply',side_effect=[OSError('TLS handshake failed'),reply]) as call:
            self.assertEqual(docs.vozovoz_generate('from-id','to-id'),['https://files.vozovoz.ru/tariff/a.xls'])
            self.assertEqual(call.call_count,2)
            self.assertEqual(call.call_args.args[0],{'from':['from-id'],'to':['to-id'],'orgId':''})

    def test_access_denials_are_not_retried_or_bypassed(self):
        with patch.object(docs,'_vozovoz_generate_reply',side_effect=OSError('CONNECT tunnel failed, response 403')) as call:
            with self.assertRaisesRegex(ValueError,'доступ.*отклонён'):
                docs.vozovoz_generate('from-id','to-id')
            self.assertEqual(call.call_count,1)
        reply=Download(b'denied',URL,'requests','text/plain',403)
        with patch.object(docs,'_vozovoz_generate_reply',return_value=reply) as call:
            with self.assertRaises(RejectedResponse):docs.vozovoz_generate('from-id','to-id')
            self.assertEqual(call.call_count,1)


class RouteFile(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.ctx=patch.object(e,'RUNTIME_DIR',Path(self.tmp.name));self.ctx.start()
        self.client=TestClient(main.app)
        self.save(ROUTE,600,3090)
        self.save(ROUTE[::-1],999,99999)

    def tearDown(self):
        self.ctx.stop();self.tmp.cleanup()

    def save(self,route,minimum,total):
        attempt=e.begin_live_attempt('Возовоз',*route)
        e.save_live_update('Возовоз',*route,{'min':{'kind':'exact','price':minimum},'w100':{'kind':'exact','price':total}},
                          {'origin':route[0],'destination':route[1],'source_url':URL},attempt)
        e.finish_live_attempt('Возовоз',*route,attempt,rows=2)

    def download(self,**params):
        return self.client.get('/api/export/route',params={'origin':ROUTE[0],'destination':ROUTE[1],
                               'companies':'Возовоз','live_only':True,**params})

    def test_separate_file_contains_only_selected_route_and_company(self):
        with patch('app.excel_export.export_bytes',side_effect=AssertionError('Full book must not be loaded')):
            response=self.download()
        self.assertEqual(response.status_code,200)
        self.assertIn('attachment;',response.headers['content-disposition'])
        self.assertIn('filename*=UTF-8',response.headers['content-disposition'])
        wb=load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames,['Маршрут','Источники'])
        ws=wb['Маршрут'];self.assertEqual(ws['B1'].value,'Москва → Санкт-Петербург')
        self.assertEqual(ws['D5'].value,'Возовоз');self.assertEqual(ws.max_column,4)
        self.assertEqual(ws.max_row,len(e.COMMON_PROFILES)+5)
        for n,p in enumerate(e.COMMON_PROFILES,6):
            if p['id']=='min':self.assertEqual(ws.cell(n,4).value,600)
            elif p['id']=='w100':self.assertEqual(ws.cell(n,4).value,3090)
            else:self.assertIsNone(ws.cell(n,4).value)
        self.assertNotIn(99999,[c.value for row in ws for c in row])
        self.assertEqual(ws.freeze_panes,'D6');wb.close()

    def test_unit_selection_and_minimum_currency_are_preserved(self):
        wb=load_workbook(io.BytesIO(self.download(view='per_kg').content));ws=wb.active
        for n,p in enumerate(e.COMMON_PROFILES,6):
            if p['id']=='w100':
                self.assertEqual(ws.cell(n,4).value,30.9);self.assertEqual(ws.cell(n,3).value,'₽/кг')
            if p['id']=='min':
                self.assertEqual(ws.cell(n,4).value,600);self.assertEqual(ws.cell(n,3).value,'₽')
        wb.close()
        self.assertEqual(self.download(view='invalid').status_code,400)


if __name__=='__main__':unittest.main()
