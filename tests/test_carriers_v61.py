from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch
from bs4 import BeautifulSoup
from openpyxl import Workbook
from app import online_tariffs as a,pek,carrier_catalogs as catalogs,v42_collectors as c

FIXTURE=Path(__file__).parent/'fixtures/bsk_route_v61.html'
class BskRouteTests(TestCase):
    def test_official_form_prices_not_terminal_list(self):
        values,meta=a.parse_bsk_route(FIXTURE.read_bytes(),'Москва','Санкт-Петербург')
        self.assertEqual(len(values),29);self.assertEqual(values['w100']['price'],2850)
        self.assertEqual(values['w001']['price'],850);self.assertEqual(values['min']['price'],850)
        self.assertEqual(values['w20000']['rate_per_kg'],3639);self.assertIn('Проверьте у БСК',values['w20000']['calculation_basis'])
        self.assertIn('Воронеж',meta['calculation_basis'])
    def test_reverse_duplicate_and_wrong_units_rejected(self):
        raw=FIXTURE.read_text()
        with self.assertRaises(ValueError):a.parse_bsk_route(raw,'Санкт-Петербург','Москва')
        soup=BeautifulSoup(raw,'lxml');soup.body.append(BeautifulSoup(str(soup.h4.parent),'lxml'))
        with self.assertRaises(ValueError):a.parse_bsk_route(str(soup),'Москва','Санкт-Петербург')
        with self.assertRaises(ValueError):a.parse_bsk_route(raw.replace('стоимость за 1кг ₽','стоимость за 1кг USD'),'Москва','Санкт-Петербург')
    def test_dynamic_city_ids_no_price_constants(self):
        raw=FIXTURE.read_bytes().replace(b'value="75"',b'value="975"').replace(b'value="105"',b'value="9105"')
        with patch.object(catalogs,'catalog_bytes',return_value=raw):
            url=catalogs.bsk_route_url('Москва','Санкт-Петербург')
        self.assertIn('ship_city%5B%5D=975',url);self.assertIn('dest_city%5B%5D=9105',url)
    def test_empty_weight_not_volume_or_interpolation(self):
        soup=BeautifulSoup(FIXTURE.read_bytes(),'lxml');soup.select('tbody tr th')[2].string='—'
        values,_=a.parse_bsk_route(str(soup),'Москва','Санкт-Петербург')
        self.assertNotIn('w100',values);self.assertIn('w200',values)

class PekReaderTests(TestCase):
    def setUp(self):
        self.tmp=TemporaryDirectory();self.path=Path(self.tmp.name)/'tariff.xlsx';pek._CACHE.clear()
    def tearDown(self):self.tmp.cleanup()
    def write(self,duplicate=False,currency='руб',price=600):
        wb=Workbook();ws=wb.active;ws.title='Перевозка'
        ws.cell(20,4,'Фиксированная стоимость, руб');ws.cell(20,9,'Минимум, '+currency);ws.cell(20,10,'За 1 кг, руб')
        ws.cell(21,2,'Город отправитель');ws.cell(21,3,'Город получатель')
        for i,title in enumerate(['до 5 кг','до 10 кг','до 15 кг','до 20 кг','до 50 кг'],4):ws.cell(21,i,title)
        for i,title in enumerate(['до 100 кг','до 200 кг','до 300 кг','до 400 кг','до 500 кг','до 750 кг','до 1000 кг','до 1500 кг','до 2000 кг','до 3000 кг','до 5000 кг','до 20000 кг'],10):ws.cell(21,i,title)
        ws.append([None,'Москва','Санкт-Петербург']+[price]*5+[1400]+[10]*12)
        ws.append([None,'Санкт-Петербург','Москва']+[700]*5+[900]+[12]*12)
        if duplicate:ws.append([None,'Москва','Санкт-Петербург']+[800]*5+[900]+[12]*12)
        wb.save(self.path);wb.close()
    def test_route_fixed_and_floor_and_file_change(self):
        self.write();v,m=pek.parse_route(self.path,'Москва','Санкт-Петербург')
        self.assertEqual(v['w001']['price'],600);self.assertEqual(v['w005']['price'],600)
        self.assertEqual(v['w100']['price'],1400);self.assertEqual(v['w100']['rate_per_kg'],10)
        self.assertEqual(v['w20000']['price'],200000);self.assertEqual(m['source_row'],22)
        reverse,_=pek.parse_route(self.path,'Санкт-Петербург','Москва');self.assertEqual(reverse['w100']['price'],1200)
        self.write(price=850);v,_=pek.parse_route(self.path,'Москва','Санкт-Петербург');self.assertEqual(v['w001']['price'],850)
    def test_unknown_duplicate_wrong_currency(self):
        self.write()
        with self.assertRaises(ValueError):pek.parse_route(self.path,'Москва','Казань')
        self.write(duplicate=True)
        with self.assertRaises(ValueError):pek.parse_route(self.path,'Москва','Санкт-Петербург')
        self.write(currency='USD')
        with self.assertRaises(ValueError):pek.parse_route(self.path,'Москва','Санкт-Петербург')
    def test_online_skips_generic_bulk_extraction(self):
        self.write()
        from app import legacy
        with patch.object(c,'_bounded_source_download',return_value={'status':'downloaded','file':self.path.name,'path':str(self.path),'freshness':'live'}),patch.object(legacy,'extract_file') as extract:
            result=c._refresh_one_document_source({'id':'PEK004','company':'ПЭК'})
        extract.assert_not_called();self.assertEqual(result['status']['status'],'downloaded')
