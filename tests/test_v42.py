from __future__ import annotations

import tempfile, time, unittest
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import v42_engine as e
from app import v42_collectors as col
from app import v42_main as main


class V42Tests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.old={k:v['live'] for k,v in e.ROUTE_CONFIG.items()}
        e.ROUTE_CONFIG[('Санкт-Петербург','Москва')]['live']=Path(self.tmp.name)/'spb.json'
        e.ROUTE_CONFIG[('Москва','Санкт-Петербург')]['live']=Path(self.tmp.name)/'msk.json'
        main.COLLECT_JOBS.clear(); main.EXACT_REVISION=0
        self.client=TestClient(main.app)

    def tearDown(self):
        for k,p in self.old.items(): e.ROUTE_CONFIG[k]['live']=p
        self.tmp.cleanup()

    def test_health_and_two_routes(self):
        d=self.client.get('/health').json()
        self.assertEqual(d['version'],'61.0')
        self.assertEqual(d['engine'],'v51_verified_sources')
        self.assertEqual(d['port_hint'],8423)
        self.assertGreater(d['cities_count'],200)

    def test_options_declares_all_17_live_adapters(self):
        d=self.client.get('/api/options',params={'origin':'Санкт-Петербург'}).json()
        self.assertEqual(len(d['companies']),17)
        self.assertEqual(set(d['online_policy']['live_adapters']),set(e.COMPANIES))
        self.assertIn('LAST GOOD',d['online_policy']['fallback'])
        self.assertTrue(all(x.get('refresh_status')=='not_run' for x in d['integrations']))

    def test_fresh_install_has_no_embedded_prices(self):
        for o,d,counts in [
            ('Санкт-Петербург','Москва',(0,0,17)),
            ('Москва','Санкт-Петербург',(0,0,17)),
        ]:
            x=self.client.get('/api/compare',params={'origin':o,'destination':d,'profile':'w100'}).json()
            self.assertEqual((x['exact_count'],x['lower_bound_count'],x['missing_count']),counts)
            self.assertEqual(x['online_exact_count'],0)
            self.assertFalse(any(item.get('online') for item in x['items']))
            self.assertTrue(all(item.get('price') is None for item in x['items']))

    def test_freshness_is_per_profile_and_failed_attempt_demotes_old_live(self):
        route=('Москва','Санкт-Петербург')
        aid=e.begin_live_attempt('Werner',*route)
        e.save_live_update('Werner',*route,{'w100':{'kind':'exact','price':1234}},
                           {'source_type':'test live','source_url':'https://example.test','captured_at':e._now()},attempt_id=aid)
        e.finish_live_attempt('Werner',*route,aid,rows=1)
        q100=e.quote('Werner',*route,'w100'); q200=e.quote('Werner',*route,'w200')
        self.assertTrue(q100['online']); self.assertEqual(q100['comparison_value'],1234.0)
        self.assertEqual(q100['freshness'],'online_live')
        self.assertFalse(q200['online']); self.assertEqual(q200['freshness'],'missing')

        aid2=e.begin_live_attempt('Werner',*route)
        e.finish_live_attempt('Werner',*route,aid2,rows=0,error='network failed')
        stale=e.quote('Werner',*route,'w100')
        self.assertFalse(stale['online'])
        self.assertEqual(stale['comparison_value'],1234.0)
        self.assertEqual(stale['freshness'],'last_good_live')
        self.assertEqual(stale['refresh_status'],'failed')
        self.assertIn('Сохранённая цена',stale['message'])

    def test_live_updates_are_route_scoped(self):
        aid=e.begin_live_attempt('CTSgroup','Санкт-Петербург','Москва')
        e.save_live_update('CTSgroup','Санкт-Петербург','Москва',{'w100':{'kind':'exact','price':1111}},
                           {'source_type':'test','source_url':'x'},attempt_id=aid)
        e.finish_live_attempt('CTSgroup','Санкт-Петербург','Москва',aid,rows=1)
        self.assertEqual(e.quote('CTSgroup','Санкт-Петербург','Москва','w100')['comparison_value'],1111.0)
        self.assertEqual(e.quote('CTSgroup','Москва','Санкт-Петербург','w100')['comparison_value'],None)

    def test_run_collect_passes_every_selected_company(self):
        selected=list(e.COMPANIES)
        key='Санкт-Петербург|Москва'
        main.COLLECT_JOBS[key]={'status':'queued','profile':'w100'}
        fake=[{'company':c,'ok':True,'rows':1,'message':'ok'} for c in selected]
        with patch.object(main,'collect_selected',return_value=fake) as fn:
            main._run_collect(key,selected,'Санкт-Петербург','Москва')
        self.assertEqual(fn.call_args.args[0],selected)
        self.assertEqual(main.COLLECT_JOBS[key]['targets'],selected)
        self.assertEqual(main.COLLECT_JOBS[key]['status'],'done')

    def test_collect_selected_has_adapter_path_for_all_17(self):
        # Keep this test network-free: simulate fresh document rows, calculator rows and the two dedicated collectors.
        doc_vals={'w100':{'kind':'exact','price':1000.0}}
        meta={'source_type':'fresh official','source_url':'https://example.test','captured_at':e._now()}
        from app import online_tariffs
        with patch.object(online_tariffs,'collect',return_value=(doc_vals,meta)), \
             patch.object(online_tariffs,'fetch',return_value=('<html>Запросить стоимость</html>'.encode(),meta)), \
             patch.object(col,'_collect_document_batch',return_value={}), \
             patch.object(col,'_profile_values_from_fresh_documents',return_value=(doc_vals,meta)), \
             patch.object(col,'_collect_rail_export',return_value=(doc_vals,meta)), \
             patch.object(col,'_collect_live_calculator',return_value=(doc_vals,meta)), \
             patch('app.official_documents.collect',return_value=(doc_vals,meta)), \
             patch.object(col,'collect_cts',return_value={'company':'CTSgroup','ok':True,'rows':1,'message':'ok'}), \
             patch.object(col,'collect_newline',return_value={'company':'Новая Линия','ok':True,'rows':1,'message':'ok'}):
            # Dedicated collectors normally save their rows themselves; mock that side effect here.
            original_cts=col.collect_cts; original_nl=col.collect_newline
            def cts(o,d):
                aid=e.live_company_state(o,d,'CTSgroup')['current_attempt_id']
                e.save_live_update('CTSgroup',o,d,doc_vals,meta,attempt_id=aid)
                return {'company':'CTSgroup','ok':True,'rows':1,'message':'ok'}
            def nl(o,d):
                aid=e.live_company_state(o,d,'Новая Линия')['current_attempt_id']
                e.save_live_update('Новая Линия',o,d,doc_vals,meta,attempt_id=aid)
                return {'company':'Новая Линия','ok':True,'rows':1,'message':'ok'}
            with patch.object(col,'collect_cts',side_effect=cts), patch.object(col,'collect_newline',side_effect=nl):
                rows=col.collect_selected(list(e.COMPANIES),'Санкт-Петербург','Москва','w100')
        self.assertEqual(len(rows),17)
        self.assertEqual({x['company'] for x in rows},set(e.COMPANIES))
        self.assertTrue(all(e.live_company_state('Санкт-Петербург','Москва',c).get('attempt_status')==('unavailable' if c=='Грузопоток' else 'success') for c in e.COMPANIES))

    def test_automatic_refresh_skips_recently_verified_rows(self):
        aid=e.begin_live_attempt('Werner','Москва','Санкт-Петербург')
        e.save_live_update('Werner','Москва','Санкт-Петербург',{'w100':{'kind':'exact','price':1234}}, {'source_url':'x'},attempt_id=aid)
        e.finish_live_attempt('Werner','Москва','Санкт-Петербург',aid,rows=1)
        with patch.object(main,'collect_selected') as collect:
            result=self.client.post('/api/collect',json={'origin':'Москва','destination':'Санкт-Петербург','companies':['Werner'],'force':False}).json()
        self.assertEqual(result['status'],'fresh')
        collect.assert_not_called()

    def test_bundled_fallback_can_never_be_promoted_to_live(self):
        from app import legacy as core
        from app import legacy_backend as lb
        fresh=Path(self.tmp.name)/'fresh.html'; fresh.write_text('<html></html>',encoding='utf-8')
        results={
            'sources':[{'id':'MGC001','company':'Мейджик','status':'parsed','freshness':'live','files':['fresh.html']}],
            'files':[{'source_id':'MGC001','file':'fresh.html','path':str(fresh),'freshness':'live'}],
            'rows':[], 'errors':[]
        }
        old_row={'status':'ok','price':999.0,'freshness':'collected_official_file','source_file':'bundled-old.html','source_url':'https://example.test'}
        with patch.object(core,'load_results',return_value=results), patch.object(lb,'collected_published_tariff',return_value=old_row):
            with self.assertRaisesRegex(RuntimeError,'LAST GOOD'):
                col._profile_values_from_fresh_documents('Мейджик','Санкт-Петербург','Москва')

    def test_settings_are_shared_with_live_adapters(self):
        self.assertEqual(main.SETTINGS_PATH.name,'settings.json')
        import app.legacy_backend as lb
        self.assertEqual(main.SETTINGS_PATH,lb.SETTINGS_PATH)

    def test_diagnostics_exposes_live_evidence(self):
        aid=e.begin_live_attempt('ДЛ','Санкт-Петербург','Москва')
        e.save_live_update('ДЛ','Санкт-Петербург','Москва',{'w100':{'kind':'exact','price':1550}},
                           {'source_type':'Официальный калькулятор ДЛ','source_url':'https://example.test/live','captured_at':e._now(),'transport':'edge'},attempt_id=aid)
        e.finish_live_attempt('ДЛ','Санкт-Петербург','Москва',aid,rows=1)
        d=self.client.get('/api/diagnostics',params={'origin':'Санкт-Петербург','destination':'Москва'}).json()
        row=next(x for x in d['live_evidence'] if x['company']=='ДЛ')
        self.assertTrue(row['online'])
        self.assertEqual(row['attempt_id'],aid)
        self.assertEqual(row['transport'],'edge')
        self.assertLess(e.age_seconds(row['captured_at']),5)
        self.assertEqual(row['data_origin'],'online')

    def test_excel_both_routes_and_version(self):
        for o,d in [('Санкт-Петербург','Москва'),('Москва','Санкт-Петербург')]:
            r=self.client.get('/api/export/excel',params={'layout':'matrix','origin':o,'destination':d,'view':'per_kg'})
            self.assertEqual(r.status_code,200)
            wb=load_workbook(BytesIO(r.content),data_only=True); ws=wb.active
            self.assertEqual(ws['B1'].value,f'{o} → {d}')
            self.assertEqual(ws['B2'].value,'61.0')

    def test_specialized_collectors_are_direction_specific(self):
        rates=(600.0,[(250.0,12.0),(750.0,11.0),(1250.0,10.0),(2500.0,9.0),(5000.0,8.0)],'Москва-Юг')
        pdf=Path(self.tmp.name)/'newline.pdf'; pdf.write_bytes(b'%PDF-mock')
        def bounded(source, **kwargs):
            self.assertIn('FROM=175',source['url'])
            return {'status':'downloaded','path':str(pdf),'connector':'requests','file':pdf.name,'sha256':'mock'}
        with patch('app.newline_calculator.collect_missing',return_value=({}, {}, {})), patch.object(col,'_bounded_source_download',side_effect=bounded), patch.object(col,'_newline_pdf_values',return_value=rates), patch.object(col,'_confirm_pdf_origin'), patch.object(col.carrier_catalogs,'newline_origin',return_value=('175','Санкт-Петербург')):
            aid=e.begin_live_attempt('Новая Линия','Санкт-Петербург','Москва')
            col.collect_newline('Санкт-Петербург','Москва')
            e.finish_live_attempt('Новая Линия','Санкт-Петербург','Москва',aid,rows=20)

    def test_cts_current_table_uses_requested_origin(self):
        cells=['Москва','1','300','7.5','8','8.3','8.5','8.7','9','9.2','9.4','9.6','9.8','9.9']
        path=Path(self.tmp.name)/'cts.html'
        path.write_text('<h2>Тарифы из города Санкт-Петербург</h2><table><tr>'+''.join('<td>'+v+'</td>' for v in cells)+'</tr></table>',encoding='utf-8')
        with patch.object(col,'_bounded_source_download',return_value={'status':'downloaded','path':str(path),'file':path.name,'connector':'requests'}):
            aid=e.begin_live_attempt('CTSgroup','Санкт-Петербург','Москва')
            result=col.collect_cts('Санкт-Петербург','Москва')
            e.finish_live_attempt('CTSgroup','Санкт-Петербург','Москва',aid,rows=result['rows'])
        self.assertEqual(e.quote('CTSgroup','Санкт-Петербург','Москва','w100')['price'],990.0)
        self.assertEqual(result['rows'],29)

    def test_dellin_appkey_calls_official_terminal_and_calculator_api(self):
        class R:
            def __init__(self,data): self._data=data
            def raise_for_status(self): return None
            def json(self): return self._data
        replies=[
            R({'terminals':[{'id':101,'city':'Санкт-Петербург','default':True}]}),
            R({'terminals':[{'id':202,'city':'Москва','default':True}]}),
            R({'metadata':{'status':200},'data':{'price':'1543.25'}}),
        ]
        with patch.object(col.requests,'post',side_effect=replies) as post:
            row=col._dellin_api_live('real-user-key','Санкт-Петербург','Москва',100.0,0.5)
        self.assertEqual(row['price'],1543.25)
        self.assertEqual(row['freshness'],'live')
        self.assertEqual(post.call_count,3)
        self.assertIn('/v1/public/request_terminals.json',post.call_args_list[0].args[0])
        calc_call=post.call_args_list[2]
        self.assertIn('/v2/calculator.json',calc_call.args[0])
        payload=calc_call.kwargs['json']
        self.assertEqual(payload['appkey'],'real-user-key')
        self.assertEqual(payload['delivery']['derival']['terminalID'],101)
        self.assertEqual(payload['delivery']['arrival']['terminalID'],202)

    def test_dellin_appkey_is_preferred_before_public_widget(self):
        from app import public_web, legacy_backend as lb
        api={'status':'ok','price':1600.0,'price_is_minimum':False,'source_type':'Официальный API Деловых Линий — LIVE','source_url':'https://api.dellin.ru/v2/calculator.json','freshness':'live','browser':'API'}
        with patch.object(lb,'load_settings',return_value={'dellin_appkey':'k'}), \
             patch.object(col,'_dellin_api_live',return_value=api) as api_fn, \
             patch.object(public_web,'calculate_from_public_site') as browser:
            vals,meta=col._collect_live_calculator('ДЛ','Санкт-Петербург','Москва','w100')
        api_fn.assert_called_once()
        browser.assert_not_called()
        self.assertEqual(vals['w100']['price'],1600.0)
        self.assertEqual(meta['transport'],'API')

    def test_route_minimum_parser_is_network_only_for_baikal(self):
        html="""<html><head><title>Перевозка Санкт-Петербург Москва</title></head><body><h1>Санкт-Петербург — Москва</h1><ul><li><a>в Москва от 588.00 руб.</a></li></ul></body></html>"""
        with patch.object(col,'_download_html_bytes',return_value=(html.encode('utf-8'),'requests',col._route_page_url('Байкал Сервис','Санкт-Петербург','Москва'))) as dl:
            row=col._live_route_minimum('Байкал Сервис','Санкт-Петербург','Москва')
        dl.assert_called_once()
        self.assertEqual(row['price'],588.0)
        self.assertEqual(row['freshness'],'live')
        self.assertEqual(row['browser'],'requests')

    def test_route_minimum_parser_is_network_only_for_dellin(self):
        html="""<html><head><title>Грузоперевозки Санкт-Петербург Москва — стоимость от 1026 ₽</title></head><body><h1>Санкт-Петербург — Москва</h1></body></html>"""
        with patch.object(col,'_download_html_bytes',return_value=(html.encode('utf-8'),'curl',col._route_page_url('ДЛ','Санкт-Петербург','Москва'))) as dl:
            row=col._live_route_minimum('ДЛ','Санкт-Петербург','Москва')
        dl.assert_called_once()
        self.assertEqual(row['price'],1026.0)
        self.assertEqual(row['freshness'],'live')
        self.assertEqual(row['browser'],'curl')

    def test_dellin_uses_public_calculator_not_401_pdf(self):
        from app import public_web
        live={'ok':True,'price':1480.0,'price_is_minimum':False,'source_url':'https://widgets.dellin.ru/calculator/','source_label':'Официальный виджет-калькулятор Деловых Линий','browser':'edge'}
        with patch.object(col,'_browser_enabled',return_value=True), patch.object(public_web,'calculate_from_public_site',return_value=live) as fn:
            vals,meta=col._collect_live_calculator('ДЛ','Санкт-Петербург','Москва','w100')
        self.assertEqual(vals['w100']['kind'],'exact')
        self.assertEqual(vals['w100']['price'],1480.0)
        self.assertEqual(fn.call_args.args[:3],('ДЛ','Санкт-Петербург','Москва'))
        self.assertIn('widgets.dellin.ru',meta['source_url'])

    def test_vozovoz_without_api_key_uses_route_public_calculator(self):
        from app import public_web, legacy_backend as lb
        live={'ok':True,'price':1720.0,'price_is_minimum':False,'source_url':'https://vozovoz.ru/order/create/sankt-peterburg_moskva/','source_label':'Официальный публичный калькулятор Возовоза','browser':'edge'}
        with patch.object(col,'_browser_enabled',return_value=True), patch.object(lb,'load_settings',return_value={}), \
             patch.object(lb,'calculate_vozovoz') as demo, \
             patch.object(public_web,'calculate_from_public_site',return_value=live) as fn:
            vals,meta=col._collect_live_calculator('Возовоз','Санкт-Петербург','Москва','w100')
        demo.assert_not_called()
        self.assertEqual(vals['w100']['price'],1720.0)
        self.assertTrue(fn.call_args.kwargs['route_preselected'])
        self.assertIn('sankt-peterburg_moskva',fn.call_args.kwargs['url_override'])

    def test_vozovoz_fresh_route_minimum_is_live_lower_bound(self):
        from app import public_web, legacy_backend as lb
        fail={'ok':False,'message':'calculator blocked'}
        minimum={'status':'ok','price':720.0,'price_is_minimum':True,'source_type':'Официальная маршрутная страница Возовоза — LIVE','source_url':'https://vozovoz.ru/order/create/sankt-peterburg_moskva/','freshness':'live','browser':'requests'}
        with patch.object(col,'_browser_enabled',return_value=True), patch.object(lb,'load_settings',return_value={}), \
             patch.object(public_web,'calculate_from_public_site',return_value=fail), \
             patch.object(col,'_vozovoz_live_route_minimum',return_value=minimum):
            vals,meta=col._collect_live_calculator('Возовоз','Санкт-Петербург','Москва','w100')
        self.assertEqual(vals['w100']['kind'],'lower_bound')
        self.assertEqual(vals['w100']['price'],720.0)
        self.assertIn('маршрутная',meta['source_type'])

    def test_baikal_tries_main_public_calculator_before_iframe(self):
        from app import public_web, legacy_backend as lb
        live={'ok':True,'price':3275.0,'price_is_minimum':False,'source_url':'https://www.baikalsr.ru/services/types/avtomobilnye-perevozki-gruzov/','source_label':'Официальный публичный калькулятор Байкал Сервис','browser':'edge'}
        with patch.object(col,'_browser_enabled',return_value=True), patch.object(lb,'load_settings',return_value={}), patch.object(public_web,'calculate_from_public_site',return_value=live) as fn:
            vals,meta=col._collect_live_calculator('Байкал Сервис','Санкт-Петербург','Москва','w100')
        self.assertEqual(vals['w100']['price'],3275.0)
        self.assertIn('/city/spb__moscow/',fn.call_args.kwargs['url_override'])
        self.assertIn('baikalsr.ru',meta['source_url'])


    def test_compare_counts_live_lower_bound_as_online(self):
        aid=e.begin_live_attempt('Возовоз','Санкт-Петербург','Москва')
        e.save_live_update('Возовоз','Санкт-Петербург','Москва',{'w100':{'kind':'lower_bound','price':720.0}},
                           {'source_type':'route live','source_url':'https://example.test','captured_at':e._now(),'transport':'requests'},attempt_id=aid)
        e.finish_live_attempt('Возовоз','Санкт-Петербург','Москва',aid,rows=1)
        d=self.client.get('/api/compare',params={'origin':'Санкт-Петербург','destination':'Москва','profile':'w100','companies':'Возовоз'}).json()
        self.assertEqual(d['online_count'],1)
        self.assertEqual(d['online_exact_count'],0)
        self.assertEqual(d['online_lower_bound_count'],1)

    def test_baikal_falls_back_to_fresh_route_minimum(self):
        from app import public_web, legacy_backend as lb
        fail={'ok':False,'message':'widget blocked'}
        minimum={'status':'ok','price':588.0,'price_is_minimum':True,'source_type':'Официальная маршрутная страница Байкал Сервис — LIVE','source_url':'https://www.baikalsr.ru/city/spb__moscow/','freshness':'live','browser':'requests'}
        with patch.object(lb,'load_settings',return_value={}), patch.object(public_web,'calculate_from_public_site',return_value=fail), patch.object(col,'_live_route_minimum',return_value=minimum):
            vals,meta=col._collect_live_calculator('Байкал Сервис','Санкт-Петербург','Москва','w100')
        self.assertEqual(vals['w100']['kind'],'lower_bound')
        self.assertEqual(vals['w100']['price'],588.0)

    def test_dellin_falls_back_to_fresh_route_minimum(self):
        from app import public_web
        fail={'ok':False,'message':'widget blocked'}
        minimum={'status':'ok','price':1026.0,'price_is_minimum':True,'source_type':'Официальная маршрутная страница ДЛ — LIVE','source_url':'https://www.dellin.ru/directions/gruzoperevozki-po-rossii/sankt-peterburg-moskva/','freshness':'live','browser':'requests'}
        with patch.object(public_web,'calculate_from_public_site',return_value=fail), patch.object(col,'_live_route_minimum',return_value=minimum):
            vals,meta=col._collect_live_calculator('ДЛ','Санкт-Петербург','Москва','w100')
        self.assertEqual(vals['w100']['kind'],'lower_bound')
        self.assertEqual(vals['w100']['price'],1026.0)

    def test_collect_returns_visible_job_id(self):
        with patch.object(main,'collect_selected',return_value=[]):
            r=self.client.post('/api/collect',json={'companies':['ДЛ'],'origin':'Санкт-Петербург','destination':'Москва','profile':'w100'}).json()
        self.assertTrue(str(r.get('job_id') or '').startswith('live-'))
        self.assertEqual(r.get('origin'),'Санкт-Петербург')
        self.assertEqual(r.get('destination'),'Москва')


    def test_dellin_route_page_is_non_amp(self):
        url=col._route_page_url('ДЛ','Санкт-Петербург','Москва')
        self.assertIn('/directions/gruzoperevozki-po-rossii/',url)
        self.assertNotIn('/amp/',url)

    def test_vozovoz_route_minimum_uses_multi_transport_html_downloader(self):
        html='''<html><head><title>Санкт-Петербург Москва</title></head><body><h1>Санкт-Петербург — Москва</h1><div>Цена за услугу «Перевозка между городами» от 720 ₽</div></body></html>'''
        with patch.object(col,'_download_html_bytes',return_value=(html.encode('utf-8'),'browser-edge','https://vozovoz.ru/order/create/sankt-peterburg_moskva/')) as dl:
            row=col._vozovoz_live_route_minimum('Санкт-Петербург','Москва')
        self.assertEqual(row['price'],720.0)
        self.assertEqual(row['browser'],'browser-edge')
        dl.assert_called_once()

    def test_cts_browser_adapter_no_longer_requires_two_selects_as_primary(self):
        import inspect
        from app import public_web
        src=inspect.getsource(public_web._calculate_many_cts_table)
        self.assertIn('_set_origin_control_by_text',src)
        self.assertIn('http://cts-group.ru/prices',src)

    def test_second_route_collect_is_blocked_while_first_route_runs(self):
        active_key=main._route_key('Москва','Санкт-Петербург')
        main.COLLECT_JOBS[active_key]={'status':'running','job_id':'live-active','origin':'Москва','destination':'Санкт-Петербург','created_at':'2026-09-12T01:00:00+03:00'}
        r=self.client.post('/api/collect',json={'companies':['ДЛ'],'origin':'Санкт-Петербург','destination':'Москва','profile':'w100'})
        self.assertEqual(r.status_code,200)
        d=r.json()
        self.assertEqual(d['status'],'busy')
        self.assertFalse(d['ok'])
        self.assertEqual(d['job_id'],'live-active')

    def test_launcher_uses_independent_v42_port(self):
        root=Path(__file__).resolve().parents[1]
        cmd=(root/'START_WINDOWS_V61_0.cmd').read_text(encoding='utf-8',errors='ignore')
        self.assertIn('8423',cmd)
        self.assertIn('launch.py',cmd)
        self.assertIn('app.v42_main:app',(root/'launch.py').read_text())


if __name__=='__main__': unittest.main()
