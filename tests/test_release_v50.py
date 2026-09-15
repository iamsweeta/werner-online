"""Regression evidence: official files fetched on 2026-09-15, test-only."""
import io,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from fastapi.testclient import TestClient
from app import v42_engine as e,tariff_documents as docs,price_library,official_documents,v42_main
from app.excel_export import export_bytes
from app.cities import UnpublishedTariff
FIX=Path(__file__).parent/'fixtures'
ROUTE=('Москва','Санкт-Петербург')

class CurrentDocuments(unittest.TestCase):
 def test_all_werner_published_routes_and_ranges(self):
  for name,origin in [('werner_moscow','Москва'),('werner_spb','Санкт-Петербург')]:
   raw=(FIX/(name+'_download_20260915.xlsx')).read_bytes()
   routes,errors=price_library.parse_routes(raw,'live.xlsx','Werner')
   self.assertFalse(errors);self.assertEqual(len(routes),127)
   source=load_workbook(io.BytesIO(raw),data_only=True)
   for (o,d),item in routes.items():
    self.assertEqual(o,origin);self.assertEqual(len(item['values']),28)
    for pid,v in item['values'].items():
     p=e.PROFILE_BY_ID[pid];row=int(v['source_row'].rsplit('!',1)[1]);w=p['weight_kg']
     col=2 if w<=5 else 3 if w<=20 else 4 if w<=40 else 11 if w<=99 else 10 if w<=199 else 9 if w<=499 else 8 if w<=999 else 7 if w<=1999 else 6 if w<=2999 else 5
     value=float(str(source.active.cell(row,col).value).replace(',','.'))
     self.assertEqual(v.get('rate_per_kg') if w>40 else v['price'],value)
   source.close()
 def test_fortuna_formats_rail_duplicates_and_every_published_row(self):
  raw=(FIX/'fortuna_download_20260915.xlsx').read_bytes()
  routes,errors=price_library.parse_routes(raw,'live.xlsx','Фортуна')
  self.assertEqual(len(routes),438)
  self.assertTrue(all('договорная' in x['message'] or 'только железнодорожная' in x['message'] for x in errors),errors)
  self.assertEqual(routes[('Москва','Екатеринбург')]['values']['w100']['rate_per_kg'],19.9)
  self.assertEqual(routes[('Москва','Красноярск')]['values']['w100']['rate_per_kg'],49.6)
  source=load_workbook(io.BytesIO(raw),data_only=True)
  for route,item in routes.items():
   self.assertEqual(len(item['values']),29,route)
   for pid,v in item['values'].items():
    sheet,num=v['source_row'].rsplit('!',1);row=source[sheet][int(num)];w=e.PROFILE_BY_ID[pid]['weight_kg']
    if pid=='min':expected=min(float(row[i].value) for i in range(3,8));self.assertAlmostEqual(v['price'],expected,places=6);continue
    if w<=40:
     col=4 if w<=4.99 else 5 if w<=19.9 else 6
     self.assertAlmostEqual(v['price'],float(row[col].value),places=6)
    else:
     col=8 if w>=3000 else 9 if w>=2000 else 10 if w>=1500 else 11 if w>=1000 else 12 if w>=500 else 13 if w>=200 else 14
     self.assertAlmostEqual(v['rate_per_kg'],float(row[col].value),places=6)
  source.close()
  with self.assertRaisesRegex(UnpublishedTariff,'договорная'):docs.parse_document(raw,'x.xlsx','Фортуна','Москва','Бийск')
 def test_werner_direct_file_does_not_depend_on_html_page(self):
  raw=(FIX/'werner_spb_download_20260915.xlsx').read_bytes()
  def fetch(c,o,d,url,*a,**kw):
   self.assertEqual(url,'https://wernerus.ru/prices/46_from_36_RUB.xlsx')
   return raw,{'source_url':url,'captured_at':e._now()}
  with patch.object(official_documents,'fetch',side_effect=fetch) as f:
   values,meta=official_documents.collect('Werner',*ROUTE[::-1])
   self.assertEqual(len(values),28);self.assertEqual(values['w100']['rate_per_kg'],9.4);self.assertEqual(f.call_count,1)
 def test_reader_rejects_populated_columns_outside_limit(self):
  from openpyxl import Workbook
  w=Workbook();w.active.cell(1,151,'not empty');b=io.BytesIO();w.save(b)
  with self.assertRaisesRegex(ValueError,'150 заполненных'):list(docs.workbook_rows(b.getvalue()))

class NoEmbeddedPrices(unittest.TestCase):
 def test_old_price_packs_are_never_read_and_excel_is_blank(self):
  with tempfile.TemporaryDirectory() as tmp,patch.object(e,'RUNTIME_DIR',Path(tmp)),patch.object(e,'ROUTE_CONFIG',{}):
   poison=Path(tmp)/'old.json';poison.write_text(json.dumps({'profiles':{'w100':{'Werner':{'kind':'exact','price':987654,'rate_per_kg':9876.54}}}}))
   e.ROUTE_CONFIG[ROUTE]={'slug':'test','pack':poison,'live':Path(tmp)/'live.json'}
   self.assertFalse(e._base_pack(*ROUTE)['profiles'])
   for live_only in [True,False]:
    out=export_bytes(*ROUTE,list(e.COMPANIES),live_only=live_only,all_loaded=False)
    w=load_workbook(io.BytesIO(out),data_only=True)
    for title in ['WernerNEW','ДЛ','Фортуна']:
     self.assertTrue(all(c.value is None for row in w[title].iter_rows(min_row=2,max_row=255,min_col=4,max_col=63) for c in row))
    self.assertIsNone(w['Графики']['O19'].value);self.assertIsNone(w['Графики']['E19'].value);w.close()
 def test_template_contains_no_price_data_or_formulas(self):
  w=load_workbook(e.DATA_DIR/'export_template.xlsx',data_only=False)
  for title in ['WernerNEW','ДЛ','Фортуна']:
   self.assertFalse(any(c.value is not None for row in w[title].iter_rows(min_row=2,min_col=4) for c in row))
  self.assertFalse(any(c.data_type=='f' for s in w for row in s for c in row));w.close()
 def test_two_matrix_export_modes_and_unchanged_customer_units(self):
  with tempfile.TemporaryDirectory() as tmp,patch.object(e,'RUNTIME_DIR',Path(tmp)),patch.object(e,'ROUTE_CONFIG',{}):
   aid=e.begin_live_attempt('Werner',*ROUTE)
   e.save_live_update('Werner',*ROUTE,{'w005':{'kind':'exact','price':196},'w100':{'kind':'exact','price':1500,'rate_per_kg':10,'minimum':1500}},{},aid);e.finish_live_attempt('Werner',*ROUTE,aid,rows=2)
   client=TestClient(v42_main.app)
   for view,small,large,unit in [('total',196,1500,'₽'),('per_kg',39.2,15,'₽/кг')]:
    r=client.get('/api/export/excel',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'Werner','layout':'matrix','view':view})
    w=load_workbook(io.BytesIO(r.content),data_only=True);rows=list(w['Тарифы'].values)
    self.assertEqual(next(r for r in rows if r[0]=='3–5 кг')[2:4],(unit,small));self.assertEqual(next(r for r in rows if r[0]=='до 100 кг')[2:4],(unit,large))
    self.assertEqual(next(r for r in rows if r[0]=='МИН')[2:4],('₽',196));w.close()
   w=load_workbook(io.BytesIO(export_bytes(*ROUTE,['Werner'],all_loaded=False)),data_only=True)
   self.assertEqual(w['WernerNEW']['F2'].value,196);self.assertEqual(w['WernerNEW']['N2'].value,10);self.assertIsNone(w['WernerNEW']['D2'].value);w.close()
