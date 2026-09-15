from __future__ import annotations

import io
import json
import tempfile
import threading
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

from bs4 import BeautifulSoup
from fastapi.testclient import TestClient

from app import online_tariffs as a, v42_engine as e, v42_collectors as c, v42_main as main

FIXTURES=Path(__file__).parent/'fixtures'


class OnlineParserTests(unittest.TestCase):
    def test_actual_fortuna_outbound_section_not_inbound(self):
        raw=(FIXTURES/'FRT001.xlsx').read_bytes()
        forward=a.parse_workbook('Фортуна',raw,'Москва','Санкт-Петербург')
        reverse=a.parse_workbook('Фортуна',raw,'Санкт-Петербург','Москва')
        self.assertEqual(forward['w100']['price'],1390)
        self.assertEqual(reverse['w100']['price'],1540)
        self.assertEqual(forward['w001']['price'],580)  # 1 kg belongs to 1–4.99, not <=0.99.
        self.assertEqual(len(forward),29)

    def test_expedition_chooses_origin_sheet_and_minimum(self):
        raw=(FIXTURES/'EP000.xlsx').read_bytes()
        forward=a.parse_workbook('ЭкспедицияПлюс',raw,'Москва','Санкт-Петербург')
        reverse=a.parse_workbook('ЭкспедицияПлюс',raw,'Санкт-Петербург','Москва')
        self.assertEqual(forward['w100']['price'],1495)
        self.assertEqual(reverse['w100']['price'],1509.63)
        self.assertEqual(reverse['w001']['price'],650)

    def test_bsk_weight_not_volume_and_th_cells(self):
        forward=a.parse_bsk((FIXTURES/'BSD001.html').read_bytes(),'Москва','Санкт-Петербург')
        reverse=a.parse_bsk((FIXTURES/'BSD002.html').read_bytes(),'Санкт-Петербург','Москва')
        self.assertEqual(forward['w100']['price'],1200)
        self.assertEqual(reverse['w100']['price'],1160)
        self.assertEqual(reverse['w001']['price'],450)

    def test_fastrans_fixed_and_heavy_both_directions(self):
        for tag,o,d,price,fixed in [('001','Москва','Санкт-Петербург',1470,720),('002','Санкт-Петербург','Москва',1540,750)]:
            raw=(FIXTURES/f'FT{tag}.html').read_bytes()
            vals=a.parse_fastrans(raw,o,d)
            self.assertEqual(vals['w100']['price'],price)
            self.assertEqual(vals['w005']['price'],fixed)
            self.assertIn('w20000',vals)
            with self.assertRaises(ValueError):a.parse_fastrans(raw,d,o)

    def test_proline_selects_correct_hidden_direction(self):
        raw=(FIXTURES/'PR002.html').read_bytes()
        self.assertEqual(a.parse_proline(raw,'Москва','Санкт-Петербург')['w100']['price'],1400)
        reverse=a.parse_proline(raw,'Санкт-Петербург','Москва')
        self.assertEqual(reverse['w100']['price'],1310)
        self.assertNotIn('w001',reverse)  # No guessing an unpublished small shipment minimum.

    def test_atec_excludes_address_delivery_prices(self):
        vals=a.parse_atec((FIXTURES/'AT001.html').read_bytes(),'Москва','Санкт-Петербург')
        self.assertEqual(vals['w020']['price'],490)
        self.assertEqual(vals['w050']['price'],675)
        self.assertEqual(vals['w100']['price'],1350)

    def test_magic_real_json_and_wrong_destination(self):
        raw=(FIXTURES/'magic_spb.json').read_bytes()
        vals=a.parse_magic(raw,'Москва')
        self.assertEqual(vals['w100']['price'],1550)
        self.assertEqual(vals['w001']['price'],550)
        self.assertEqual(vals['w010']['price'],750)
        with self.assertRaises(ValueError):a.parse_magic(raw,'Санкт-Петербург')

    def test_current_workbook_is_discovered_not_date_guessed(self):
        raw=b'<a href="/upload/price_new.xlsx">'+ 'Скачать общий прайс-лист'.encode()+b'</a><a href="/extra.xlsx">Extra</a>'
        self.assertEqual(a.discover_price_link(raw,'https://fte.ru/price/','Фортуна'),'https://fte.ru/upload/price_new.xlsx')
        with self.assertRaises(ValueError):a.discover_price_link(raw.replace(b'/upload/price_new.xlsx',b'https://wrong.test/file.xlsx'),'https://fte.ru/price/','Фортуна')

    def test_fetch_rejects_redirect_to_another_company(self):
        fixture=FIXTURES/'FT001.html'
        with patch.object(c,'_bounded_source_download',return_value={'status':'downloaded','path':str(fixture),'file':fixture.name,'final_url':'https://other.test/'}):
            with self.assertRaises(ValueError):a.fetch('ФастТранс','Москва','Санкт-Петербург','https://fastrans.ru/city/')

    def test_cts_rejects_missing_or_reverse_origin(self):
        for header in ['', '<h2>Тарифы из города Москва</h2>']:
            soup=BeautifulSoup(header+'<table><tr><td>Москва</td><td>1</td></tr></table>','lxml')
            with self.assertRaises(RuntimeError):c._find_route_table(soup,'Санкт-Петербург','Москва')

    def test_price_validation_and_weight_boundaries(self):
        for value in [True,False,-1,0,float('nan'),float('inf'),'401 Unauthorized']:
            with self.subTest(value=value),self.assertRaises(ValueError):a.number(value)
        v=a.values_from_tiers([(0,100,10),(101,200,8)],minimum=500)
        self.assertEqual(v['w100']['price'],1000)
        self.assertEqual(v['w200']['price'],1600)
        self.assertNotIn('w250',v)


class LifecycleTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.old={k:v['live'] for k,v in e.ROUTE_CONFIG.items()}
        for i,cfg in enumerate(e.ROUTE_CONFIG.values()):cfg['live']=Path(self.tmp.name)/f'{i}.json'
        main.COLLECT_JOBS.clear()
        self.client=TestClient(main.app)
        self.route=('Москва','Санкт-Петербург')
    def tearDown(self):
        for k,p in self.old.items():e.ROUTE_CONFIG[k]['live']=p
        main.COLLECT_JOBS.clear();self.tmp.cleanup()
    def test_invalid_inputs_fail_before_starting_a_worker(self):
        for body in [{'companies':[]},{'companies':['unknown']},{'profile':'missing'}]:
            with patch.object(main,'collect_selected') as worker:
                r=self.client.post('/api/collect',json=body)
                self.assertEqual(r.status_code,400);worker.assert_not_called()
    def test_new_weight_is_not_blocked_by_previous_weight_cooldown(self):
        aid=e.begin_live_attempt('Werner',*self.route,'w100')
        e.save_live_update('Werner',*self.route,{'w100':{'kind':'exact','price':1200}},{},aid)
        e.finish_live_attempt('Werner',*self.route,aid,rows=1)
        with patch.object(main,'threading') as thread_module:
            result=self.client.post('/api/collect',json={'origin':self.route[0],'destination':self.route[1],'companies':['Werner'],'profile':'w200','force':False}).json()
        self.assertEqual(result['status'],'queued');thread_module.Thread.assert_called_once()

    def test_old_attempt_cannot_overwrite_new_price(self):
        old=e.begin_live_attempt('Werner',*self.route);new=e.begin_live_attempt('Werner',*self.route)
        e.save_live_update('Werner',*self.route,{'w100':{'kind':'exact','price':1200}},{},new)
        with self.assertRaises(ValueError):e.save_live_update('Werner',*self.route,{'w100':{'kind':'exact','price':900}},{},old)
        e.finish_live_attempt('Werner',*self.route,new,rows=1)
        self.assertEqual(e.quote('Werner',*self.route,'w100')['price'],1200)
    def test_live_expires_and_failed_retry_preserves_last_good(self):
        aid=e.begin_live_attempt('Werner',*self.route)
        stale=(datetime.now(timezone.utc)-timedelta(hours=2)).isoformat()
        e.save_live_update('Werner',*self.route,{'w100':{'kind':'exact','price':1200}},{'captured_at':stale},aid)
        e.finish_live_attempt('Werner',*self.route,aid,rows=1)
        q=e.quote('Werner',*self.route,'w100');self.assertFalse(q['online']);self.assertEqual(q['price'],1200)
        aid=e.begin_live_attempt('Werner',*self.route);e.finish_live_attempt('Werner',*self.route,aid,rows=0,error='HTTP 503')
        q=e.quote('Werner',*self.route,'w100');self.assertEqual(q['price'],1200);self.assertEqual(q['refresh_error'],'HTTP 503')
    def test_partial_progress_visible_before_job_finishes_and_active_job_resumes(self):
        emitted=threading.Event();release=threading.Event()
        def collect(companies,o,d,profile,on_progress):
            on_progress({'company':'Werner','ok':True,'rows':1,'message':'ready'});emitted.set();release.wait(3)
            return [{'company':'Werner','ok':True,'rows':1,'message':'ready'},{'company':'ДЛ','ok':False,'rows':0,'message':'HTTP 401'}]
        with patch.object(main,'collect_selected',side_effect=collect):
            accepted=self.client.post('/api/collect',json={'companies':['Werner','ДЛ'],'origin':self.route[0],'destination':self.route[1]}).json()
            self.assertTrue(emitted.wait(2))
            job=self.client.get('/api/active-collect').json()
            self.assertEqual(job['job_id'],accepted['job_id']);self.assertEqual(job['completed_companies'],1);self.assertEqual(job['status'],'running')
            release.set()
            import time
            until=time.monotonic()+2
            while main.COLLECT_JOBS[main._route_key(*self.route)]['status']=='running' and time.monotonic()<until:time.sleep(.01)
        job=main.COLLECT_JOBS[main._route_key(*self.route)]
        self.assertEqual(job['success_count'],1);self.assertEqual(job['failed_count'],1);self.assertEqual(job['status'],'done')
    def test_export_respects_live_filter_and_includes_sources(self):
        from openpyxl import load_workbook
        params={'origin':self.route[0],'destination':self.route[1],'companies':'Werner','live_only':'true'}
        response=self.client.get('/api/export/excel',params={**params,'layout':'matrix'})
        self.assertEqual(response.status_code,200)
        workbook=load_workbook(io.BytesIO(response.content),data_only=True)
        self.assertEqual(workbook['Тарифы']['D6'].value,None)
        self.assertIn('Источники',workbook.sheetnames)
        workbook.close()

    def test_calculator_minimum_does_not_silently_calculate_100kg(self):
        with self.assertRaisesRegex(RuntimeError,'конкретный вес'):c._collect_live_calculator('ДЛ',*self.route,'min')
    def test_dellin_wrong_terminal_and_component_price_are_not_accepted(self):
        class Response:
            def __init__(self,data):self.data=data
            def raise_for_status(self):pass
            def json(self):return self.data
        with patch.object(c.requests,'post',return_value=Response({'terminals':[{'id':1,'city':'Казань'}]})):
            with self.assertRaisesRegex(RuntimeError,'терминал'):c._dellin_api_live('key',*self.route,100,.5)
        responses=[Response({'terminals':[{'id':1,'city':self.route[0]}]}),Response({'terminals':[{'id':2,'city':self.route[1]}]}),Response({'services':[{'price':50}],'errors':'No total'})]
        with patch.object(c.requests,'post',side_effect=responses):
            with self.assertRaises(RuntimeError):c._dellin_api_live('key',*self.route,100,.5)

if __name__=='__main__':unittest.main()
