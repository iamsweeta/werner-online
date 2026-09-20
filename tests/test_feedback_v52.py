import io
import json
import os
import subprocess
import sys
import tempfile
import time
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e, v42_main as main, tariff_documents as t
from app import document_imports as imports, price_library as library, bulk_refresh as bulk

ROOT=Path(__file__).resolve().parent.parent
ROUTE=('Москва','Санкт-Петербург')


class Feedback(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.patches=[patch.object(e,'RUNTIME_DIR',Path(self.tmp.name)),patch.object(e,'ROUTE_CONFIG',{})]
        for p in self.patches:p.start()
        self.client=TestClient(main.app)
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.tmp.cleanup()
    def online(self):
        stamp=(datetime.now().astimezone()-timedelta(days=2)).isoformat()
        aid=e.begin_live_attempt('ДЛ',*ROUTE)
        e.save_live_update('ДЛ',*ROUTE,{'w100':{'kind':'exact','price':1480,'rate_per_kg':14.8}},
                          {'captured_at':stamp,'source_url':'https://www.dellin.ru/pricelist_pdf/?city=3'},aid)
        e.finish_live_attempt('ДЛ',*ROUTE,aid,rows=1)
        return stamp
    def test_old_price_survives_failed_refresh_restart_and_route_excel(self):
        stamp=self.online()
        aid=e.begin_live_attempt('ДЛ',*ROUTE)
        self.assertEqual(e.quote('ДЛ',*ROUTE,'w100')['price'],1480)
        e.finish_live_attempt('ДЛ',*ROUTE,aid,rows=0,error='HTTP 401')
        e._cached_json.cache_clear()
        q=e.quote('ДЛ',*ROUTE,'w100')
        self.assertEqual(q['captured_at'],stamp);self.assertEqual(q['price'],1480)
        self.assertFalse(q['online']);self.assertTrue(q['retained_previous']);self.assertEqual(q['refresh_status'],'failed')
        response=self.client.get('/api/export/route',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'ДЛ'})
        self.assertEqual(response.status_code,200)
        wb=load_workbook(io.BytesIO(response.content),data_only=True)
        self.assertEqual(wb.sheetnames,['Маршрут','Источники'])
        self.assertTrue(any(1480 in r for r in wb['Маршрут'].values));wb.close()
        # New interpreter reads the same on-disk result; no collector is invoked.
        code="from app import v42_engine as e;import json;e.ROUTE_CONFIG={};print(json.dumps(e.quote('ДЛ','Москва','Санкт-Петербург','w100')))"
        env={**os.environ,'TARIFF_DATA_DIR':self.tmp.name}
        result=subprocess.run([sys.executable,'-c',code],cwd=ROOT,env=env,capture_output=True,text=True,check=True)
        saved=json.loads(result.stdout)
        self.assertEqual(saved['price'],1480);self.assertEqual(saved['captured_at'],stamp)
        self.assertFalse(saved['online'])
        self.assertIsNone(e.quote('ДЛ',*ROUTE,'w20000')['price'])
    def test_saved_bulk_and_failed_bulk_keep_price_without_live_label(self):
        self.online()
        for result in [{'ok':True,'snapshot':True},{'ok':False,'message':'HTTP 401'}]:
            rows=bulk.capture_company('ДЛ',*ROUTE,result)['items']
            q=next(r for r in rows if r['profile_id']=='w100')
            self.assertEqual(q['price'],1480);self.assertFalse(q['collected_online'])
    def test_on_request_retains_historical_value_and_explains_it(self):
        self.online();aid=e.begin_live_attempt('ДЛ',*ROUTE)
        e.save_live_update('ДЛ',*ROUTE,{}, {'unavailable_profiles':{'w100':'Индивидуальный расчёт'}},aid)
        e.finish_live_attempt('ДЛ',*ROUTE,aid,rows=0)
        q=e.quote('ДЛ',*ROUTE,'w100')
        self.assertEqual(q['price'],1480);self.assertFalse(q['online'])
        self.assertEqual(q['latest_availability'],'on_request');self.assertIn('Индивидуальный расчёт',q['message'])
    def test_apostrophe_pdf_upload_commit_download_and_company_isolation(self):
        raw=(ROOT/'tests/fixtures/dellin_original_example.pdf').read_bytes()
        response=self.client.post('/api/import/preview',data={'company':'ДЛ','origin':ROUTE[0],'destination':ROUTE[1]},files={'file':("pricelist.pdf'",raw,'application/octet-stream')})
        self.assertEqual(response.status_code,200,response.text)
        preview=response.json();self.assertEqual(preview['meta']['extension'],'.pdf')
        self.assertEqual(preview['meta']['original_filename'],'pricelist.pdf')
        self.assertEqual(len(preview['rows']),28)
        self.assertTrue(any('исправлено' in w for w in preview['warnings']))
        result=self.client.post('/api/import/commit',json={'token':preview['token']})
        self.assertEqual(result.status_code,200,result.text)
        q=e.quote('ДЛ',*ROUTE,'w100');self.assertEqual(q['price'],1480)
        self.assertTrue(q['uploaded']);self.assertFalse(q['online'])
        downloaded=self.client.get('/api/import-file/'+q['source_file'])
        self.assertEqual(downloaded.content,raw);self.assertIn('application/pdf',downloaded.headers['content-type'])
        self.assertIsNone(e.quote('ДЛ',*ROUTE[::-1],'w100')['price'])
        self.assertIsNone(e.quote('ПЭК',*ROUTE,'w100')['price'])
    def test_multi_document_path_accepts_quoted_pdf_and_finds_moscow_origin(self):
        raw=(ROOT/'tests/fixtures/dellin_original_example.pdf').read_bytes()
        with t.parsing_session():
            candidates=library._candidates(raw,'pricelist.pdf','ДЛ',None)
        self.assertIn(ROUTE,candidates)
        with patch.object(library,'_candidates',return_value=[ROUTE]):
            job=library.start_many([("pricelist.pdf'",raw)],'ДЛ','Москва')
            deadline=time.monotonic()+30
            while job['status']=='parsing' and time.monotonic()<deadline:
                time.sleep(.02);job=library.status(job['token'])
            self.assertEqual(job['status'],'ready',job.get('message'))
            library.commit(job['token'])
        self.assertEqual(e.quote('ДЛ',*ROUTE,'w100')['price'],1480)
    def test_quotes_do_not_bypass_content_or_extension_checks(self):
        for name in ["pricelist.pdf'",'"pricelist.PDF"','pricelist.pdf’']:
            self.assertEqual(t.normalize_filename(name).lower(),'pricelist.pdf')
            with self.assertRaisesRegex(ValueError,'не является PDF'):t.validate_file(b'<html>401</html>',name)
        with self.assertRaises(ValueError):t.validate_file(b'%PDF-fake',"pricelist.pdf.exe'")


if __name__=='__main__':unittest.main()
