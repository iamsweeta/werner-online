import io,json,tempfile,unittest
from pathlib import Path
from unittest.mock import patch
from types import SimpleNamespace
from datetime import datetime,timedelta
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e,v42_main as main,v42_collectors as collectors
from app import newline_calculator as nl
from app import document_imports as imports

ROUTE=('Казань','Екатеринбург')
FIX=Path(__file__).parent/'fixtures'


class MinimumAndExport(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.ctx=patch.object(e,'RUNTIME_DIR',Path(self.tmp.name));self.ctx.start()
        self.client=TestClient(main.app)
    def tearDown(self):self.ctx.stop();self.tmp.cleanup()
    def save(self,values,company='Пролайн',route=ROUTE,meta=None):
        aid=e.begin_live_attempt(company,*route)
        e.save_live_update(company,*route,{pid:(value if isinstance(value,dict) else {'kind':'exact','price':value}) for pid,value in values.items()},
                           {'origin':route[0],'destination':route[1],**(meta or {})},aid)
        e.finish_live_attempt(company,*route,aid,rows=len(values))
    def test_minimum_without_explicit_min_is_company_shipment_price(self):
        self.save({'w001':600,'w100':3090})
        row=e.quote('Пролайн',*ROUTE,'min')
        self.assertEqual(row['price'],600);self.assertTrue(row['online'])
        self.assertEqual(row['comparison_value'],600);self.assertIsNone(row['effective_rate_per_kg'])
        self.assertEqual(row['minimum_source_profile'],'w001')
        self.save({'min':3090},'ДЛ')
        self.assertEqual(e.quote('ДЛ',*ROUTE,'min')['comparison_value'],3090)
        self.assertIsNone(e.quote('Пролайн',*ROUTE[::-1],'min')['price'])
    def test_minimum_never_uses_cheaper_archived_price_against_live(self):
        self.save({'w001':900,'w100':3000},meta={'captured_at':(datetime.now().astimezone()-timedelta(days=1)).isoformat()})
        self.save({'w100':3090})
        row=e.quote('Пролайн',*ROUTE,'min')
        self.assertEqual(row['price'],900);self.assertFalse(row['online'])
    def test_confirmed_file_minimum_keeps_uploaded_provenance(self):
        self.save({'w001':500},meta={'captured_at':(datetime.now().astimezone()-timedelta(days=1)).isoformat()})
        imports.e._robust_json_write(imports.route_path(*ROUTE),{'profiles':{'w001':{'Пролайн':{'kind':'exact','price':600,'source_file':'test.pdf','original_filename':'price.pdf'}}},'companies':{}})
        row=e.quote('Пролайн',*ROUTE,'min')
        self.assertEqual(row['price'],600);self.assertTrue(row['uploaded']);self.assertFalse(row['online'])
        self.assertEqual(row['source_file'],'test.pdf')
    def test_default_export_preserves_template_tabs_columns_and_uses_app_values(self):
        self.save({'w001':600,'w100':{'kind':'exact','price':3090,'rate_per_kg':30.9},'w10000':{'kind':'exact','price':123000,'rate_per_kg':12.3}})
        self.save({'w001':700,'w100':4000},route=ROUTE[::-1])
        r=self.client.get('/api/export/excel',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'Пролайн','live_only':True,'all_loaded':True})
        self.assertEqual(r.status_code,200)
        wb=load_workbook(io.BytesIO(r.content))
        template=load_workbook(e.DATA_DIR/'export_template.xlsx')
        self.assertTrue(set(template.sheetnames)-{'Грузопоток'}<=set(wb.sheetnames));template.close()
        ws=wb['Пролайн']
        self.assertEqual([ws.cell(2,c).value for c in [2,3,4,13,14,31]],['Казань','Екатеринбург',600,600,30.9,12.3])
        self.assertEqual(ws['M1'].value,'МИН');self.assertIn('20000',ws['AF1'].value)
        self.assertEqual(ws['AH2'].value,None)
        self.assertTrue(any(row[1:3]==('Екатеринбург','Казань') for row in ws.values))
        self.assertEqual(wb['ДЛ']['D2'].value,None)
        self.assertEqual(wb['Графики']['B16'].value,ROUTE[0]);self.assertEqual(wb['Графики']['F16'].value,ROUTE[1])
        self.assertEqual(wb['Графики']['N27'].value,'=IF(ISNUMBER(\'рабочий\'!M12),\'рабочий\'!M12,"")')
        self.assertEqual(wb['рабочий']['G1'].data_type,'f')
        self.assertEqual(wb['склейка']['D2'].value,'Казань → Екатеринбург')
        self.assertEqual(wb.active.title,'Графики')
        self.assertGreater(len(wb['Графики']._charts),0)
        wb.close()
        cached=load_workbook(io.BytesIO(r.content),data_only=True)
        self.assertEqual(cached['Графики']['N27'].value,600)
        self.assertEqual(cached['Графики']['O27'].value,30.9)
        self.assertEqual(cached['рабочий']['G1'].value,2)
        cached.close()
    def test_export_filters_stale_values_and_does_not_copy_template_prices(self):
        self.save({'w001':12345},meta={'captured_at':(datetime.now().astimezone()-timedelta(days=1)).isoformat()})
        r=self.client.get('/api/export/excel',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'Пролайн','live_only':True,'all_loaded':False})
        wb=load_workbook(io.BytesIO(r.content),data_only=True)
        self.assertEqual(wb['Пролайн']['D2'].value,None)
        self.assertEqual(wb['Пролайн']['M2'].value,None)
        self.assertTrue(all(row[4] is None for row in list(wb['Источники'].values)[1:]));wb.close()


class NewLineMissingWeights(unittest.TestCase):
    def test_open_ended_published_column_is_retained(self):
        text='Казань    3 дн.    600 руб.    14 руб.    13 руб.    12 руб.    11 руб.    10 руб.    9 руб.    8 руб.'
        reader=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **kw:text)])
        with patch('pypdf.PdfReader',return_value=reader):
            minimum,rates,city=collectors._newline_pdf_values(b'pdf','Казань')
        self.assertEqual(rates[-1],(float('inf'),8));self.assertEqual(city,'Казань')
    def test_calculator_uses_shipping_only_and_checks_weight(self):
        row=json.loads((FIX/'newline_calc_100.json').read_bytes())
        self.assertEqual(nl.parse_reply(row,100),1551)
        with self.assertRaises(ValueError):nl.parse_reply(row,10000)
        row['result_nocourier']['shipping'][0]['code']='address'
        with self.assertRaises(ValueError):nl.parse_reply(row,100)
    def test_individual_quote_is_not_replaced_with_fabricated_number(self):
        row=json.loads((FIX/'newline_calc_20000.json').read_bytes())
        self.assertIsNone(nl.parse_reply(row,20000))
        with tempfile.TemporaryDirectory() as tmp,patch.object(e,'RUNTIME_DIR',Path(tmp)):
            aid=e.begin_live_attempt('Новая Линия',*ROUTE)
            e.save_live_update('Новая Линия',*ROUTE,{'w001':{'kind':'exact','price':600}},
                              {'unavailable_profiles':{'w20000':'Индивидуальный расчёт'}},aid)
            e.finish_live_attempt('Новая Линия',*ROUTE,aid,rows=1)
            q=e.quote('Новая Линия',*ROUTE,'w20000')
            self.assertEqual(q['display_text'],'по запросу');self.assertIsNone(q['price']);self.assertFalse(q['online'])

if __name__=='__main__':unittest.main()
