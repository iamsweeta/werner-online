import io
import json
import tempfile
import threading
import unittest
import zipfile
from datetime import datetime,timedelta
from pathlib import Path
from unittest.mock import patch
from contextlib import closing

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import bulk_refresh as b, v42_engine as e, v42_main as main


class BulkCollection(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.runtime=patch.object(e,'RUNTIME_DIR',self.root);self.runtime.start()
        self.routes=patch.object(e,'ROUTE_CONFIG',{});self.routes.start()
        self.managers=[]
    def tearDown(self):
        for m in self.managers:
            if m.worker:m.worker.join(20)
            if m.export_worker:m.export_worker.join(20)
        self.routes.stop();self.runtime.stop();self.tmp.cleanup()
    def manager(self,collector=None,**kwargs):
        m=b.BulkManager(self.root/'bulk',collector=collector or self.collector,**kwargs)
        self.managers.append(m);return m
    def save(self,company,o,d,values=None,meta=None):
        aid=e.begin_live_attempt(company,o,d)
        values=values if values is not None else {'w001':700,'w100':{'kind':'exact','price':2200,'rate_per_kg':22}}
        e.save_live_update(company,o,d,{p:(v if isinstance(v,dict) else {'kind':'exact','price':v}) for p,v in values.items()},
                          {'source_url':'https://example.invalid/test-only','origin':o,'destination':d,**(meta or {})},aid)
        e.finish_live_attempt(company,o,d,aid,rows=len(values))
        return {'company':company,'ok':True,'rows':len(values),'message':'TEST FIXTURE'}
    def collector(self,companies,o,d,profile_id,on_progress,*,full_grid):
        self.assertTrue(full_grid)
        result=[]
        for company in companies:
            row=self.save(company,o,d);on_progress(row);result.append(row)
        return result
    def finished(self,m):
        m.worker.join(15);self.assertFalse(m.worker.is_alive())
        if m.export_worker:m.export_worker.join(30);self.assertFalse(m.export_worker.is_alive())
        job=m.status();self.assertEqual(job['status'],'done',job)
        self.assertEqual(job['export_status'],'ready',job)
        return job
    def test_reference_and_all_pairs_are_real_distinct_city_names(self):
        routes=b.plan_routes()
        self.assertEqual(len(routes),254)
        self.assertIn(('Москва','Минск'),routes)
        self.assertIn(('Санкт-Петербург','Нижний Новгород'),routes)
        self.assertTrue(all(o!=d and 'город' not in (o,d) for o,d in routes))
        cities=b.city_names();plan=b.route_plan('all')
        self.assertEqual(plan['routes'],len(cities)*(len(cities)-1))
        self.assertEqual(plan['checks'],len(cities)*(len(cities)-1)*17)
        self.assertEqual(b.plan_routes('origins',['Казань','Казань'],['Казань','Уфа']),[('Казань','Уфа')])
        for scope,origins in [('unknown',None),('origins',[]),('origins',['НЕ ГОРОД'])]:
            with self.assertRaises(ValueError):b.plan_routes(scope,origins)
    def test_all_companies_are_exported_independently_of_screen_filters(self):
        m=self.manager();job=m.start('origins',['Казань'],['Уфа']);job=self.finished(m)
        self.assertEqual(job['completed_checks'],17);self.assertEqual(job['completed_routes'],1)
        self.assertEqual(job['outcomes']['partial'],17)
        with closing(load_workbook(m.download(job['job_id']),data_only=True)) as wb:
            for company in e.COMPANIES:
                from app.excel_export import SHEETS
                ws=wb[SHEETS.get(company,company)]
                self.assertEqual(ws['B2'].value,'Казань');self.assertEqual(ws['C2'].value,'Уфа')
                self.assertEqual(ws['D2'].value,700);self.assertEqual(ws['M2'].value,700)
                self.assertEqual(ws['N2'].value,22)
            self.assertIn('Сбор',wb.sheetnames);self.assertEqual(wb.active.title,'Графики')
            self.assertIn('Проверено при сборе',[r[5] for r in list(wb['Источники'].values)[1:]])
            self.assertEqual(wb['Пролайн']['E2'].value,None)
    def test_snapshot_excludes_imports_old_attempts_and_later_overwrites(self):
        o,d='Казань','Уфа';c='Пролайн'
        self.save(c,o,d,{'w100':1})
        result=self.save(c,o,d,{'w001':700})
        imported={'profiles':{'w001':{c:{'kind':'exact','price':1}}}}
        with patch('app.document_imports.pack',return_value=imported):payload=b.capture_company(c,o,d,result)
        by={i['profile_id']:i for i in payload['items']}
        self.assertEqual(by['w001']['price'],700);self.assertEqual(by['min']['price'],700)
        self.assertIsNone(by['w100']['price']);self.assertFalse(by['w001']['uploaded'])
        self.save(c,o,d,{'w001':9000})
        self.assertEqual(by['w001']['price'],700)
        failed=b.capture_company(c,o,d,{'company':c,'ok':False,'message':'Offline'})
        self.assertTrue(all(i['price'] is None for i in failed['items']))
    def test_pause_resume_checkpoints_and_route_busy_conflict(self):
        entered=threading.Event();release=threading.Event();calls=[]
        def collect(companies,o,d,*args,**kwargs):
            calls.append((o,d));entered.set();release.wait(5)
            return self.collector(companies,o,d,*args,**kwargs)
        m=self.manager(collect)
        with patch('app.excel_export.export_bytes',return_value=b'test export'):
            job=m.start('origins',['Казань'],['Уфа','Омск']);self.assertTrue(entered.wait(3))
            with self.assertRaises(b.BusyError):m.start('reference')
            self.assertEqual(m.pause(job['job_id'])['status'],'pausing')
            release.set();m.worker.join(10)
            paused=m.status();self.assertEqual(paused['status'],'paused');self.assertEqual(paused['completed_checks'],17)
            m.resume(job['job_id']);done=self.finished(m)
            self.assertEqual(done['completed_checks'],34);self.assertEqual(len(calls),2)
        with patch.object(m,'route_busy',return_value=True):
            with self.assertRaises(b.BusyError):m.start('origins',['Казань'],['Уфа'])
    def test_restart_preserves_successful_carriers_and_only_resumes_unfinished(self):
        m=self.manager()
        with patch.object(m,'_launch'):
            job=m.start('origins',['Казань'],['Уфа'])
        result=self.save('Werner','Казань','Уфа')
        m._save_result(job['job_id'],0,'Казань','Уфа',result)
        received=[]
        def collector(companies,*args,**kwargs):
            received.extend(companies);return self.collector(companies,*args,**kwargs)
        restored=self.manager(collector)
        self.assertEqual(restored.status()['status'],'paused')
        with patch('app.excel_export.export_bytes',return_value=b'test'):
            restored.resume(job['job_id']);done=self.finished(restored)
        self.assertEqual(len(received),16);self.assertNotIn('Werner',received)
        self.assertEqual(done['completed_checks'],17)
    def test_partitioned_export_lists_every_route_and_never_mixes_parts(self):
        m=self.manager();calls=[]
        def export(o,d,companies,**kwargs):
            calls.append(kwargs['routes_override']);return f'{o} -> {d}'.encode()
        with patch.object(b,'SINGLE_WORKBOOK_LIMIT',1),patch.object(b,'PART_SIZE',1),patch('app.excel_export.export_bytes',side_effect=export):
            m.start('origins',['Казань'],['Уфа','Омск']);job=self.finished(m)
        self.assertEqual(calls,[[('Казань','Уфа')],[('Казань','Омск')]])
        with zipfile.ZipFile(m.download(job['job_id'])) as z:
            self.assertEqual(len([n for n in z.namelist() if n.endswith('.xlsx')]),2)
            manifest=z.read('routes.csv').decode('utf-8-sig')
            self.assertIn('Казань;Уфа;Завершена;tariffs_001.xlsx',manifest)
            self.assertIn('Казань;Омск;Завершена;tariffs_002.xlsx',manifest)
    def test_failed_company_does_not_abort_other_companies_and_retries_only_failures(self):
        received=[];first=[True]
        def collector(companies,o,d,profile,on_progress,*,full_grid):
            received.append(list(companies));output=[]
            for c in companies:
                if first[0] and c=='ДЛ':
                    aid=e.begin_live_attempt(c,o,d);e.finish_live_attempt(c,o,d,aid,rows=0,error='offline')
                    result={'company':c,'ok':False,'rows':0,'message':'offline'}
                else:result=self.save(c,o,d,{pid:{'kind':'exact','price':max(700,e.PROFILE_BY_ID[pid]['weight_kg']*22),'rate_per_kg':22} for pid in b.WEIGHTS})
                output.append(result);on_progress(result)
            first[0]=False;return output
        m=self.manager(collector)
        with patch('app.excel_export.export_bytes',return_value=b'test'):
            m.start('origins',['Казань'],['Уфа']);job=self.finished(m)
            self.assertEqual(job['outcomes'],{'complete':16,'failed':1})
            m.resume(job['job_id'],retry=True);job=self.finished(m)
        self.assertEqual(received[1],['ДЛ']);self.assertEqual(job['outcomes'],{'complete':17})
    def test_api_validates_requests_downloads_and_blocks_overlapping_route_refresh(self):
        m=self.manager();client=TestClient(main.app)
        with patch.object(main,'BULK',m),patch.object(m,'_launch'):
            plan=client.post('/api/bulk/plan',json={'scope':'reference'})
            self.assertEqual(plan.status_code,200);self.assertEqual(plan.json()['companies'],17)
            response=client.post('/api/bulk',json={'scope':'origins','origins':['Казань'],'destinations':['Уфа']})
            self.assertEqual(response.status_code,200);ident=response.json()['job_id']
            self.assertEqual(client.post('/api/bulk',json={}).status_code,409)
            active=client.get('/api/active-collect').json();self.assertEqual(active['mode'],'bulk')
            response=client.post('/api/collect',json={'origin':'Казань','destination':'Уфа'})
            self.assertEqual(response.json()['bulk_job_id'],ident)
            self.assertEqual(client.get(f'/api/bulk/{ident}/download').status_code,400)
            self.assertEqual(client.post('/api/bulk/plan',json={'scope':'origins','origins':['x']}).status_code,400)
            self.assertEqual(client.get('/api/bulk/no-such-job').status_code,400)


class CalculatorGrid(unittest.TestCase):
    def test_lower_bound_is_fetched_once_and_is_not_copied_to_all_weights(self):
        from app.v42_collectors import _collect_calculator_grid
        with patch('app.v42_collectors._collect_live_calculator',return_value=({'w001':{'kind':'lower_bound','price':600}}, {'source_url':'https://example.invalid'})) as calc:
            values,meta=_collect_calculator_grid('Байкал Сервис','Казань','Уфа')
        self.assertEqual(calc.call_count,1);self.assertEqual(set(values),{'min'})
        self.assertEqual(values['min']['kind'],'lower_bound');self.assertTrue(meta['partial_errors'])
    def test_exact_grid_requests_each_control_weight_without_scaling(self):
        from app.v42_collectors import _collect_calculator_grid
        def calculator(c,o,d,pid,**kwargs):return {pid:{'kind':'exact','price':100+e.PROFILE_BY_ID[pid]['weight_kg']}},{'source_url':'https://example.invalid/'+pid}
        with patch('app.v42_collectors._collect_live_calculator',side_effect=calculator) as calc:
            values,meta=_collect_calculator_grid('ДЛ','Казань','Уфа')
        self.assertEqual(calc.call_count,28);self.assertEqual(len(values),28)
        self.assertEqual(values['w20000']['price'],20100)
        self.assertTrue(values['w100']['source_url'].endswith('w100'))


if __name__=='__main__':unittest.main()
