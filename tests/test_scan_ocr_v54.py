"""OCR regression prices below are assertions from a test scan, never live data."""
import io,json,os,tempfile,time,unittest
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import tariff_documents as t, scan_ocr as ocr, v42_engine as e, v42_main as main, price_library as lib
from app.scan_ocr_worker import numeric
FIX=Path(__file__).parent/'fixtures'

class ScanOCR(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw=(FIX/'dellin_scan_two_rows.pdf').read_bytes()
        # Integrated MuPDF OCR must work without tesseract / pdftoppm on PATH.
        with patch.dict(os.environ,{'PATH':'','TESSDATA_PREFIX':'/not-installed'}):
            cls.recognized=ocr.prepare(cls.raw,'ДЛ')
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.context=patch.object(e,'RUNTIME_DIR',Path(self.tmp.name));self.context.start()
        self.client=TestClient(main.app)
    def tearDown(self):self.context.stop();self.tmp.cleanup()
    def test_actual_rotated_raster_prices_date_and_tax(self):
        self.assertTrue(ocr.needed(self.raw));r=self.recognized
        self.assertEqual(r['origin'],'Санкт-Петербург');self.assertEqual(r['document_date'],'2026-08-03')
        self.assertEqual(r['tax_basis'],'С НДС (по источнику)');self.assertEqual(r['errors'],[])
        self.assertEqual([r['destination'] for r in r['rows']],['Абакан','Альметьевск'])
        v,m=ocr.parse(r,'Санкт-Петербург','Абакан')
        self.assertEqual(len(v),28);self.assertNotIn('w20000',v)
        self.assertEqual(v['w001']['price'],680);self.assertEqual(v['min']['price'],1520)
        self.assertEqual(v['w100']['rate_per_kg'],43.6);self.assertEqual(v['w100']['price'],4360)
        self.assertTrue(m['ocr']);self.assertEqual(m['source_page'],1)
    def test_inconsistent_direction_company_and_unknown_city_rejected(self):
        with self.assertRaisesRegex(ValueError,'Санкт-Петербург'):ocr.parse(self.recognized,'Москва','Абакан')
        with self.assertRaisesRegex(ValueError,'не найдена'):ocr.parse(self.recognized,'Санкт-Петербург','Москва')
        with self.assertRaisesRegex(ValueError,'OCR поддерживает'):ocr.prepare(self.raw,'Werner')
    def test_number_parser_never_guesses_missing_comma(self):
        for text in ['1410','14/10','14,1','14,IO','14,10 руб','-14,10','0,00','1e3']:
            self.assertIsNone(numeric(text,True),text)
        self.assertEqual(numeric('14,10',True),14.1);self.assertEqual(numeric('14.10',True),14.1)
        self.assertEqual(numeric('1 330'),1330)
        for text in ['01330','1 33','-680','680x','']:
            self.assertIsNone(numeric(text))
    def test_shared_library_preview_confirm_route_and_excel(self):
        # The actual OCR result above is reused to test persistence and APIs.
        with patch.object(ocr,'prepare',return_value=self.recognized):
            response=self.client.post('/api/price-documents/preview',data={'company':'ДЛ'},files={'file':("pricelist.pdf'",self.raw)})
            self.assertEqual(response.status_code,200,response.text);job=response.json();deadline=time.monotonic()+20
            while job['status']=='parsing' and time.monotonic()<deadline:
                time.sleep(.01);job=self.client.get('/api/price-documents/preview/'+job['token']).json()
            self.assertEqual(job['status'],'ready',job);self.assertEqual(job['matched'],2)
            self.assertTrue(any('OCR' in w for w in job['warnings']));self.assertTrue(job['meta']['ocr'])
            self.assertIsNone(e.quote('ДЛ','Санкт-Петербург','Абакан','w100')['price'])
            response=self.client.post('/api/price-documents/commit',json={'token':job['token']});self.assertEqual(response.status_code,200,response.text)
        saved=response.json();self.assertEqual(self.client.get('/api/import-file/'+saved['source_file']).content,self.raw)
        quote=e.quote('ДЛ','Санкт-Петербург','Абакан','w100');self.assertEqual(quote['price'],4360);self.assertTrue(quote['ocr']);self.assertFalse(quote['online'])
        self.assertIsNone(e.quote('ДЛ','Абакан','Санкт-Петербург','w100')['price'])
        response=self.client.get('/api/export/route',params={'origin':'Санкт-Петербург','destination':'Абакан','companies':'ДЛ'})
        self.assertEqual(response.status_code,200);w=load_workbook(io.BytesIO(response.content),data_only=True)
        self.assertEqual(w.sheetnames,['Маршрут','Источники']);self.assertTrue(any(4360 in r for r in w['Маршрут'].values));w.close()
    def test_single_route_preview_uses_same_ocr(self):
        with patch.object(ocr,'prepare',return_value=self.recognized):
            response=self.client.post('/api/import/preview',data={'company':'ДЛ','origin':'Санкт-Петербург','destination':'Абакан'},files={'file':('scan.pdf',self.raw)})
            self.assertEqual(response.status_code,200,response.text);p=response.json();self.assertTrue(p['meta']['ocr'])
            response=self.client.post('/api/import/commit',json={'token':p['token']});self.assertEqual(response.status_code,200)
            self.assertEqual(len(lib.list_files()),1)
    def test_session_reuses_ocr_without_second_worker(self):
        with t.parsing_session():
            t._SESSION.get()[('ocr',self.raw,'ДЛ')]=self.recognized
            with patch('app.scan_ocr.subprocess.Popen',side_effect=AssertionError('repeated OCR')):
                self.assertEqual(t.parse_document(self.raw,'scan.pdf','ДЛ','Санкт-Петербург','Абакан')[0]['w001']['price'],680)
                self.assertEqual(t.parse_document(self.raw,'scan.pdf','ДЛ','Санкт-Петербург','Альметьевск')[0]['w001']['price'],560)

if __name__=='__main__':unittest.main()
