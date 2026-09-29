"""Bulk regressions. All numbers here are synthetic test inputs, never app seeds."""
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook
from app import bulk_refresh as b, v42_engine as e, manual_prices as manual
from app import document_imports as imports
from app.excel_export import export_bytes

ROUTE = ('Москва', 'Санкт-Петербург')


class BulkRegression(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.patches = [patch.object(e, 'RUNTIME_DIR', self.root), patch.object(e, 'ROUTE_CONFIG', {})]
        for p in self.patches: p.start()
        self.m = b.BulkManager(self.root / 'bulk')

    def tearDown(self):
        for worker in (self.m.worker, self.m.export_worker):
            if worker: worker.join(60)
        for p in reversed(self.patches): p.stop()
        self.tmp.cleanup()

    def job(self, **kwargs):
        with patch.object(self.m, '_launch'):
            job = self.m.start('origins', [ROUTE[0]], [ROUTE[1]], **kwargs)
        with self.m.db() as db:
            db.execute("UPDATE jobs SET status='paused' WHERE id=?", (job['job_id'],))
        return job['job_id']

    def export(self, ident):
        self.m.prepare_export(ident)
        self.m.export_worker.join(60)
        self.assertFalse(self.m.export_worker.is_alive())
        job = self.m.status(ident)
        self.assertEqual(job['export_status'], 'ready', job)
        return job

    def live(self, values, stamp=None):
        aid = e.begin_live_attempt('Werner', *ROUTE)
        e.save_live_update('Werner', *ROUTE, values, {'captured_at': stamp} if stamp else {}, aid)
        e.finish_live_attempt('Werner', *ROUTE, aid, rows=len(values))
        return {'company': 'Werner', 'ok': True}

    def document(self):
        imports.e._robust_json_write(imports.route_path(*ROUTE), {
            'profiles': {'w100': {'Werner': {'kind': 'exact', 'price': 1200, 'rate_per_kg': 12,
                'original_filename': 'TEST_ONLY.csv', 'source_file': 'TEST_ONLY.csv'}}}, 'companies': {}})

    def test_manual_changes_invalidate_export_even_without_documents(self):
        ident = self.job(include_imports=False)
        with patch('app.excel_export.export_bytes', return_value=b'fixture'):
            self.export(ident)
            manual.put('Werner', *ROUTE, 'w100', 14.5, 'rub_per_kg')
            self.assertTrue(self.m.status(ident)['export_outdated'])
            with self.assertRaisesRegex(ValueError, 'Пересоберите'): self.m.download(ident)
            self.export(ident)
            manual.remove('Werner', *ROUTE, 'w100')
            self.assertTrue(self.m.status(ident)['export_outdated'])

    def test_bulk_fills_missing_published_rate_from_automatic_document(self):
        self.document()
        result = self.live({'w100': {'kind': 'exact', 'price': 1700}})
        ident = self.job()
        self.m._save_result(ident, 0, *ROUTE, result)
        self.export(ident)
        wb = load_workbook(self.m.download(ident), data_only=True)
        self.assertEqual(wb['WernerNEW']['N2'].value, 12)
        self.assertIsNone(wb['WernerNEW']['O2'].value)
        self.assertTrue(any(r[5] == 'Файл пользователя' and r[9] == 'TEST_ONLY.csv'
                            for r in list(wb['Источники'].values)[1:]))
        wb.close()

    def test_document_filter_excludes_document_in_full_workbook(self):
        self.document()
        raw = export_bytes(*ROUTE, ['Werner'], all_loaded=False, include_imports=False)
        wb = load_workbook(io.BytesIO(raw), data_only=True)
        self.assertIsNone(wb['WernerNEW']['N2'].value)
        self.assertFalse(any(r[5] == 'Файл пользователя' for r in list(wb['Источники'].values)[1:]))
        wb.close()

    def test_saved_report_keeps_published_online_rate_over_automatic_document(self):
        self.document()
        self.live({'w001': {'kind': 'exact', 'price': 700},
                   'w100': {'kind': 'exact', 'price': 1700, 'rate_per_kg': 17}})
        ident = self.job(mode='saved')
        rows = self.m._route_rows(ident, 0, *ROUTE, imports.pack(*ROUTE), current=True)
        item = next(i for r in rows if r['profile']['id']=='w100' for i in r['items'] if i['company']=='Werner')
        self.assertEqual(item['published_rate_per_kg'], 17)
        self.assertFalse(item.get('uploaded'))
        self.assertFalse(item['collected_online'])

    def test_collection_evidence_does_not_capture_manual_override(self):
        result = self.live({'w001': {'kind': 'exact', 'price': 600}})
        manual.put('Werner', *ROUTE, 'w001', 800, 'rub')
        payload = b.capture_company('Werner', *ROUTE, result)
        item = next(i for i in payload['items'] if i['profile_id'] == 'w001')
        self.assertEqual(item['price'], 600)
        self.assertTrue(item['collected_online'])
        self.assertFalse(item.get('manual'))

    def test_same_second_same_total_new_rate_is_not_old_collection_evidence(self):
        stamp = e._now()
        result = self.live({'w100': {'kind': 'exact', 'price': 1500, 'rate_per_kg': 10}}, stamp)
        ident = self.job()
        self.m._save_result(ident, 0, *ROUTE, result)
        self.live({'w100': {'kind': 'exact', 'price': 1500, 'rate_per_kg': 12}}, stamp)
        rows = self.m._route_rows(ident, 0, *ROUTE, {}, current=True)
        item = next(i for r in rows if r['profile']['id'] == 'w100' for i in r['items'] if i['company'] == 'Werner')
        self.assertEqual(item['published_rate_per_kg'], 12)
        self.assertFalse(item['collected_online'])

    def test_failed_refresh_invalidates_old_success_audit_without_erasing_price(self):
        self.live({'w001': {'kind': 'exact', 'price': 600}})
        ident = self.job()
        with patch('app.excel_export.export_bytes', return_value=b'fixture'):
            self.export(ident)
            aid = e.begin_live_attempt('Werner', *ROUTE)
            e.finish_live_attempt('Werner', *ROUTE, aid, rows=0, error='HTTP 502')
            self.assertTrue(self.m.status(ident)['export_outdated'])
            self.assertEqual(e.quote('Werner', *ROUTE, 'w001')['price'], 600)

    def test_ready_export_survives_moving_entire_data_directory(self):
        ident = self.job()
        with patch('app.excel_export.export_bytes', return_value=b'fixture'):
            self.export(ident)
        old = self.m.download(ident)
        # Simulate pre-62 absolute paths and an upgrade to a new installation.
        with self.m.db() as db: db.execute('UPDATE jobs SET export_file=? WHERE id=?', (str(old), ident))
        moved = self.root / 'moved' / 'bulk'
        shutil.copytree(self.m.directory, moved)
        shutil.rmtree(self.m.directory)
        restored = b.BulkManager(moved)
        self.assertEqual(restored.download(ident).read_bytes(), b'fixture')

    def test_missing_export_has_rebuild_state_instead_of_broken_download(self):
        ident = self.job()
        with patch('app.excel_export.export_bytes', return_value=b'fixture'): self.export(ident)
        self.m.download(ident).unlink()
        job = self.m.status(ident)
        self.assertIsNone(job['download_url'])
        self.assertEqual(job['export_status'], 'error')

    def test_cannot_export_an_old_job_while_another_collection_runs(self):
        ident = self.job()
        with patch.object(self.m, '_launch'): self.m.start('origins', ['Казань'], ['Уфа'])
        with self.assertRaises(b.BusyError): self.m.prepare_export(ident)

    def test_change_during_export_never_offers_obsolete_file(self):
        ident = self.job()
        def changed(*args, **kwargs):
            manual.put('Werner', *ROUTE, 'w100', 15, 'rub_per_kg')
            return b'old fixture'
        with patch('app.excel_export.export_bytes', side_effect=changed):
            job = self.export(ident)
        self.assertTrue(job['export_outdated'])
        self.assertIsNone(job['download_url'])

    def test_streamed_sources_preserve_text_and_numbers_without_formulas(self):
        self.live({'w100': {'kind': 'exact', 'price': 1500, 'rate_per_kg': 12.5,
                           'minimum': 1500, 'source_url': '=1+2',
                           'calculation_basis': 'Кавычки " & <тариф> — руб/кг'}})
        ident=self.job()
        self.export(ident)
        wb=load_workbook(self.m.download(ident),data_only=False)
        rows=[r for r in wb['Источники'].iter_rows(min_row=2) if r[0].value=='Werner' and r[14].value==12.5]
        self.assertEqual(len(rows),1)
        row=rows[0]
        self.assertEqual(row[7].value,'=1+2');self.assertEqual(row[7].data_type,'s')
        self.assertEqual(row[8].value,'Кавычки " & <тариф> — руб/кг')
        self.assertEqual(row[4].value,1500);self.assertEqual(row[14].data_type,'n')
        self.assertEqual(wb['Источники'].auto_filter.ref,f'A1:Q{wb["Источники"].max_row}')
        wb.close()


class FullMainWorkbook(unittest.TestCase):
    setUp=BulkRegression.setUp
    tearDown=BulkRegression.tearDown
    export=BulkRegression.export
    # Run separately for the full 254-route acceptance check.
    def test_full_main_saved_job_and_every_price_cell(self):
        import time
        from app.excel_export import SHEETS
        began=time.monotonic()
        routes=b.plan_routes('reference')
        for ri, (o,d) in enumerate(routes):
            profiles={};meta={}
            for ci,c in enumerate(e.COMPANIES):
                meta[c]={'current_attempt_id':'fixture','attempt_status':'success'}
                for pid,value,rate in [('w001',500+ri*17+ci,None),('w005',700+ri*17+ci,None),
                                       ('w100',100*(10+ri+ci/100),10+ri+ci/100),
                                       ('w20000',20000*(9+ri+ci/100),9+ri+ci/100)]:
                    profiles.setdefault(pid,{})[c]={'kind':'exact','price':value,'rate_per_kg':rate,
                        'attempt_id':'fixture','captured_at':'2026-01-01T00:00:00+00:00',
                        'source_url':'https://example.invalid/TEST-ONLY'}
            e._write_live(o,d,{'route':{'origin':o,'destination':d},'companies':meta,'profiles':profiles})
        def forbidden_network(*args,**kwargs):raise AssertionError('Saved reports must not use the network')
        self.m.collector=forbidden_network
        job=self.m.start('reference',mode='saved')
        self.m.worker.join(180)
        self.assertFalse(self.m.worker.is_alive())
        self.assertIsNotNone(self.m.export_worker,self.m.status())
        self.m.export_worker.join(90)
        status=self.m.status(job['job_id'])
        self.assertEqual(status['status'],'done',status)
        self.assertEqual(status['completed_checks'],254*17)
        self.assertEqual(status['completed_routes'],254)
        self.assertEqual(status['export_status'],'ready',status)
        wb=load_workbook(self.m.download(job['job_id']),data_only=True,read_only=True)
        checked=0
        for ci,c in enumerate(e.COMPANIES):
            ws=wb[SHEETS.get(c,c)]
            for ri,row in enumerate(ws.iter_rows(min_row=2,max_row=255,max_col=32,values_only=True)):
                self.assertEqual(row[1:3],routes[ri])
                for pi,p in enumerate(e.COMMON_PROFILES):
                    expected={'w001':500+ri*17+ci,'w005':700+ri*17+ci,'min':500+ri*17+ci,
                              'w100':10+ri+ci/100,'w20000':9+ri+ci/100}.get(p['id'])
                    self.assertEqual(row[pi+3],expected,(c,ri,p['id']))
                    checked+=1
        self.assertEqual(checked,254*17*29)
        self.assertEqual(wb['Графики']['E19'].value,500)
        self.assertEqual(wb['Графики']['O19'].value,10)
        self.assertIn('Источники',wb.sheetnames)
        self.assertIn('Сбор',wb.sheetnames)
        wb.close()
        print(f'FULL MAIN: {len(routes)} routes, 17 companies, {checked} price cells verified in {time.monotonic()-began:.2f}s',flush=True)


if __name__ == '__main__': unittest.main()
