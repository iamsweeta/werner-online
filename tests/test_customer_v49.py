"""Customer acceptance examples. All amounts here are test fixtures, not LIVE prices."""
import io
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from app import v42_engine as e, v42_main as main, tariff_documents as docs
from app import bulk_refresh as bulk, price_library, official_documents
from app.tariff_model import tariff_value
from app.excel_export import _value, export_bytes
from app.cities import main_cities

ROUTE=('Москва','Санкт-Петербург')

def interval_csv(rows,company='Werner',route=ROUTE):
    lines=[';'.join(docs.INTERVAL_HEADER)]
    lines+=[';'.join(map(str,[company,*route,*r])) for r in rows]
    return ('\ufeff'+'\n'.join(lines)).encode()


def werner_fixture(origin='Москва'):
    wb=Workbook();ws=wb.active
    ws.append(['Сайт: wernerus.ru'])
    ws.append(['Цены указаны в рублях, без НДС, на 10.08.2026 г.'])
    ws.append([f'Перевозка сборных грузов из города {origin}'])
    ws.append(['Пункт назначения','Фиксированные тарифы',None,None,'Стоимость, руб./кг',None,None,None,None,None,None,'Стоимость, руб./м³'])
    ws.append([None,'до 5кг\nдо 0,03м³','до 20кг\nдо 0,1м³','до 40кг\nдо 0,19м³','от 3000','2000-2999','1000-1999','500-999','200-499','100-199','41-99','от 15'])
    ws.append(['Санкт-Петербург',196,403,460,8.8,8.9,9.4,9.9,10.2,10.4,10.9,2200])
    ws.append(['Абакан',4294,4294,4294,46.3,46.5,47,48.4,49.1,50.5,61.3,11575])
    out=io.BytesIO();wb.save(out);return out.getvalue()


class CustomerTariffs(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.runtime=patch.object(e,'RUNTIME_DIR',self.root);self.runtime.start()
        self.routes=patch.object(e,'ROUTE_CONFIG',{});self.routes.start()
        self.client=TestClient(main.app)
    def tearDown(self):self.routes.stop();self.runtime.stop();self.tmp.cleanup()
    def parse(self,rows):return docs.parse_document(interval_csv(rows),'test.csv','Werner',*ROUTE)[0]
    def pack(self,values):return {'companies':{'Werner':{}},'profiles':{p:{'Werner':v} for p,v in values.items()}}
    def test_fixed_customer_examples_repeat_without_division(self):
        values=self.parse([(0,5,196,'руб',''),(5,20,403,'руб',''),(20,50,550,'руб','')])
        self.assertEqual([values[p]['price'] for p in ['w001','w003','w005']],[196]*3)
        self.assertEqual([values[p]['price'] for p in ['w010','w015','w020']],[403]*3)
        self.assertEqual([values[p]['price'] for p in ['w035','w040','w050']],[550]*3)
        self.assertNotIn('w100',values)
        for p,v in values.items():
            q=e._route_quote('Werner',*ROUTE,p,({}, {},self.pack(values)))
            self.assertEqual(tariff_value(q,e.PROFILE_BY_ID[p]),v['price'])
    def test_published_rate_survives_minimum_in_both_excel_layouts(self):
        vals=self.parse([(0,50,600,'руб',''),(50,200,10,'руб/кг',1500)])
        pack=self.pack(vals)
        with patch('app.document_imports.pack',return_value=pack):
            q=e.quote('Werner',*ROUTE,'w100')
            self.assertEqual(q['price'],1500);self.assertEqual(q['effective_rate_per_kg'],15)
            self.assertEqual(tariff_value(q,e.PROFILE_BY_ID['w100']),10)
            for layout in ['customer','matrix']:
                r=self.client.get('/api/export/excel',params={'origin':ROUTE[0],'destination':ROUTE[1],'companies':'Werner','layout':layout})
                self.assertEqual(r.status_code,200)
                wb=load_workbook(io.BytesIO(r.content),data_only=True)
                if layout=='customer':
                    self.assertEqual(wb['WernerNEW']['N2'].value,10);self.assertEqual(wb['WernerNEW']['D2'].value,600)
                    self.assertEqual(wb['WernerNEW']['M2'].value,600)
                    self.assertNotIn('Грузопоток',wb.sheetnames)
                    audit=next(r for r in wb['Источники'].values if r[3]=='до 100 кг')
                    self.assertEqual((audit[4],audit[14],audit[15]),(1500,10,1500))
                else:
                    row=next(r for r in wb['Тарифы'].values if r[0]=='до 100 кг');self.assertEqual(row[2:4],('₽',1500))
                    row=next(r for r in wb['Тарифы'].values if r[0]=='3–5 кг');self.assertEqual(row[2:4],('₽',600))
                wb.close()
    def test_calculator_total_cannot_be_presented_as_published_rate(self):
        item={'comparison_value':1500,'price':1500,'effective_rate_per_kg':15}
        self.assertIsNone(tariff_value(item,e.PROFILE_BY_ID['w100']))
        self.assertEqual(_value(item,e.PROFILE_BY_ID['w100'],False,True),None)
        self.assertEqual(tariff_value(item,e.PROFILE_BY_ID['w050']),1500)
    def test_finite_ranges_gaps_units_and_overlap(self):
        values=self.parse([(0,5,196,'руб',''),(10,20,403,'руб',''),(50,100,10,'руб/кг',600)])
        self.assertNotIn('w010',values);self.assertNotIn('w200',values)
        for rows in [[(0,10,100,'руб',''),(5,20,200,'руб','')],[(0,50,100,'USD','')],[(50,100,10,'руб/кг',0)],[(0,5,100,'руб',500)],[(5,0,100,'руб','')]]:
            with self.subTest(rows=rows),self.assertRaises(ValueError):self.parse(rows)
        self.assertIn('w20000',self.parse([(50,'∞',10,'руб/кг',600)]))
    def test_legacy_control_point_is_not_expanded_or_divided(self):
        raw='Компания;Откуда;Куда;Вес, кг;Цена, руб\nWerner;Москва;Санкт-Петербург;5;196\nWerner;Москва;Санкт-Петербург;100;1500'.encode()
        vals,_=docs.parse_document(raw,'old.csv','Werner',*ROUTE)
        self.assertEqual(set(vals),{'w005','w100'});self.assertNotIn('rate_per_kg',vals['w100'])
    def test_all_carriers_can_import_interval_template_including_reverse(self):
        for c in e.COMPANIES:
            with self.subTest(company=c):
                raw=interval_csv([(0,5,196,'руб',''),(50,100,10,'руб/кг',1500)],c,ROUTE[::-1])
                values,_=docs.parse_document(raw,'test.csv',c,*ROUTE[::-1]);self.assertEqual(values['w100']['rate_per_kg'],10)
                pairs,errors=price_library.parse_routes(raw,'test.csv',c)
                self.assertEqual(set(pairs),{ROUTE[::-1]});self.assertFalse(errors)
    def test_native_werner_detects_origin_destination_and_original_ranges(self):
        raw=werner_fixture();vals,meta=docs.parse_document(raw,'native.xlsx','Werner',*ROUTE)
        self.assertEqual([vals[p]['price'] for p in ['w001','w003','w005','w010','w015','w020']],[196,196,196,403,403,403])
        self.assertEqual(vals['w040']['price'],460);self.assertEqual(vals['w050']['price'],545)
        self.assertEqual(vals['w100']['rate_per_kg'],10.4);self.assertEqual(vals['w20000']['rate_per_kg'],8.8)
        self.assertEqual(meta['document_date'],'2026-08-10')
        routes,errors=price_library.parse_routes(raw,'native.xlsx','Werner')
        self.assertEqual(set(routes),{ROUTE,('Москва','Абакан')});self.assertFalse(errors)
        with self.assertRaises(ValueError):docs.parse_document(raw,'native.xlsx','Werner',*ROUTE[::-1])
    def test_native_werner_changed_header_is_rejected(self):
        wb=load_workbook(io.BytesIO(werner_fixture()));wb.active['K5']='50–100';out=io.BytesIO();wb.save(out)
        with self.assertRaises(ValueError):docs.parse_document(out.getvalue(),'bad.xlsx','Werner',*ROUTE)
    def test_werner_online_discovers_current_link_and_checks_future_date(self):
        url='https://wernerus.ru/prices/46_from_35_RUB.xlsx'
        page=f'<a href="{url}">Скачать</a>'.encode();raw=werner_fixture()
        def fetch(c,o,d,u,*a,**kw):return (page if u.endswith('/prices/') else raw),{'source_url':u}
        with patch.object(official_documents,'fetch',side_effect=fetch) as f:
            values,meta=official_documents.collect('Werner',*ROUTE)
            self.assertEqual(f.call_args.args[3],url);self.assertEqual(meta['source_url'],url)
            self.assertEqual(values['w005']['price'],196)
        wb=load_workbook(io.BytesIO(raw));wb.active['A2']='Цены указаны в рублях, без НДС, на 01.01.2099 г.'
        out=io.BytesIO();wb.save(out);raw=out.getvalue()
        with patch.object(official_documents,'fetch',side_effect=fetch),self.assertRaisesRegex(ValueError,'будущие'):
            official_documents.collect('Werner',*ROUTE)
    def test_same_total_different_rates_requires_conflict_choice(self):
        out=io.BytesIO()
        with zipfile.ZipFile(out,'w') as z:
            for rate in [10,12]:z.writestr(f'{rate}.csv',interval_csv([(50,100,rate,'руб/кг',1500)]))
        conflicts=[];routes,errors=price_library.parse_routes(out.getvalue(),'parts.zip','Werner',conflicts=conflicts)
        self.assertFalse(errors);self.assertEqual(len(conflicts),1)
        self.assertEqual({r['price'] for r in conflicts[0]['options']},{1500})
        self.assertEqual({r['rate_per_kg'] for r in conflicts[0]['options']},{10,12})
    def test_file_rate_fills_calculator_total_but_not_published_online_rate(self):
        vals=self.parse([(50,100,10,'руб/кг',1500)])
        aid=e.begin_live_attempt('Werner',*ROUTE)
        e.save_live_update('Werner',*ROUTE,{'w100':{'kind':'exact','price':1700}},{},aid)
        e.finish_live_attempt('Werner',*ROUTE,aid,rows=1)
        items=bulk.capture_company('Werner',*ROUTE,{'ok':True})['items']
        rows=[{'profile':p,'items':[items[i]]} for i,p in enumerate(e.COMMON_PROFILES)]
        filled=bulk.fill_document_gaps(rows,self.pack(vals),*ROUTE)
        q=next(r['items'][0] for r in filled if r['profile']['id']=='w100')
        self.assertTrue(q['uploaded']);self.assertEqual(q['published_rate_per_kg'],10)
        for r in rows:
            if r['profile']['id']=='w100':r['items'][0]['published_rate_per_kg']=11
        filled=bulk.fill_document_gaps(rows,self.pack(vals),*ROUTE)
        q=next(r['items'][0] for r in filled if r['profile']['id']=='w100')
        self.assertTrue(q['collected_online']);self.assertEqual(q['published_rate_per_kg'],11)
    def test_main_scope_matches_127_destinations_each_and_keeps_extended_routes(self):
        routes=bulk.plan_routes();self.assertEqual(len(routes),254)
        self.assertEqual(len(main_cities()),128)
        for o in ROUTE:
            dest={d for a,d in routes if a==o};self.assertEqual(len(dest),127);self.assertEqual(dest,set(main_cities())-{o})
            options=main.options(o);self.assertEqual(options['origins'],list(ROUTE));self.assertEqual(set(options['destinations']),dest)
        self.assertEqual(len(bulk.plan_routes('extended_reference')),318)
        self.assertEqual(len(main.options('Москва',catalog='all')['origins']),233)
        self.assertEqual(len(e.COMPANIES),17);self.assertNotIn('Грузопоток',e.COMPANIES)
    def test_interval_pdf_uses_explicit_units(self):
        text='Компания: Werner\nОткуда: Москва\nКуда: Санкт-Петербург\nВес от, кг | Вес до, кг | Тариф | Единица | Минимум, руб\n0 | 5 | 196 | руб |\n5 | 20 | 403 | руб |\n50 | 100 | 10 | руб/кг | 1500'
        reader=SimpleNamespace(pages=[SimpleNamespace(extract_text=lambda **kw:text)])
        with patch.object(docs,'pdf_reader',return_value=reader):values,meta=docs.parse_generic_pdf(b'test','Werner',*ROUTE)
        self.assertEqual(values['w003']['price'],196);self.assertEqual(values['w100']['rate_per_kg'],10)
    def test_explicit_minimum_is_kept_when_small_weights_are_missing(self):
        values=self.parse([(100,200,10,'руб/кг',600)])
        q=e._route_quote('Werner',*ROUTE,'min',({}, {},self.pack(values)))
        self.assertEqual(q['price'],600);self.assertEqual(values['w200']['price'],2000)
    def test_new_interval_template_has_empty_prices_and_all_main_routes(self):
        wb=load_workbook(io.BytesIO(price_library.template('Werner')),data_only=True)
        rows=list(wb['Диапазоны'].values)
        self.assertEqual(list(rows[0]),docs.INTERVAL_HEADER)
        self.assertEqual(len(rows)-1,254*28)
        self.assertTrue(all(r[5] is None and r[7] is None for r in rows[1:]))
        self.assertEqual({r[6] for r in rows[1:]},{'руб','руб/кг'});wb.close()
    def test_upgrade_keeps_checkpoints_and_invalidates_old_excel(self):
        m=bulk.BulkManager(self.root/'bulk')
        with patch.object(m,'_launch'):job=m.start('origins',['Казань'],['Уфа'])
        with m.db() as db:
            config=json.loads(db.execute('SELECT config FROM jobs').fetchone()[0]);config.pop('tariff_schema');config['companies'].append('Грузопоток')
            db.execute("UPDATE jobs SET config=?,status='done',export_status='ready',export_file='old.xlsx'",(json.dumps(config),))
        restored=bulk.BulkManager(self.root/'bulk');status=restored.status(job['job_id'])
        self.assertEqual(status['companies'],e.COMPANIES);self.assertEqual(status['total_routes'],1)
        self.assertIsNone(status['download_url']);self.assertEqual(status['export_status'],'idle')


if __name__=='__main__':unittest.main()
