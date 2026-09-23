import io,json,os,tempfile,subprocess,sys,unittest
from pathlib import Path
from datetime import datetime,timezone,timedelta,date
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e,v42_main as main,manual_prices as manual,document_imports as imp,price_library as lib,bulk_refresh as bulk,business_time as clock,official_documents as official

ROOT=Path(__file__).resolve().parent.parent
O,D='Москва','Санкт-Петербург'

class Cells(unittest.TestCase):
 def setUp(self):
  self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
  self.ps=[patch.object(e,'RUNTIME_DIR',self.root),patch.object(e,'ROUTE_CONFIG',{})]
  for p in self.ps:p.start()
  self.client=TestClient(main.app)
 def tearDown(self):
  for p in reversed(self.ps):p.stop()
  self.tmp.cleanup()
 def data(self,**values):return {'company':'ДЛ','origin':O,'destination':D,'profile':'w100','value':'14,5','unit':'rub_per_kg',**values}
 def test_edit_exact_cell_units_both_exports_restart_and_delete(self):
  response=self.client.post('/api/manual-price',json=self.data());self.assertEqual(response.status_code,200,response.text)
  q=e.quote('ДЛ',O,D,'w100');self.assertEqual(q['price'],1450);self.assertEqual(q['published_rate_per_kg'],14.5)
  self.assertTrue(q['manual']);self.assertFalse(q['online']);self.assertFalse(q['uploaded'])
  self.assertIsNone(e.quote('ДЛ',D,O,'w100')['price']);self.assertIsNone(e.quote('ПЭК',O,D,'w100')['price'])
  self.assertIsNone(e.quote('ДЛ',O,D,'w200')['price']);self.assertIsNone(e.quote('ДЛ',O,D,'min')['price'])
  response=self.client.post('/api/manual-price',json=self.data(profile='w001',value='600',unit='rub'));self.assertEqual(response.status_code,200)
  self.assertEqual(e.quote('ДЛ',O,D,'min')['price'],600)
  code="from app import v42_engine as e;import json;e.ROUTE_CONFIG={};print(json.dumps(e.quote('ДЛ','Москва','Санкт-Петербург','w100')))"
  child=subprocess.run([sys.executable,'-c',code],cwd=ROOT,env={**os.environ,'TARIFF_DATA_DIR':self.tmp.name},text=True,capture_output=True,check=True)
  self.assertEqual(json.loads(child.stdout)['price'],1450)
  response=self.client.get('/api/export/route',params={'origin':O,'destination':D,'companies':'ДЛ'})
  wb=load_workbook(io.BytesIO(response.content),data_only=True);self.assertEqual(wb.sheetnames,['Маршрут','Источники'])
  self.assertTrue(any(1450 in r for r in wb['Маршрут'].values));self.assertTrue(any('Введено вручную' in r for r in wb['Источники'].values));wb.close()
  from app.excel_export import export_bytes
  wb=load_workbook(io.BytesIO(export_bytes(O,D,['ДЛ'],routes_override=[(O,D)])),data_only=True)
  self.assertEqual(wb['ДЛ']['D2'].value,600);self.assertEqual(wb['ДЛ']['N2'].value,14.5);wb.close()
  self.assertEqual(self.client.request('DELETE','/api/manual-price',json=self.data()).status_code,200)
  self.assertIsNone(e.quote('ДЛ',O,D,'w100')['price'])
 def test_invalid_input_never_writes_a_price(self):
  for value in ['=1+2','NaN','Infinity',-1,0,True,'abc','1e99']:
   r=self.client.post('/api/manual-price',json=self.data(value=value));self.assertIn(r.status_code,[400,422],(value,r.text))
  self.assertEqual(self.client.post('/api/manual-price',json=self.data(profile='w005')).status_code,400)
  self.assertEqual(self.client.post('/api/manual-price',json=self.data(company='unknown')).status_code,400)
  self.assertIsNone(e.quote('ДЛ',O,D,'w100')['price'])
 def test_file_older_than_month_is_immediately_available_and_application_receipt(self):
  csv='Компания;Откуда;Куда;Вес, кг;Цена, руб\nДЛ;Москва;Санкт-Петербург;1;710\n'.encode()
  p=imp.preview(csv,'monthly.csv','ДЛ',O,D);saved=imp.commit(p['token'])
  file_id=p['token'];old=(datetime.now().astimezone()-timedelta(days=65)).isoformat()
  with lib.db() as db:
   row=db.execute('SELECT meta FROM files WHERE id=?',(file_id,)).fetchone();m=json.loads(row['meta']);m.update(uploaded_at=old,captured_at=old)
   db.execute('UPDATE files SET meta=? WHERE id=?',(json.dumps(m),file_id))
  e._cached_json.cache_clear();q=e.quote('ДЛ',O,D,'w001');self.assertEqual(q['price'],710);self.assertTrue(q['uploaded']);self.assertFalse(q['online'])
  manual.put('ДЛ',O,D,'w001',800,'rub');manual.put('ДЛ',O,D,'w200',20,'rub_per_kg')
  response=self.client.post('/api/route-documents/select',json={'company':'ДЛ','origin':O,'destination':D,'document_id':file_id})
  receipt=response.json();self.assertEqual(response.status_code,200,response.text)
  self.assertEqual(receipt['filename'],'monthly.csv');self.assertEqual(receipt['applied_prices'],1);self.assertEqual(receipt['replaced_manual'],1)
  self.assertEqual(e.quote('ДЛ',O,D,'w001')['price'],710);self.assertFalse(e.quote('ДЛ',O,D,'w001').get('manual'))
  self.assertEqual(e.quote('ДЛ',O,D,'w200')['price'],4000)
  bad=self.client.post('/api/route-documents/select',json={'company':'ПЭК','origin':O,'destination':D,'document_id':file_id});self.assertEqual(bad.status_code,400)
 def test_failed_update_never_removes_manual_or_saved_monthly_data(self):
  aid=e.begin_live_attempt('ДЛ',O,D);old=(datetime.now().astimezone()-timedelta(days=45)).isoformat()
  e.save_live_update('ДЛ',O,D,{'w100':{'kind':'exact','price':1440,'rate_per_kg':14.4}},{'captured_at':old},aid);e.finish_live_attempt('ДЛ',O,D,aid,rows=1)
  aid=e.begin_live_attempt('ДЛ',O,D);e.finish_live_attempt('ДЛ',O,D,aid,rows=0,error='HTTP 502')
  self.assertEqual(e.quote('ДЛ',O,D,'w100')['price'],1440)
  manual.put('ДЛ',O,D,'w100',16,'rub_per_kg');self.assertEqual(e.quote('ДЛ',O,D,'w100')['price'],1600)
  manual.remove('ДЛ',O,D,'w100');self.assertEqual(e.quote('ДЛ',O,D,'w100')['price'],1440)

class MoscowDate(unittest.TestCase):
 def test_moscow_midnight_independent_of_host_timezone(self):
  class Fake:
   @classmethod
   def now(cls,tz):return datetime(2026,9,23,21,30,tzinfo=timezone.utc).astimezone(tz)
  with patch.object(clock,'datetime',Fake):self.assertEqual(clock.tariff_today(),date(2026,9,24))
  with patch.object(official,'tariff_today',return_value=date(2026,9,24)):
   official._check_current_document({'document_date':'2026-09-24'})
   with self.assertRaisesRegex(ValueError,'будущие'):official._check_current_document({'document_date':'2026-09-25'})

if __name__=='__main__':unittest.main()
