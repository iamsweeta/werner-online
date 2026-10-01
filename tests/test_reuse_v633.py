"""One uploaded original, several routes, shared exports and durable restart."""
import io,json,os,subprocess,sys,tempfile,time,unittest
from pathlib import Path
from unittest.mock import patch
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from app import v42_engine as e,v42_main as main,price_library as lib,document_imports as imp,route_import_jobs as jobs
from tests.test_library_v53 import document,ROUTE,ROOT

class ReuseDocument(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.folder=Path(self.tmp.name)
        self.patches=[patch.object(e,'RUNTIME_DIR',self.folder),patch.object(e,'ROUTE_CONFIG',{})]
        for p in self.patches:p.start()
        self.client=TestClient(main.app)
        p=imp.preview(document(),'multi-route.csv','Werner',*ROUTE)
        self.saved=imp.commit(p['token']);self.ident=self.saved['document_id']
    def tearDown(self):
        for p in reversed(self.patches):p.stop()
        self.tmp.cleanup()
    def extract(self,route=ROUTE[::-1],key='d'*32):
        result=self.client.post('/api/price-documents/'+self.ident+'/extract',json={'origin':route[0],'destination':route[1],'job_id':key})
        self.assertEqual(result.status_code,202,result.text)
        deadline=time.monotonic()+10
        while time.monotonic()<deadline:
            data=self.client.get('/api/import/jobs/'+key).json()
            if data['status'] not in {'parsing','queued'}:return data
            time.sleep(.01)
        self.fail('Parsing did not finish')
    def test_saved_file_adds_reverse_without_upload_then_both_exports_and_restart(self):
        self.assertIsNone(e.quote('Werner',*ROUTE[::-1],'w100')['price'])
        original=(imp.root()/'files'/self.saved['meta']['source_file']).read_bytes()
        job=self.extract();self.assertEqual(job['status'],'ready',job)
        self.assertIsNone(e.quote('Werner',*ROUTE[::-1],'w100')['price'],'preview must not apply prices')
        self.assertEqual(job['document_id'],self.ident)
        self.assertFalse(list(jobs.root().glob('*.upload')),'reuse references the original')
        token=job['preview']['token']
        result=self.client.post('/api/import/commit',json={'token':token});self.assertEqual(result.status_code,200,result.text)
        self.assertTrue(result.json()['reused_document']);self.assertEqual(result.json()['document_id'],self.ident)
        self.assertEqual(jobs.status('d'*32)['status'],'committed')
        self.assertEqual(len(lib.list_files()),1);self.assertEqual(lib.list_files()[0]['route_count'],2)
        self.assertEqual((imp.root()/'files'/self.saved['meta']['source_file']).read_bytes(),original)
        self.assertEqual(len(list((imp.root()/'files').iterdir())),1)
        self.assertEqual(e.quote('Werner',*ROUTE,'w100')['price'],1200)
        reverse=e.quote('Werner',*ROUTE[::-1],'w100');self.assertEqual(reverse['price'],1300)
        stored=lib.pack(*ROUTE[::-1])['profiles']['w100']['Werner']
        self.assertEqual(stored['origin'],ROUTE[1]);self.assertEqual(stored['destination'],ROUTE[0])
        self.assertEqual(reverse['document_id'],self.ident);self.assertTrue(reverse['document_selected'])
        count=lib.list_files()[0]['values_count'];again=self.client.post('/api/import/commit',json={'token':token}).json()
        self.assertTrue(again['already_applied']);self.assertEqual(lib.list_files()[0]['values_count'],count)
        for endpoint,sheet in [('/api/export/route','Маршрут'),('/api/export/excel','WernerNEW')]:
            r=self.client.get(endpoint,params={'origin':ROUTE[1],'destination':ROUTE[0],'companies':'Werner'})
            self.assertEqual(r.status_code,200)
            wb=load_workbook(io.BytesIO(r.content),data_only=True)
            expected=1300 if sheet=='Маршрут' else 13
            self.assertTrue(any(expected in row for row in wb[sheet].values));wb.close()
        code="import json;from app import v42_engine as e,price_library as l;print(json.dumps([e.quote('Werner','Уфа','Казань','w100')['price'],l.list_files()[0]['route_count']]))"
        child=subprocess.run([sys.executable,'-c',code],cwd=ROOT,env={**os.environ,'TARIFF_DATA_DIR':str(self.folder)},capture_output=True,text=True,check=True)
        self.assertEqual(json.loads(child.stdout),[1300,2])
    def test_missing_route_reports_error_without_inventing_or_erasing_prices(self):
        data=self.extract(('Москва','Уфа'))
        self.assertEqual(data['status'],'error');self.assertTrue(data['can_retry'])
        self.assertEqual(lib.list_files()[0]['route_count'],1)
        self.assertEqual(e.quote('Werner',*ROUTE,'w100')['price'],1200)
        self.assertIsNone(e.quote('Werner','Москва','Уфа','w100')['price'])
        self.assertTrue((imp.root()/'files'/self.saved['meta']['source_file']).exists())
    def test_native_dellin_pdf_reuses_original_for_another_destination(self):
        raw=(ROOT/'tests/fixtures/dellin_original_example.pdf').read_bytes()
        first=imp.preview(raw,"pricelist.pdf'",'ДЛ','Москва','Санкт-Петербург')
        saved=imp.commit(first['token']);self.ident=saved['document_id']
        job=self.extract(('Москва','Петропавловск-Камчатский'))
        self.assertEqual(job['status'],'ready',job)
        result=imp.commit(job['preview']['token'])
        self.assertEqual(result['document_id'],self.ident)
        self.assertEqual(result['meta']['original_filename'],'pricelist.pdf')
        self.assertEqual(e.quote('ДЛ','Москва','Санкт-Петербург','w100')['price'],1480)
        self.assertEqual(e.quote('ДЛ','Москва','Петропавловск-Камчатский','w100')['price'],6010)
        self.assertEqual(lib.document_routes(self.ident)['meta']['route_count'],2)
        self.assertEqual(imp.source_file(saved['meta']['source_file']).read_bytes(),raw)
    def test_disabled_file_cannot_be_applied_after_preview(self):
        job=self.extract();self.assertEqual(job['status'],'ready')
        lib.remove(self.ident)
        response=self.client.post('/api/import/commit',json={'token':job['preview']['token']})
        self.assertEqual(response.status_code,400)
        self.assertIsNone(e.quote('Werner',*ROUTE[::-1],'w100')['price'])
    def test_corrupt_original_and_invalid_ids_do_not_start_job(self):
        for ident in ['not-an-id','a'*32]:
            r=self.client.post('/api/price-documents/'+ident+'/extract',json={'origin':'Уфа','destination':'Казань'})
            self.assertEqual(r.status_code,400)
        path=imp.root()/'files'/self.saved['meta']['source_file'];path.write_bytes(b'changed')
        r=self.client.post('/api/price-documents/'+self.ident+'/extract',json={'origin':'Уфа','destination':'Казань'})
        self.assertEqual(r.status_code,400);self.assertIn('целостности',r.json()['detail'])

if __name__=='__main__':unittest.main()
