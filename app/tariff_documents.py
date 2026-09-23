"""Read user documents and current carrier exports without consulting price caches."""
from __future__ import annotations

import io
import math
import re
import zipfile
import csv
from contextlib import contextmanager
from contextvars import ContextVar
from pathlib import PurePosixPath

from pypdf import PdfReader
from openpyxl import load_workbook

from .cities import city_pattern, normalize_city
from .v42_engine import COMMON_PROFILES, COMPANIES
from .online_tariffs import number, norm, bounds, values_from_tiers
from .interval_tariffs import INTERVAL_HEADER, is_interval_header, parse_interval_workbook, parse_werner, price_signature

MAX_BYTES = 20 * 1024 * 1024
MAX_EXPANDED = 120 * 1024 * 1024
MAX_PAGES = 80
_SESSION=ContextVar('tariff_document_session',default=None)


@contextmanager
def parsing_session():
    if _SESSION.get() is not None:
        yield
        return
    token=_SESSION.set({})
    try:yield
    finally:_SESSION.reset(token)


def pdf_has_text_fonts(page):
    """Conservative resource check: outlined/image PDFs need no text extraction."""
    def inspect(resources,seen,depth):
        if depth>12:return True
        try:
            if hasattr(resources,'get_object'):resources=resources.get_object()
            if id(resources) in seen:return True
            seen=seen|{id(resources)}
            if resources.get('/Font'):return True
            objects=resources.get('/XObject') or {}
            if hasattr(objects,'get_object'):objects=objects.get_object()
            for obj in objects.values():
                obj=obj.get_object();kind=obj.get('/Subtype')
                if kind=='/Image':continue
                if kind!='/Form':return True
                if inspect(obj.get('/Resources',{}),seen,depth+1):return True
            return False
        except AttributeError:
            # Ordinary empty resource dictionaries are safe; unknown objects are not.
            return bool(resources)
        except Exception:return True
    return inspect(page.get('/Resources',{}),set(),0)


def pdf_reader(raw):
    from pypdf import PdfReader
    cache=_SESSION.get();key=('pdf',raw)
    if cache is not None and key in cache:return cache[key]
    reader=PdfReader(io.BytesIO(raw))
    if cache is not None:
        for page in reader.pages:
            original=page.extract_text if pdf_has_text_fonts(page) else lambda *a,**k:'';texts={}
            def extract(*args,_fn=original,_texts=texts,**kwargs):
                k=repr((args,kwargs))
                if k not in _texts:_texts[k]=_fn(*args,**kwargs)
                return _texts[k]
            page.extract_text=extract
        cache[key]=reader
    return reader


def checked_zip(raw):
    z = zipfile.ZipFile(io.BytesIO(raw))
    infos = z.infolist()
    if len(infos) > 2000 or sum(x.file_size for x in infos) > MAX_EXPANDED:
        z.close()
        raise ValueError('Архив слишком велик после распаковки (лимит 120 МБ / 2000 файлов)')
    if len({x.filename for x in infos})!=len(infos):
        z.close();raise ValueError('В архиве повторяются имена файлов. Переименуйте документы и создайте ZIP заново.')
    for x in infos:
        p = PurePosixPath(x.filename.replace('\\', '/'))
        if p.is_absolute() or '..' in p.parts or x.flag_bits & 1 or x.file_size > MAX_EXPANDED:
            z.close()
            raise ValueError('Архив содержит небезопасные пути или зашифрованные файлы')
    return z


def normalize_filename(filename):
    """Repair browser/server quote artifacts without trusting the extension alone."""
    name = str(filename or '').replace('\\', '/').rsplit('/', 1)[-1]
    name = name.strip(" \t\r\n\"'‘’“”`")
    name = ''.join(c for c in name if ord(c) >= 32)
    ext = PurePosixPath(name).suffix
    return name[:max(0, 180-len(ext))] + ext if len(name) > 180 else name


def validate_file(raw, filename):
    if not raw or len(raw) > MAX_BYTES:
        raise ValueError('Загрузите непустой документ размером не более 20 МБ')
    ext = PurePosixPath(normalize_filename(filename)).suffix.lower()
    if ext not in {'.pdf', '.xls', '.xlsx', '.zip', '.csv'}:
        raise ValueError('Поддерживаются PDF (включая сканы ДЛ), XLS, XLSX, CSV и ZIP с прайс-листами')
    cache=_SESSION.get();key=('valid',raw,ext)
    if cache is not None and key in cache:return ext
    if ext == '.pdf':
        if not raw.startswith(b'%PDF-'):
            raise ValueError('Содержимое файла не является PDF')
        reader = pdf_reader(raw)
        if reader.is_encrypted:
            raise ValueError('Сначала сохраните PDF без пароля')
        if len(reader.pages) > MAX_PAGES:
            raise ValueError(f'В PDF больше {MAX_PAGES} страниц: загрузите раздел с тарифами')
    elif ext in {'.xlsx', '.zip'}:
        with checked_zip(raw) as z:
            if ext == '.xlsx' and '[Content_Types].xml' not in z.namelist():
                raise ValueError('Содержимое файла не является XLSX')
    elif ext=='.csv':
        if b'\0' in raw:raise ValueError('CSV должен содержать текстовую таблицу')
    elif not raw.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        raise ValueError('Содержимое файла не является XLS')
    if cache is not None:cache[key]=True
    return ext


def dated(text):
    hit = re.search(r'(?:действительны на|действуют с|действует с|тарифы с|рублях[^\n]{0,30}\bна)\s*(\d{2})\.(\d{2})\.(\d{4})', text, re.I)
    if not hit:
        return None
    from datetime import date
    return date(int(hit[3]), int(hit[2]), int(hit[1])).isoformat()


def parse_dellin(raw, origin, destination):
    reader = pdf_reader(raw)
    first = reader.pages[0].extract_text() or ''
    header = norm(first.split('Направление')[0])
    if not re.search(r'тарифы на межтерминальную перевозку из\s+' + city_pattern(origin) + r'(?![а-яё])', header, re.I):
        raise ValueError(f'ДЛ: заголовок PDF не подтверждает отправление из {origin}')
    # This parser intentionally accepts only the published 5 + 1 + 14 + 14 layout.
    compact = re.sub(r'\s+', '', first[:1500]).replace('–', '-').replace('—', '-')
    if not all(h in compact for h in ('5000-10000', '3000-4999', '100-149', '36-99', 'Стоимость,руб/кг')):
        raise ValueError('ДЛ: изменился заголовок весовой таблицы; автоматическое сопоставление остановлено')
    matches = []
    texts = []
    for page_no, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ''
        texts.append(text)
        end = re.search(r'Тарифы на доставку|Тарифы на услуги|Особые требования', text, re.I)
        section = text[:end.start()] if end else text
        lines = [' '.join(x.split()) for x in section.splitlines() if x.strip()]
        for i, line in enumerate(lines):
            row_start=None
            for count in (1,2,3):
                label=re.sub(r'-\s+', '-', ' '.join(lines[i:i+count]))
                if normalize_city(label)==destination:
                    row_start=i+count
                    break
            if row_start is None:
                continue
            cells = lines[row_start:row_start + 34]
            if len(cells) != 34:
                raise ValueError('ДЛ: неполная строка назначения в PDF')
            try:
                nums = [number(x) for x in cells]
            except ValueError as exc:
                raise ValueError('ДЛ: число колонок или структура строки изменились') from exc
            matches.append((nums, page_no))
        if end:
            break
    if len(matches) != 1:
        raise ValueError(f'ДЛ: не найдена единственная межтерминальная строка {origin} → {destination}')
    nums, page = matches[0]
    kg_bounds = [(5000,10000),(3000,4999),(2500,2999),(2000,2499),(1500,1999),
                 (1200,1499),(1000,1199),(800,999),(500,799),(300,499),
                 (200,299),(150,199),(100,149),(36,99)]
    fixed_bounds = [(0,1),(2,3),(4,5),(6,15),(16,35)]
    vals = values_from_tiers([(lo,hi,rate) for (lo,hi),rate in zip(kg_bounds,nums[6:20])],
                            fixed=[(lo,hi,rate) for (lo,hi),rate in zip(fixed_bounds,nums[:5])], minimum=nums[5])
    return vals, {'parser':'ДЛ · первая межтерминальная таблица', 'source_page':page,
                  'document_date':dated('\n'.join(texts)), 'route_verified':True,
                  'calculation_basis':'ДЛ: фиксированный тариф до 35 кг (объём до 0,1 м³); далее max(минимум, вес × ₽/кг). До 10 000 кг; без адресной доставки и допуслуг.'}


def workbook_rows(raw):
    cache=_SESSION.get();key=('workbook',raw)
    if cache is None:yield from _workbook_rows(raw);return
    if key not in cache:cache[key]=list(_workbook_rows(raw))
    yield from cache[key]


def workbook_view(raw):
    from types import SimpleNamespace
    class View:
        def __init__(self):
            self.sheets={title:SimpleNamespace(values=rows,title=title) for title,rows in workbook_rows(raw)}
            self.sheetnames=list(self.sheets)
        def __getitem__(self,key):return self.sheets[key]
        def __iter__(self):return iter(self.sheets.values())
        def close(self):pass
    return View()


def _workbook_rows(raw):
    """Bound workbook work; never evaluate formulas or execute macros."""
    if raw.startswith(b'PK'):
        with checked_zip(raw):
            pass
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            if len(wb.worksheets) > 300:
                raise ValueError('В книге слишком много листов')
            for sh in wb.worksheets:
                if (sh.max_row or 0) > 50000 or (sh.max_column or 0) > 512:
                    raise ValueError('Таблица превышает лимит 50 000 строк / 512 оформленных столбцов')
                rows=[]
                for row in sh.iter_rows(values_only=True):
                    # Official files often format empty columns far beyond the
                    # table. Bound populated columns, not decorative cells.
                    end=len(row)
                    while end and row[end-1] is None:end-=1
                    if end>150:raise ValueError('Таблица содержит более 150 заполненных столбцов')
                    rows.append(tuple(row[:end]))
                width=max((len(row) for row in rows),default=0)
                yield sh.title, [row+(None,)*(width-len(row)) for row in rows]
        finally:
            wb.close()
    elif raw.startswith(bytes.fromhex('d0cf11e0a1b11ae1')):
        import xlrd
        wb = xlrd.open_workbook(file_contents=raw, on_demand=True)
        try:
            if wb.nsheets > 300:
                raise ValueError('В книге слишком много листов')
            for i in range(wb.nsheets):
                sh = wb.sheet_by_index(i)
                if sh.nrows > 50000 or sh.ncols > 150:
                    raise ValueError('Таблица превышает лимит 50 000 строк / 150 столбцов')
                yield sh.name, [sh.row_values(j) for j in range(sh.nrows)]
                wb.unload_sheet(i)
        finally:
            wb.release_resources()
    else:
        try:text=raw.decode('utf-8-sig')
        except UnicodeDecodeError:text=raw.decode('cp1251')
        try:dialect=csv.Sniffer().sniff(text[:8192],delimiters=';,\t')
        except csv.Error:raise ValueError('Не удалось определить разделитель CSV: используйте точку с запятой')
        rows=[]
        for row in csv.reader(io.StringIO(text),dialect):
            if len(rows)>=50000 or len(row)>150:raise ValueError('CSV превышает лимит 50 000 строк / 150 столбцов')
            rows.append(row)
        yield 'CSV',rows


def parse_vozovoz(raw, origin, destination):
    sheets = [rows for title,rows in workbook_rows(raw) if norm(title) == 'тарифы межтерм.перевозок']
    if len(sheets) != 1:
        raise ValueError('Возовоз: нужен лист «Тарифы межтерм.перевозок» обычной перевозки')
    rows = sheets[0]
    text = '\n'.join(' '.join(str(v or '') for v in r) for r in rows)
    if 'vozovoz.ru' not in text or 'российских рублях' not in text:
        raise ValueError('Возовоз: не подтверждены источник и валюта RUB')
    headers = [r for r in rows[:20] if len(r) >= 25 and norm(r[0]) == 'отправление' and norm(r[1]) == 'назначение']
    kg = [r for r in rows[:20] if len(r) >= 25 and norm(r[7]) == 'кг' and '99' in str(r[8])]
    fixed = [r for r in rows[:20] if len(r) >= 25 and 'кг' in str(r[2]) and 'ДШВ' in str(r[2])]
    if len(headers) != 1 or len(kg) != 1 or len(fixed) != 1 or 'мин.' not in norm(headers[0][24]):
        raise ValueError('Возовоз: структура столбцов изменилась')
    matches = [(i+1,r) for i,r in enumerate(rows) if len(r)>=25 and
               normalize_city(str(r[0]).strip(' *')) == origin and normalize_city(str(r[1]).strip(' *')) == destination and norm(r[7]) == 'кг']
    if len(matches) != 1:
        raise ValueError(f'Возовоз: нет единственной исходящей строки {origin} → {destination}')
    row_no, row = matches[0]
    tiers = [(*bounds(kg[0][j]), number(row[j])) for j in range(8,24) if row[j] not in ('',None,'-','—')]
    fix = [(*bounds(str(fixed[0][j]).split('кг')[0]), number(row[j])) for j in range(2,7) if row[j] not in ('',None,'-','—')]
    minimum = number(row[24])
    vals = values_from_tiers(tiers, fixed=fix, minimum=minimum)
    rounding = bool(re.search(r'округляется до 10 руб.*большую сторону', norm(text)))
    if rounding:
        for v in vals.values():
            v['price'] = math.ceil(round(v['price'],2)/10 - 1e-9) * 10.0
    return vals, {'parser':'Возовоз · Тарифы межтерм.перевозок', 'source_row':row_no,
                  'document_date':dated(text), 'route_verified':True,
                  'calculation_basis':'Возовоз: обычная межтерминальная перевозка; фиксированные тарифы до 40 кг с ограничениями ДШВ из файла, далее max(минимум, вес × ₽/кг). '
                  + ('Округление вверх до 10 ₽ по правилу файла. ' if rounding else '') + 'Без скидок, доставки до адреса, страхования и допуслуг; объём не включён.'}


def parse_vozovoz_zip(raw, origin, destination):
    found = []
    errors = []
    with checked_zip(raw) as z:
        docs = [x for x in z.infolist() if PurePosixPath(x.filename).suffix.lower() in {'.xls','.xlsx'}]
        if not docs or len(docs)>30:
            raise ValueError('В ZIP ожидается до 30 файлов XLS/XLSX Возовоза')
        for info in docs:
            if info.file_size > MAX_BYTES:
                raise ValueError('Excel внутри ZIP превышает 20 МБ')
            try:
                result,meta = parse_vozovoz(z.read(info),origin,destination)
                found.append((result,{**meta,'archive_member':info.filename}))
            except ValueError as exc:
                errors.append(str(exc))
    if len(found)!=1:
        raise ValueError('В ZIP не найдена единственная таблица выбранного направления. '+('; '.join(errors[:2]) if not found else 'Найдены дубликаты маршрута.'))
    return found[0]


GENERIC_HEADER = ['Компания','Откуда','Куда','Вес, кг','Цена, руб']


def is_template_header(row):
    return is_interval_header(row) or [norm(x) for x in row[:5]] == [norm(x) for x in GENERIC_HEADER]


def _profile(value):
    s = norm(value)
    if s in {'мин','мин.','минимум'}:
        return 'min'
    # Exact control weights only: no interpolation, no inferred minimum.
    if re.fullmatch(r'\d+(?:[.,]\d+)?',s):
        weight = float(s.replace(',','.'))
        return next((p['id'] for p in COMMON_PROFILES if not p.get('is_minimum_profile') and p['weight_kg']==weight),None)
    return None


def parse_generic_workbook(raw, company, origin, destination):
    vals = {}; matched = 0
    for title, rows in workbook_rows(raw):
        header_idx = next((i for i,r in enumerate(rows[:20]) if [norm(x) for x in r[:5]] == [norm(x) for x in GENERIC_HEADER]),None)
        if header_idx is None:
            continue
        for i,r in enumerate(rows[header_idx+1:],header_idx+2):
            if len(r)<5 or not any(x not in ('',None) for x in r[:5]):continue
            if str(r[0]).strip()!=company or normalize_city(str(r[1]))!=origin or normalize_city(str(r[2]))!=destination:continue
            if r[4] in ('',None):continue
            pid=_profile(r[3])
            if not pid:raise ValueError(f'Лист {title}, строка {i}: укажите контрольный вес из приложения или МИН')
            if pid in vals:raise ValueError('В таблице повторяется один вес выбранного маршрута')
            vals[pid]={'kind':'exact','price':number(r[4]),'source_row':i}
            matched+=1
    if not matched:
        raise ValueError('Формат таблицы не распознан или нет выбранной компании и маршрута. Используйте XLSX-шаблон: Компания, Откуда, Куда, Вес, кг, Цена, руб. Для произвольных таблиц автоматическое угадывание колонок отключено.')
    return vals,{'parser':'Таблица по шаблону · стоимость отправки', 'route_verified':True,
                 'calculation_basis':'Цена за отправку для контрольного веса из пользовательской таблицы; без интерполяции.'}


def parse_pek(raw,origin,destination):
    from .legacy_backend import _strict_weight_range
    sheets=[r for title,r in workbook_rows(raw) if title=='Перевозка']
    if len(sheets)!=1 or len(sheets[0])<22:raise ValueError('ПЭК: не найден лист «Перевозка» с заголовками')
    rows=sheets[0];group,header=rows[19],rows[20]
    if norm(header[1])!='город отправитель' or norm(header[2])!='город получатель' or 'руб' not in norm(group[8]):
        raise ValueError('ПЭК: заголовки валюты и направления не подтверждены')
    matches=[(i+1,r) for i,r in enumerate(rows) if len(r)>=21 and normalize_city(str(r[1]))==origin and normalize_city(str(r[2]))==destination]
    if len(matches)!=1:raise ValueError('ПЭК: нет единственной строки выбранного направления')
    row_no,row=matches[0];tiers=[];fixed=[];active=''
    minimum=number(row[8]) if row[8] not in ('',None) else None
    for i in range(3,21):
        if group[i]:active=norm(group[i])
        if i==8 or row[i] in ('',None,'-','—'):continue
        b=_strict_weight_range(str(header[i] or ''))
        if not b:raise ValueError('ПЭК: не распознан весовой диапазон')
        (tiers if '1 кг' in active else fixed).append((b[0] or 0,b[1] if b[1] is not None else math.inf,number(row[i])))
    return values_from_tiers(tiers,fixed=fixed,minimum=minimum),{'parser':'ПЭК · Перевозка XLSX','source_row':row_no,'route_verified':True,
            'calculation_basis':'ПЭК: фиксированный тариф малого груза с ограничением объёма из файла; далее max(минимум, вес × ₽/кг). Допуслуги не включены.'}


def parse_generic_pdf(raw, company, origin, destination):
    """Strict two-column text table, with explicit carrier and ordered route."""
    reader=pdf_reader(raw);vals={};page_numbers=[]
    all_text='\n'.join(p.extract_text() or '' for p in reader.pages)
    if not all_text.strip():
        raise ValueError('В PDF нет текстового слоя. Для скана ДЛ выберите компанию ДЛ; для другого макета используйте текстовый PDF или XLSX.')
    fields={}
    for key in ('Компания','Откуда','Куда'):
        matches=re.findall(r'^\s*'+key+r'\s*:\s*(.*?)\s*$',all_text,re.M|re.I)
        if len(set(matches))!=1:raise ValueError('Для общего PDF нужны строки «Компания: …», «Откуда: …», «Куда: …» и таблица «Вес, кг | Цена, руб». Используйте XLSX-шаблон для другого макета.')
        fields[key]=matches[0]
    if fields['Компания']!=company or normalize_city(fields['Откуда'])!=origin or normalize_city(fields['Куда'])!=destination:
        raise ValueError('Компания или направление в PDF не совпадают с выбранными')
    intervals=[]
    for page_no,page in enumerate(reader.pages,1):
        active=False
        for line in (page.extract_text(extraction_mode='layout') or '').splitlines():
            cells=re.split(r'\s*[|;]\s*|\s{2,}',line.strip())
            if [norm(c) for c in cells]==[norm(c) for c in INTERVAL_HEADER[3:]]:active=True;continue
            if not active or not line.strip():continue
            if len(cells) not in (4,5):active=False;continue
            from .interval_tariffs import interval_row
            intervals.append(interval_row(cells,page_no=page_no))
    if intervals:
        from .interval_tariffs import project_intervals
        return project_intervals(intervals),{'parser':'PDF · интервалы с явными единицами','route_verified':True,
            'document_date':dated(all_text),'calculation_basis':'Нижняя граница исключена, верхняя включена. Фиксированная сумма или ставка руб/кг из PDF. Без объёмных тарифов и допуслуг.'}
    for page_no,page in enumerate(reader.pages,1):
        text=page.extract_text(extraction_mode='layout') or '';active=False
        for line in text.splitlines():
            if re.fullmatch(r'\s*Вес,?\s*кг\s*(?:\||;|\s{2,})\s*Цена,?\s*(?:руб\.?|₽)\s*',line,re.I):
                active=True;continue
            if not active or not line.strip():continue
            cells=re.split(r'\s*[|;]\s*|\s{2,}',line.strip())
            if len(cells)!=2:
                active=False;continue
            pid=_profile(cells[0])
            if not pid:
                active=False;continue
            if pid in vals:raise ValueError('В PDF повторяется один контрольный вес; разделы неоднозначны')
            vals[pid]={'kind':'exact','price':number(cells[1]),'source_page':page_no}
            if page_no not in page_numbers:page_numbers.append(page_no)
    if not vals:raise ValueError('Не распознана таблица «Вес, кг | Цена, руб»; используйте XLSX-шаблон')
    return vals,{'parser':'PDF · контрольный вес и цена отправки','source_pages':page_numbers,
                 'document_date':dated(all_text),'route_verified':True,
                 'calculation_basis':'Цена отправки для указанного в текстовом PDF контрольного веса; без интерполяции.'}


def parse_document(raw, filename, company, origin, destination):
    if _SESSION.get() is None:
        with parsing_session():
            return parse_document(raw, filename, company, origin, destination)
    filename=normalize_filename(filename)
    values,meta=_parse_document(raw,filename,company,origin,destination)
    from .source_conditions import document_conditions
    conditions=document_conditions(raw,PurePosixPath(filename).suffix)
    return values,{**conditions,**{k:v for k,v in meta.items() if v is not None or k not in conditions}}


def _parse_document(raw, filename, company, origin, destination):
    ext=validate_file(raw,filename)
    if company not in COMPANIES:raise ValueError('Выберите компанию из списка')
    if ext=='.zip':
        if company!='Возовоз':raise ValueError('ZIP поддерживается для архива тарифов Возовоза; для других компаний выберите PDF/XLSX/XLS')
        return parse_vozovoz_zip(raw,origin,destination)
    if ext=='.pdf':
        from . import scan_ocr
        if scan_ocr.needed(raw):
            return scan_ocr.parse(scan_ocr.prepare(raw,company,origin=origin,destination=destination),origin,destination)
        first=pdf_reader(raw).pages[0].extract_text() or ''
        if 'Тарифы на межтерминальную перевозку из' in first:
            if company!='ДЛ':raise ValueError('Это PDF Деловых Линий: выберите компанию ДЛ')
            return parse_dellin(raw,origin,destination)
        if company in {'КИТ','Новая Линия'}:
            from . import v42_collectors as c
            c._confirm_pdf_origin(raw,origin,company)
            if company=='КИТ':
                # The existing geometry parser accepts paths; temporary file has no fallback identity.
                import tempfile
                from pathlib import Path
                with tempfile.TemporaryDirectory() as tmp:
                    p=Path(tmp)/'user.pdf';p.write_bytes(raw)
                    vals=c._kit_pdf_values(p,origin,destination)
            else:
                minimum,rates,variant=c._newline_pdf_values(raw,destination)
                vals=values_from_tiers([(0,hi,r) for hi,r in rates],minimum=minimum)
            return vals,{'parser':company+' · официальный макет PDF','route_verified':True,'document_date':dated(first)}
        return parse_generic_pdf(raw,company,origin,destination)
    # A standard template is available to all 17 companies, including native parsers.
    if any(is_interval_header(r) for _,rows in workbook_rows(raw) for r in rows[:20]):
        return parse_interval_workbook(raw,company,origin,destination)
    if any([norm(x) for x in r[:5]]==[norm(x) for x in GENERIC_HEADER]
           for _,rows in workbook_rows(raw) for r in rows[:20]):
        return parse_generic_workbook(raw,company,origin,destination)
    if company=='Возовоз':return parse_vozovoz(raw,origin,destination)
    if company=='Werner':return parse_werner(raw,origin,destination)
    if company=='ПЭК' and ext=='.xlsx':return parse_pek(raw,origin,destination)
    if company in {'Фортуна','ЭкспедицияПлюс'} and ext=='.xlsx':
        from .online_tariffs import parse_workbook
        return parse_workbook(company,raw,origin,destination),{'parser':company+' · исходящая таблица XLSX','route_verified':True,
            'calculation_basis':company+': исходящие весовые тарифы; у Фортуны выбран автотариф, отдельная железнодорожная строка исключена. Фиксированные цены и опубликованные руб/кг, без допуслуг.'}
    if company=='Рейл Континент' and ext=='.xlsx':
        from .online_tariffs import parse_railcontinent_workbook
        return parse_railcontinent_workbook(raw,origin,destination),{'parser':'Рейл Континент · XLSX авто','route_verified':True}
    return parse_generic_workbook(raw,company,origin,destination)
