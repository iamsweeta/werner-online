"""Confirmed multi-route price documents. SQLite commits all routes atomically."""
from __future__ import annotations
from .business_time import tariff_today
import hashlib,json,re,sqlite3,threading,time,uuid,io,zipfile
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from . import v42_engine as e
from . import tariff_documents as t
from .cities import city_names,city_pattern

JOBS={};LOCK=threading.RLock();MAX_ROUTES=2000
_MIGRATED=set()


def root():return e.RUNTIME_DIR/'imports'


@contextmanager
def db():
    root().mkdir(parents=True,exist_ok=True)
    conn=sqlite3.connect(root()/'documents.sqlite3',timeout=30);conn.row_factory=sqlite3.Row
    try:
        with conn:
            conn.executescript('''CREATE TABLE IF NOT EXISTS files(id TEXT PRIMARY KEY,company TEXT,meta TEXT,active INTEGER);
              CREATE TABLE IF NOT EXISTS prices(file TEXT,origin TEXT,destination TEXT,profile TEXT,payload TEXT,PRIMARY KEY(file,origin,destination,profile));
              CREATE INDEX IF NOT EXISTS document_route ON prices(origin,destination);
              CREATE TABLE IF NOT EXISTS excluded(company TEXT,origin TEXT,destination TEXT,PRIMARY KEY(company,origin,destination));
              CREATE TABLE IF NOT EXISTS route_documents(company TEXT,origin TEXT,destination TEXT,file TEXT,
                PRIMARY KEY(company,origin,destination));''')
            yield conn
    finally:conn.close()


def _id(value):
    if not re.fullmatch('[a-f0-9]{32}',str(value)):raise ValueError('Документ не найден')
    return value


def pack(origin,destination):
    origin,destination=e.route_pair(origin,destination)
    migrate_legacy(origin,destination)
    if not (root()/'documents.sqlite3').exists():return {'profiles':{},'companies':{}}
    with db() as conn:
        rows=conn.execute('''SELECT f.id,f.rowid AS sequence,f.company,f.meta,f.active,p.profile,p.payload
          FROM prices p JOIN files f ON f.id=p.file WHERE p.origin=? AND p.destination=? ORDER BY f.rowid''',(origin,destination)).fetchall()
        choices=dict(conn.execute('SELECT company,file FROM route_documents WHERE origin=? AND destination=?',(origin,destination)))
        excluded={r[0] for r in conn.execute('SELECT company FROM excluded WHERE origin=? AND destination=?',(origin,destination))}
    out={'profiles':{},'companies':{},'_known_source_files':[]};documents={}
    for row in rows:
        meta=json.loads(row['meta']);value=json.loads(row['payload'])
        if meta.get('source_file'):out['_known_source_files'].append(meta['source_file'])
        entry=documents.setdefault(row['company'],{}).setdefault(row['id'],
            {'meta':{**meta,'document_id':row['id']},'active':bool(row['active']),'sequence':row['sequence'],'values':{}})
        entry['values'][row['profile']]=value
    for company,files in documents.items():
        if company in excluded:continue
        chosen=choices.get(company)
        if chosen in files and files[chosen]['active']:entry=files[chosen];pinned=True
        else:
            entry=max(files.values(),key=lambda f:(int(f['meta'].get('import_revision',0)),str(f['meta'].get('uploaded_at','')),f['sequence']))
            pinned=False
        meta={**entry['meta'],'document_selected':pinned,'_disabled':not entry['active']}
        out['companies'][company]=meta
        if not entry['active']:continue
        for pid,value in entry['values'].items():
            out['profiles'].setdefault(pid,{})[company]={**meta,**value,'document_selected':pinned,'document_id':meta['document_id']}
    return out


def migrate_legacy(origin=None,destination=None):
    """Register confirmed pre-53 route files once, preserving originals and revisions."""
    with e.STATE_LOCK:
        from .document_imports import route_path
        paths=[route_path(origin,destination)] if origin and destination else (root()/'routes').glob('*.json')
        for path in paths:
            try:signature=(str(path.resolve()),path.stat().st_mtime_ns,path.stat().st_size)
            except FileNotFoundError:continue
            if signature in _MIGRATED:continue
            data=e._read_json(path,{})
            for company,meta in data.get('companies',{}).items():
                filename=str(meta.get('source_file',''))
                if company not in e.COMPANIES or not re.fullmatch(r'[a-f0-9]{32}\.(pdf|xls|xlsx|csv|zip)',filename):continue
                if not meta.get('origin') or not meta.get('destination'):continue
                o,d=e.route_pair(meta['origin'],meta['destination'])
                if not e.is_supported_route(o,d):continue
                values={pid:rows[company] for pid,rows in data.get('profiles',{}).items() if company in rows}
                if not values:continue
                from .document_imports import _check_values
                _check_values(values)
                saved={**meta,'uploaded':True,'route_count':1,'values_count':len(values)}
                with db() as conn:
                    if conn.execute('SELECT 1 FROM files WHERE id=?',(filename.split('.')[0],)).fetchone():continue
                    conn.execute('INSERT INTO files VALUES (?,?,?,1)',(filename.split('.')[0],company,json.dumps(saved,ensure_ascii=False)))
                    for pid,value in values.items():conn.execute('INSERT INTO prices VALUES (?,?,?,?,?)',(filename.split('.')[0],o,d,pid,json.dumps(value,ensure_ascii=False)))
            _MIGRATED.add(signature)


def register_single(ident,meta,values,origin,destination):
    migrate_legacy()
    saved={**meta,'uploaded':True,'route_count':1,'values_count':len(values)}
    with db() as conn:
        conn.execute('INSERT INTO files VALUES (?,?,?,1)',(ident,meta['company'],json.dumps(saved,ensure_ascii=False)))
        for pid,value in values.items():conn.execute('INSERT INTO prices VALUES (?,?,?,?,?)',(ident,origin,destination,pid,json.dumps(value,ensure_ascii=False)))
        conn.execute('DELETE FROM excluded WHERE company=? AND origin=? AND destination=?',(meta['company'],origin,destination))
        conn.execute('INSERT OR REPLACE INTO route_documents VALUES (?,?,?,?)',(meta['company'],origin,destination,ident))
    return saved


def document_routes(ident):
    migrate_legacy()
    with db() as conn:
        file=conn.execute('SELECT company,meta FROM files WHERE id=? AND active=1',(_id(ident),)).fetchone()
        if not file:raise ValueError('Документ не найден или отключён')
        routes=conn.execute('SELECT origin,destination,COUNT(*) AS values_count FROM prices WHERE file=? GROUP BY origin,destination ORDER BY origin,destination',(ident,)).fetchall()
    return {'id':ident,'company':file['company'],'meta':json.loads(file['meta']),'routes':[dict(row) for row in routes]}


def route_documents(origin,destination):
    origin,destination=e.route_pair(origin,destination)
    if not e.is_supported_route(origin,destination):raise ValueError('Выберите два разных города')
    current=pack(origin,destination)
    with db() as conn:
        rows=conn.execute('''SELECT f.id,f.company,f.meta,COUNT(*) AS values_count FROM files f JOIN prices p ON p.file=f.id
          WHERE f.active=1 AND p.origin=? AND p.destination=? GROUP BY f.id ORDER BY f.rowid DESC''',(origin,destination)).fetchall()
        total=conn.execute('SELECT COUNT(*) FROM files WHERE active=1').fetchone()[0]
    files=[]
    for row in rows:
        meta=json.loads(row['meta']);selected=current['companies'].get(row['company'],{})
        files.append({**meta,'id':row['id'],'company':row['company'],'route_values_count':row['values_count'],
            'selected':selected.get('document_id')==row['id'] and bool(selected.get('document_selected')),
            'automatic':selected.get('document_id')==row['id'] and not selected.get('document_selected') and not selected.get('_disabled')})
    return {'origin':origin,'destination':destination,'files':files,'total_files':total}


def select_document(company,origin,destination,document_id=None):
    origin,destination=e.route_pair(origin,destination)
    if company not in e.COMPANIES or not e.is_supported_route(origin,destination):raise ValueError('Выберите компанию и направление из списка')
    migrate_legacy()
    with e.STATE_LOCK:
        with db() as conn:
            if document_id:
                match=conn.execute('''SELECT 1 FROM files f JOIN prices p ON p.file=f.id
                  WHERE f.id=? AND f.company=? AND f.active=1 AND p.origin=? AND p.destination=? LIMIT 1''',(_id(document_id),company,origin,destination)).fetchone()
                if not match:raise ValueError('В этом документе нет подтверждённых цен выбранной компании и направления')
                conn.execute('INSERT OR REPLACE INTO route_documents VALUES (?,?,?,?)',(company,origin,destination,document_id))
            else:conn.execute('DELETE FROM route_documents WHERE company=? AND origin=? AND destination=?',(company,origin,destination))
            conn.execute('DELETE FROM excluded WHERE company=? AND origin=? AND destination=?',(company,origin,destination))
        from .document_imports import bump_revision
        bump_revision()
        count=0;filename=None;cleared=0
        if document_id:
            with db() as conn:
                profiles=[r[0] for r in conn.execute('SELECT profile FROM prices WHERE file=? AND origin=? AND destination=?',(document_id,origin,destination))]
                metadata=conn.execute('SELECT meta FROM files WHERE id=?',(document_id,)).fetchone()
            filename=json.loads(metadata['meta']).get('original_filename');count=len(profiles)
            from .manual_prices import clear_covered
            cleared=clear_covered(company,origin,destination,profiles)
        return {'ok':True,'company':company,'origin':origin,'destination':destination,'document_id':document_id,'filename':filename,'applied_prices':count,'replaced_manual':cleared,'applied_at':e._now()}


def disable_route(company,origin,destination):
    if (root()/'documents.sqlite3').exists():
        with db() as conn:conn.execute('INSERT OR REPLACE INTO excluded VALUES (?,?,?)',(company,origin,destination))


def list_files():
    migrate_legacy()
    if not (root()/'documents.sqlite3').exists():return []
    with db() as conn:rows=conn.execute('SELECT id,meta FROM files WHERE active=1 ORDER BY rowid DESC').fetchall()
    return [{'id':r['id'],**json.loads(r['meta'])} for r in rows]


def remove(ident):
    with e.STATE_LOCK:
        with db() as conn:
            if not conn.execute('SELECT id FROM files WHERE id=?',(_id(ident),)).fetchone():raise ValueError('Документ не найден')
            conn.execute('UPDATE files SET active=0 WHERE id=?',(ident,))
            conn.execute('DELETE FROM route_documents WHERE file=?',(ident,))
        from .document_imports import bump_revision
        bump_revision()
    return {'ok':True}


def _candidates(raw,filename,company,origin):
    ext=Path(filename).suffix.lower();pairs=set();origins=set();destinations=set()
    if ext=='.pdf':
        from . import scan_ocr
        if scan_ocr.needed(raw):return scan_ocr.pairs(scan_ocr.prepare(raw,company),origin)
        reader=t.pdf_reader(raw);first=reader.pages[0].extract_text() or ''
        text='\n'.join(p.extract_text() or '' for p in reader.pages)
        if not text.strip():raise ValueError('В PDF нет текстового слоя. Используйте текстовый PDF или XLSX-шаблон для прайса.')
        explicit=re.search(r'^\s*Откуда\s*:\s*(.+)$',text,re.M|re.I)
        target=re.search(r'^\s*Куда\s*:\s*(.+)$',text,re.M|re.I)
        if explicit and target:pairs.add(e.route_pair(explicit[1],target[1]))
        else:
            for city in city_names():
                if re.search(r'\bиз\s+(?:г(?:орода)?\.?\s*)?'+city_pattern(city)+r'(?=\s|[,.;]|$)',first,re.I):origins.add(city)
            destinations=set(city_names())
    else:
        sheets=list(t.workbook_rows(raw));generic=False;explicit_pairs=set();text='';fortuna_pairs=set()
        for title,rows in sheets:
            if company=='Фортуна':
                active_origin=None
                for row in rows:
                    label=str(row[0] or '').strip() if row else ''
                    section=re.match(r'^из\s+г\.\s*(.+)$',label,re.I)
                    if section:active_origin=e.normalize_city(section[1]);continue
                    if re.match(r'^в\s+г\.',label,re.I):active_origin=None
                    dest=e.normalize_city(label)
                    if active_origin and e.is_supported_route(active_origin,dest):fortuna_pairs.add((active_origin,dest))
            for i,r in enumerate(rows[:20]):
                if t.is_template_header(r):
                    generic=True
                    for r in rows[i+1:]:
                        if len(r)>=5 and str(r[0]).strip()==company and r[5 if t.is_interval_header(rows[i]) else 4] not in ('',None):explicit_pairs.add(e.route_pair(r[1],r[2]))
            name=e.normalize_city(title)
            if name in city_names():origins.add(name)
            for r in rows:
                cells=[e.normalize_city(str(v or '').strip(' *')) for v in r[:4]]
                for o,d in zip(cells,cells[1:]):
                    if e.is_supported_route(o,d):pairs.add((o,d))
                destinations.update(c for c in cells if c in city_names())
            text+=' '.join(str(v or '') for r in rows[:25] for v in r[:6])+'\n'
        if generic:
            # A template row explicitly names its carrier; never fall back to a
            # native parser to read another carrier from this same template.
            return sorted(p for p in explicit_pairs if e.is_supported_route(*p) and (not origin or p[0]==origin))
        if company=='Фортуна' and fortuna_pairs:
            return sorted(p for p in fortuna_pairs if not origin or p[0]==origin)
        if not origins:
            for city in city_names():
                if re.search(r'из\s+(?:г(?:орода)?\.?\s*)?'+city_pattern(city)+r'(?=\s|[,.;]|$)',text,re.I):origins.add(city)
    if origin:origins={origin};pairs={p for p in pairs if p[0]==origin}
    if not pairs:
        if not origins:raise ValueError('Не удалось определить город отправления. Выберите его в поле «Откуда» и загрузите файл ещё раз.')
        pairs={(o,d) for o in origins for d in destinations if e.is_supported_route(o,d)}
    pairs=sorted(p for p in pairs if e.is_supported_route(*p))
    if len(pairs)>MAX_ROUTES:raise ValueError('В документе слишком много маршрутов для одной проверки. Выберите город отправления.')
    return pairs


def parse_routes(raw,filename,company,origin=None,on_progress=None,conflicts=None):
    """Read native formats and explicit templates without deriving absent prices."""
    filename=t.normalize_filename(filename)
    t.validate_file(raw,filename)
    if Path(filename).suffix.lower()=='.zip':
        output={};errors=[]
        with t.checked_zip(raw) as z:
            docs=[i for i in z.infolist() if Path(t.normalize_filename(i.filename)).suffix.lower() in {'.pdf','.xls','.xlsx','.csv'}]
            if not docs or len(docs)>30:raise ValueError('В ZIP должно быть от 1 до 30 документов PDF, Excel или CSV')
            for info in docs:
                if info.file_size>t.MAX_BYTES:raise ValueError('Документ внутри ZIP превышает 20 МБ')
                try:parsed,missed=parse_routes(z.read(info),info.filename,company,origin,on_progress,conflicts)
                except Exception as exc:errors.append({'file':info.filename,'message':str(exc)});continue
                for route,item in parsed.items():
                    item['meta']['archive_member']=info.filename
                    for pid,value in item['values'].items():
                        value.update({k:v for k,v in item['meta'].items() if k not in value})
                        value['original_filename']=info.filename
                    if route not in output:output[route]=item
                    else:
                        current=output[route]['values']
                        for pid,value in item['values'].items():
                            if pid not in current:current[pid]=value;continue
                            conflict=next((c for c in (conflicts or []) if (c['origin'],c['destination'],c['profile_id'])==(*route,pid)),None)
                            if t.price_signature(current[pid])==t.price_signature(value) and not conflict:continue
                            if conflicts is None:raise ValueError('В документах разные цены одного маршрута и веса. Выберите нужную редакцию.')
                            if conflict is None:
                                conflict={'id':str(len(conflicts)),'origin':route[0],'destination':route[1],'profile_id':pid,'options':[dict(current[pid])]}
                                conflicts.append(conflict)
                            if not any(t.price_signature(v)==t.price_signature(value) for v in conflict['options']):conflict['options'].append(dict(value))
                            if len(conflicts)>500:raise ValueError('Более 500 конфликтов цен. Оставьте актуальные редакции или выберите один город.')
                errors.extend(missed)
        if not output:raise ValueError('В архиве не найден прайс выбранной компании. '+str(errors[:2]))
        return output,errors
    with t.parsing_session():
        candidates=_candidates(raw,filename,company,origin);output={};errors=[];deadline=time.monotonic()+240
        if Path(filename).suffix.lower()=='.pdf':
            from . import scan_ocr
            if scan_ocr.needed(raw):errors.extend(scan_ocr.prepare(raw,company).get('errors',[]))
        if not candidates:raise ValueError('В документе нет маршрутов выбранной компании или города. Проверьте выбор и формат таблицы.')
        for i,(o,d) in enumerate(candidates):
            if time.monotonic()>deadline:raise ValueError('Проверка документа заняла более 4 минут. Выберите один город отправления или разделите файл.')
            try:
                values,meta=t.parse_document(raw,filename,company,o,d)
                from .document_imports import _check_values
                _check_values(values)
                output[(o,d)]={'values':values,'meta':meta}
            except Exception as exc:errors.append({'origin':o,'destination':d,'message':str(exc)[:500]})
            if on_progress:on_progress(i+1,len(candidates),len(output))
        if not output:raise ValueError('Цены не распознаны. '+ ' '.join(dict.fromkeys(x['message'] for x in errors))[:1200])
        return output,errors


def start_preview(raw,filename,company,origin=None,document_date=None):
    if company not in e.COMPANIES:raise ValueError('Выберите компанию')
    origin=e.normalize_city(origin) if origin else None
    if origin and origin not in city_names():raise ValueError('Выберите город из списка')
    try:t.validate_file(raw,filename)
    except Exception as exc:raise ValueError('Не удалось прочитать документ: '+str(exc)[:500]) from exc
    if document_date:
        if date.fromisoformat(document_date)>tariff_today():raise ValueError('Дата тарифов ещё не наступила')
    token=uuid.uuid4().hex;name=t.normalize_filename(filename)
    pending=root()/'multi_pending';pending.mkdir(parents=True,exist_ok=True)
    with LOCK:
        for key in list(JOBS):
            if JOBS[key]['status']!='parsing' and e.age_seconds(JOBS[key].get('created_at'))>1800:JOBS.pop(key,None)
        if sum(j['status']=='parsing' for j in JOBS.values())>=2:raise ValueError('Дождитесь завершения текущих документов')
        active={k for k,v in JOBS.items() if v['status']=='parsing'}
        for f in pending.iterdir():
            try:
                if f.stem not in active and time.time()-f.stat().st_mtime>1800:f.unlink(missing_ok=True)
            except FileNotFoundError:pass
        (pending/(token+Path(name).suffix.lower())).write_bytes(raw)
        JOBS[token]={'token':token,'status':'parsing','company':company,'filename':name,'created_at':e._now(),
                     'done':0,'total':0,'matched':0,'message':'Читаю документ…'}
    def run():
        from . import scan_ocr
        def ocr_progress(info):
            with LOCK:JOBS[token].update(message=info['message'],ocr_page=info['done'],ocr_total=info['total'])
        progress_token=scan_ocr._PROGRESS.set(ocr_progress)
        try:
            def progress(done,total,matched):
                with LOCK:JOBS[token].update(done=done,total=total,matched=matched)
            conflicts=[]
            parsed,errors=parse_routes(raw,name,company,origin,progress,conflicts)
            routes=[];dates=set();warnings=[]
            if name!=filename:warnings.append('Имя файла исправлено: '+name+'. Содержимое проверено отдельно.')
            for (o,d),item in parsed.items():
                stamp=item['meta'].get('document_date') or document_date
                if stamp:
                    if date.fromisoformat(stamp)>tariff_today():raise ValueError('В документе указана будущая дата тарифов')
                    dates.add(stamp)
                for value in item['values'].values():
                    if not value.get('document_date'):value['document_date']=stamp
                    if value.get('document_date'):
                        if date.fromisoformat(value['document_date'])>tariff_today():raise ValueError('В документе указана будущая дата тарифов')
                        dates.add(value['document_date'])
                item['meta']['document_date']=stamp
                routes.append({'origin':o,'destination':d,**item})
            for conflict in conflicts:
                for value in conflict['options']:
                    if not value.get('document_date'):value['document_date']=document_date
                    if value.get('document_date') and date.fromisoformat(value['document_date'])>tariff_today():raise ValueError('В документе указана будущая дата тарифов')
            if conflicts:warnings.append('В файлах есть разные цены для одного маршрута и веса. Выберите нужную версию каждой цены ниже.')
            has_ocr=any(r['meta'].get('ocr') for r in routes)
            if has_ocr:warnings.insert(0,scan_ocr.WARNING)
            warnings.append('Загрузка заменит ранее сохранённый прайс этой компании для распознанных маршрутов. Части одного прайса загружайте вместе.')
            if not dates:warnings.append('Дата тарифов не определена. Проверьте актуальность документа.')
            if any((tariff_today()-date.fromisoformat(s)).days>30 for s in dates):warnings.append('В документе есть тарифы старше 30 дней. Проверьте, что они ещё действуют.')
            warnings.append('После подтверждения документ станет выбранным источником его направлений в «Одном маршруте» и большой таблице. Вернуть онлайн-приоритет можно выбором «Автоматически».')
            metadata={'token':token,'company':company,'original_filename':name,'extension':Path(name).suffix.lower(),
                      'sha256':hashlib.sha256(raw).hexdigest(),'created_at':e._now(),'document_date':next(iter(dates)) if len(dates)==1 else None,
                      'ocr':has_ocr}
            e._robust_json_write(pending/(token+'.json'),{'meta':metadata,'routes':routes,'conflicts':conflicts})
            (pending/(token+Path(name).suffix.lower())).touch()
            with LOCK:JOBS[token].update(created_at=metadata['created_at'],status='ready',routes=routes,conflicts=conflicts,warnings=warnings,errors=errors[:30],
                skipped=len(errors),matched=len(routes),values_count=sum(len(r['values']) for r in routes),meta=metadata,message='Проверьте распознанные цены перед сохранением')
        except Exception as exc:
            with LOCK:JOBS[token].update(status='error',message=str(exc)[:1500])
        finally:scan_ocr._PROGRESS.reset(progress_token)
    threading.Thread(target=run,daemon=True).start()
    return status(token)


def start_many(files,company,origin=None,document_date=None):
    if not files or len(files)>30:raise ValueError('Выберите от 1 до 30 документов')
    if len(files)==1:return start_preview(files[0][1],files[0][0],company,origin,document_date)
    if sum(len(raw) for _,raw in files)>t.MAX_BYTES:raise ValueError('Общий размер выбранных файлов — не более 20 МБ')
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w',zipfile.ZIP_DEFLATED) as archive:
        for i,(name,raw) in enumerate(files,1):
            clean=t.normalize_filename(name)
            if Path(clean).suffix.lower() not in {'.pdf','.xlsx','.xls','.csv'}:raise ValueError('Для нескольких файлов выберите PDF, Excel или CSV. ZIP загрузите отдельно.')
            archive.writestr(f'{i:02d}_{clean}',raw)
    return start_preview(buffer.getvalue(),'Прайсы_'+company+'.zip',company,origin,document_date)


def status(token):
    with LOCK:
        result=JOBS.get(_id(token))
        if not result:raise ValueError('Предпросмотр недоступен. Загрузите файл заново.')
        return json.loads(json.dumps(result))


def commit(token,resolutions=None):
    ident=_id(token);path=root()/'multi_pending'/(ident+'.json')
    with e.STATE_LOCK:
        data=e._read_json(path,{})
        if not data or e.age_seconds(data['meta']['created_at'])>1800:raise ValueError('Предпросмотр истёк. Загрузите файл заново.')
        conflicts=data.get('conflicts',[]);resolutions=resolutions or {}
        if set(resolutions)!={c['id'] for c in conflicts}:raise ValueError('Выберите цену для каждого конфликта в предпросмотре')
        for conflict in conflicts:
            choice=resolutions[conflict['id']]
            if isinstance(choice,bool) or not isinstance(choice,int) or not 0<=choice<len(conflict['options']):raise ValueError('Некорректный выбор цены')
            route=next(r for r in data['routes'] if (r['origin'],r['destination'])==(conflict['origin'],conflict['destination']))
            route['values'][conflict['profile_id']]=conflict['options'][choice]
        meta=data['meta'];source=path.with_suffix(meta['extension']);raw=source.read_bytes()
        if hashlib.sha256(raw).hexdigest()!=meta['sha256']:raise ValueError('Файл изменился после проверки')
        target=root()/'files'/(ident+meta['extension']);target.parent.mkdir(exist_ok=True)
        target.write_bytes(raw)
        from .document_imports import revision
        saved={**meta,'import_revision':revision()+1,'source_file':target.name,'uploaded_at':e._now(),'captured_at':e._now(),
               'source_type':'Файл пользователя','data_origin':'uploaded','online':False,'uploaded':True,
               'route_count':len(data['routes']),'values_count':sum(len(r['values']) for r in data['routes'])}
        from .document_imports import _check_values
        with db() as conn:
            conn.execute('INSERT INTO files VALUES (?,?,?,1)',(ident,meta['company'],json.dumps(saved,ensure_ascii=False)))
            for route in data['routes']:
                _check_values(route['values']);o,d=route['origin'],route['destination']
                if not e.is_supported_route(o,d):raise ValueError('Некорректный маршрут')
                conn.execute('DELETE FROM excluded WHERE company=? AND origin=? AND destination=?',(meta['company'],o,d))
                conn.execute('INSERT OR REPLACE INTO route_documents VALUES (?,?,?,?)',(meta['company'],o,d,ident))
                for pid,value in route['values'].items():
                    conn.execute('INSERT INTO prices VALUES (?,?,?,?,?)',(ident,o,d,pid,json.dumps({**route['meta'],**value},ensure_ascii=False)))
        from .manual_prices import clear_covered
        for route in data['routes']:clear_covered(meta['company'],route['origin'],route['destination'],route['values'])
        path.unlink();source.unlink()
        from .document_imports import bump_revision
        bump_revision()
        with LOCK:
            if ident in JOBS:JOBS[ident]['status']='committed'
        return {'ok':True,'id':ident,**saved}


def template(company,origin=None,kind='interval'):
    """A blank multi-route workbook for carriers without a recognized layout."""
    if company not in e.COMPANIES:raise ValueError('Выберите компанию')
    if kind not in {'interval','points'}:raise ValueError('Неизвестный шаблон')
    if kind=='interval':return interval_template(company,origin)
    from .bulk_refresh import plan_routes
    routes=plan_routes('origins',[origin]) if origin else plan_routes('reference')
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb=Workbook();ws=wb.active;ws.title='Прайс-лист';ws.append(t.GENERIC_HEADER)
    for o,d in routes:
        for profile in e.COMMON_PROFILES:ws.append([company,o,d,'МИН' if profile['id']=='min' else profile['weight_kg'],None])
    ws.freeze_panes='D2';ws.auto_filter.ref=ws.dimensions
    for col,width in [('A',24),('B',26),('C',28),('D',16),('E',22)]:ws.column_dimensions[col].width=width
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='245AD6');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=32
    dv=DataValidation(type='decimal',operator='greaterThan',formula1=0,allow_blank=True)
    dv.errorTitle='Проверьте цену';dv.error='Введите положительную итоговую стоимость отправки в рублях.';dv.showErrorMessage=True;dv.errorStyle='stop'
    ws.add_data_validation(dv);dv.add(f'E2:E{ws.max_row}')
    help=wb.create_sheet('Инструкция');help.column_dimensions['A'].width=120
    for line in ['Шаблон импорта прайса; реальные цены не заполнены.',
                 'Внесите итоговую стоимость отправки в рублях, а не ставку за килограмм.',
                 'Заполняйте только подтверждённые строки. Пустая цена пропускается.',
                 'МИН — опубликованный минимум отправления. Не выдумывайте отсутствующие цены.',
                 'Названия компании, городов и контрольные веса уже подготовлены.',
                 'Удалите ненужные маршруты или оставьте их цены пустыми.',
                 'Сохраните XLSX и загрузите в разделе «Прайс-листы» для этой компании.',
                 'Для объединения нескольких частей прайса выберите все файлы в одной загрузке.',
                 'Укажите дату тарифов в приложении и проверьте цены перед сохранением.']:
        help.append([line])
    output=io.BytesIO();wb.save(output);return output.getvalue()


def interval_template(company,origin=None):
    from .bulk_refresh import plan_routes
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    routes=plan_routes('origins',[origin]) if origin else plan_routes('reference')
    wb=Workbook();ws=wb.active;ws.title='Диапазоны';ws.append(t.INTERVAL_HEADER)
    for o,d in routes:
        lo=0
        for p in e.COMMON_PROFILES:
            if p.get('is_minimum_profile'):continue
            hi=p['weight_kg']
            ws.append([company,o,d,lo,hi,None,'руб' if hi<=50 else 'руб/кг',None]);lo=hi
    ws.freeze_panes='F2';ws.auto_filter.ref=ws.dimensions
    for col,width in zip('ABCDEFGH',[23,25,28,18,18,18,20,20]):ws.column_dimensions[col].width=width
    for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='245AD6');cell.font=Font(color='FFFFFF',bold=True);cell.alignment=Alignment(wrap_text=True)
    ws.row_dimensions[1].height=32
    units=DataValidation(type='list',formula1='"руб,руб/кг"');units.showErrorMessage=True;units.errorStyle='stop';ws.add_data_validation(units);units.add(f'G2:G{ws.max_row}')
    help=wb.create_sheet('Инструкция');help.column_dimensions['A'].width=115
    for line in [
        'Шаблон диапазонов. Цены пустые; перенесите только подтверждённые тарифы из оригинала.',
        'Замените границы на интервалы перевозчика. Нижняя граница не включается, верхняя включается.',
        'Фиксированный тариф: единица руб; ставка за килограмм: руб/кг. Не делите фиксированную сумму на вес.',
        'Один интервал может заменить несколько подготовленных строк. Удалите перекрывающиеся строки.',
        'При одинаковых границах получается отдельный контрольный вес. Для открытого верха напишите ∞.',
        'Минимум, руб — только явно опубликованная минимальная плата для ставки руб/кг. Иначе оставьте пустым.',
        'В Excel сравнения до 50 кг выводится сумма, от 100 кг — опубликованная ставка, МИН — наименьшая подтверждённая сумма отправки.',
        'Контрольный вес — верхняя граница общей строки. Не переносите ставку на веса вне исходного интервала.',
        'Пустые тарифы пропускаются. Цена из файла не получает статус LIVE.',
        'Для старого формата отдельных итогов отправки доступен шаблон «Контрольные веса». Он не подтверждает ставку за кг.',
        'Сохраните файл, загрузите его для выбранной компании и проверьте предпросмотр перед подтверждением.']:
        help.append([line])
    out=io.BytesIO();wb.save(out);return out.getvalue()
