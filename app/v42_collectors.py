from __future__ import annotations

import base64, hashlib, io, json, os, re, subprocess, tempfile, threading, time, traceback

import requests
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlencode, urljoin, urlsplit, urlunsplit

from bs4 import BeautifulSoup
from .cities import city_pattern, city_slug, route_slug, UnpublishedTariff
from . import carrier_catalogs
from .network import download, download_ranges, DownloadError, RejectedResponse, explain_error, decode_console

from .v42_engine import COMMON_PROFILES, PROFILE_BY_ID, RUNTIME_DIR, save_live_update, begin_live_attempt, finish_live_attempt, normalize_city

LOG_PATH=RUNTIME_DIR/"v42_collect.log"
_BROWSER_LOCK=threading.Lock()

def _log(msg: str, console_message: str | None=None) -> None:
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    stamp=datetime.now().astimezone().isoformat(timespec="seconds")
    with LOG_PATH.open("a",encoding="utf-8") as f: f.write(f"[{stamp}] {msg}\n")
    print(f"[v42] {console_message or msg}", flush=True)

def _norm(s: str) -> str:
    return " ".join(str(s or "").lower().replace("ё","е").split())

def _num(s: Any) -> float|None:
    t=str(s or "").replace("\xa0"," ").replace(",",".")
    m=re.search(r"(?<!\d)(\d+(?:[ .]\d{3})*(?:\.\d+)?)",t)
    if not m:return None
    try:return float(m.group(1).replace(" ",""))
    except:return None

def _make_driver():
    from selenium import webdriver
    errors=[]
    try:
        opt=webdriver.EdgeOptions(); opt.add_argument("--headless=new"); opt.add_argument("--disable-gpu"); opt.add_argument("--no-first-run"); opt.add_argument("--disable-extensions"); opt.add_argument("--window-size=1400,1000")
        return webdriver.Edge(options=opt),"Edge"
    except Exception as e: errors.append(f"Edge: {e}")
    try:
        opt=webdriver.ChromeOptions(); opt.add_argument("--headless=new"); opt.add_argument("--disable-gpu"); opt.add_argument("--no-first-run"); opt.add_argument("--disable-extensions"); opt.add_argument("--window-size=1400,1000")
        return webdriver.Chrome(options=opt),"Chrome"
    except Exception as e: errors.append(f"Chrome: {e}")
    raise RuntimeError("; ".join(errors))


def _curl_fetch(url: str, *, method: str = "GET", fields: list[tuple[str,str]] | None = None, referer: str | None = None, cookie_jar: str | None = None, timeout: int = 40) -> bytes:
    exe = "curl.exe" if os.name == "nt" else "curl"
    base=[exe,"-L","--fail","--silent","--show-error","--connect-timeout","10","--max-time",str(timeout),
          "-A","Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36"]
    if os.name == "nt": base += ["--ssl-no-revoke"]
    if referer: base += ["-e",referer]
    if cookie_jar: base += ["-b",cookie_jar,"-c",cookie_jar]
    target=url
    data_args=[]
    if method.upper()=="POST": data_args=["-X","POST","--data",urlencode(fields or [])]
    elif fields:
        sep="&" if "?" in target else "?"; target += sep + urlencode(fields)
    attempts=[[]]
    if os.name == "nt":
        # Some Russian carrier sites terminate Schannel's default handshake.
        # A forced IPv4 / HTTP/1.1 / TLS1.2 retry uses the same trusted cert
        # store without disabling certificate verification.
        attempts.append(["-4","--http1.1","--tlsv1.2","--tls-max","1.2"])
    errors=[]
    for extra in attempts:
        cmd=base+extra+data_args+[target]
        try:
            cp=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+5,check=False)
        except Exception as exc:
            errors.append(str(exc)); continue
        if cp.returncode==0 and cp.stdout: return bytes(cp.stdout)
        errors.append((cp.stderr or b"curl failed").decode("utf-8","replace")[-800:])
    raise RuntimeError(" | ".join(errors[-2:]) or "curl failed")

def _powershell_fetch(url: str, *, timeout: int = 45) -> bytes:
    """Windows-native download without interpolating the URL into PowerShell code."""
    with tempfile.TemporaryDirectory(prefix="tariff_ps_") as td:
        out=Path(td)/"response.bin"
        env=os.environ.copy(); env["V42_URL"]=url; env["V42_OUT"]=str(out)
        script=(
            "$ErrorActionPreference='Stop'; $ProgressPreference='SilentlyContinue'; "
            "$h=@{'User-Agent'='Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}; "
            "Invoke-WebRequest -UseBasicParsing -Uri $env:V42_URL -Headers $h -OutFile $env:V42_OUT"
        )
        cp=subprocess.run(["powershell.exe","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-Command",script],
                          stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+10,check=False,env=env)
        if cp.returncode!=0 or not out.exists() or out.stat().st_size==0:
            err=decode_console(cp.stderr or cp.stdout or b"PowerShell download failed")[-1000:]
            raise RuntimeError(err)
        return out.read_bytes()

def _winhttp_fetch(url: str, *, timeout: int = 45) -> bytes:
    """Last Windows-native HTTPS transport using WinHTTP/COM and normal cert validation."""
    with tempfile.TemporaryDirectory(prefix="tariff_winhttp_") as td:
        out=Path(td)/"response.bin"
        env=os.environ.copy(); env["V42_URL"]=url; env["V42_OUT"]=str(out); env["V42_TIMEOUT_MS"]=str(max(5000,int(timeout*1000)))
        script=(
            "$ErrorActionPreference='Stop'; "
            "$u=$env:V42_URL; $o=$env:V42_OUT; $ms=[int]$env:V42_TIMEOUT_MS; "
            "$w=New-Object -ComObject 'WinHttp.WinHttpRequest.5.1'; "
            "$w.SetTimeouts(10000,10000,$ms,$ms); $w.Open('GET',$u,$false); "
            "$w.SetRequestHeader('User-Agent','Mozilla/5.0 (Windows NT 10.0; Win64; x64)'); $w.Send(); "
            "if($w.Status -lt 200 -or $w.Status -ge 400){throw ('HTTP '+$w.Status)}; "
            "$s=New-Object -ComObject ADODB.Stream; $s.Type=1; $s.Open(); $s.Write($w.ResponseBody); $s.SaveToFile($o,2); $s.Close()"
        )
        cp=subprocess.run(["powershell.exe","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-Command",script],
                          stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+15,check=False,env=env)
        if cp.returncode!=0 or not out.exists() or out.stat().st_size==0:
            err=decode_console(cp.stderr or cp.stdout or b"WinHTTP download failed")[-1000:]
            raise RuntimeError(err)
        return out.read_bytes()

def _browser_fetch_bytes(url: str, referer: str) -> bytes:
    driver=None
    with _BROWSER_LOCK:
        try:
            driver,browser=_make_driver(); driver.set_page_load_timeout(35); driver.get(referer); time.sleep(.5)
            result=driver.execute_async_script(r"""
              const done=arguments[arguments.length-1], u=arguments[0];
              fetch(u,{credentials:'include',cache:'no-store'}).then(async r=>{
                const b=new Uint8Array(await r.arrayBuffer()); let s='';
                for(let i=0;i<b.length;i+=0x8000)s+=String.fromCharCode(...b.subarray(i,i+0x8000));
                done({ok:r.ok,status:r.status,url:r.url,data:btoa(s)});
              }).catch(e=>done({ok:false,error:String(e)}));
            """,url) or {}
            if not result.get('ok') or not result.get('data'): raise RuntimeError(str(result.get('error') or result.get('status')))
            return base64.b64decode(result['data'])
        finally:
            if driver:
                try: driver.quit()
                except: pass

def _download_bytes(url: str, *, referer: str | None = None, timeout: int = 50) -> tuple[bytes,str]:
    result=download(url,referer=referer,timeout=timeout)
    return result.content,result.transport


def _browser_navigate_html(url: str, *, timeout: int = 40) -> tuple[bytes,str,str]:
    """Open an HTML page as a real browser navigation and return rendered HTML.

    Some carrier WAFs reject requests/curl and even same-origin ``fetch()`` while
    allowing an ordinary Edge/Chrome navigation.  This helper is HTML-only: it
    never turns downloaded binary files into LIVE evidence.
    """
    driver=None
    with _BROWSER_LOCK:
        try:
            driver,browser=_make_driver(); driver.set_page_load_timeout(timeout); driver.get(url); time.sleep(1.0)
            html=str(driver.page_source or '')
            final=str(driver.current_url or url)
            if len(html.strip())<80:
                raise RuntimeError('browser navigation returned an empty/too short page')
            body=_norm(html)
            if '401 unauthorized' in body or '<title>401' in body:
                raise RuntimeError('browser navigation returned HTTP 401 page')
            return html.encode('utf-8'),'browser-'+str(browser),final
        finally:
            if driver:
                try: driver.quit()
                except Exception: pass

def _download_html_bytes(url: str, *, referer: str | None = None, timeout: int = 50) -> tuple[bytes,str,str]:
    result=download(url,referer=referer,expected="HTML",timeout=timeout)
    return result.content,result.transport,result.url


def _browser_enabled():
    from .legacy_backend import SETTINGS_PATH
    try:return json.loads(SETTINGS_PATH.read_text(encoding='utf-8')).get('public_browser_enabled') is True
    except (OSError,ValueError):return False


def _fetch_submitted_form_html_browser(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Submit the site's real route form with fetch(), without visible windows.

    The browser is only a same-origin HTTP client here. We do not click custom
    widgets and we never read a result before validating the response route.
    """
    driver=None
    with _BROWSER_LOCK:
        try:
            driver,browser=_make_driver(); driver.set_page_load_timeout(35); driver.get(url); time.sleep(1.0)
            script=r"""
            const done=arguments[arguments.length-1];
            const origin=arguments[0], dest=arguments[1];
            const norm=s=>(s||'').trim().toLowerCase().replace(/ё/g,'е');
            (async()=>{
              try{
                const forms=[...document.forms];
                let hit=null, city=[];
                for(const form of forms){
                  const selects=[...form.querySelectorAll('select')];
                  const candidates=selects.filter(s=>{const opts=[...s.options].map(o=>norm(o.textContent));return opts.includes(norm(origin))&&opts.includes(norm(dest));});
                  if(candidates.length>=2){hit=form;city=candidates;break;}
                }
                if(!hit){ done({ok:false,error:'route form with two city selects not found'}); return; }
                function setText(s,text){ const o=[...s.options].find(x=>norm(x.textContent)===norm(text)); if(!o)return false; s.value=o.value; [...s.options].forEach(x=>x.selected=(x===o)); s.dispatchEvent(new Event('change',{bubbles:true})); return true; }
                if(!setText(city[0],origin)||!setText(city[1],dest)){done({ok:false,error:'city options not set'});return;}
                const fd=new FormData(hit);
                const submit=hit.querySelector('button[type=submit],input[type=submit],button:not([type])');
                if(submit && submit.name && !fd.has(submit.name)) fd.append(submit.name, submit.value||submit.textContent||'1');
                let action=hit.action||location.href, method=(hit.method||'GET').toUpperCase();
                let resp;
                if(method==='GET'){
                  const qs=new URLSearchParams(fd).toString();
                  const u=new URL(action,location.href); u.search=qs; resp=await fetch(u.toString(),{credentials:'include',redirect:'follow'});
                } else {
                  resp=await fetch(action,{method,body:fd,credentials:'include',redirect:'follow'});
                }
                const text=await resp.text(); done({ok:resp.ok,status:resp.status,url:resp.url,text,method,from:city[0].value,to:city[1].value});
              }catch(e){done({ok:false,error:String(e&&e.stack||e)});}
            })();
            """
            result=driver.execute_async_script(script,origin,destination) or {}
            if not result.get("ok") or not result.get("text"):
                raise RuntimeError(str(result.get("error") or f"HTTP {result.get('status')}"))
            return str(result["text"]),str(result.get("url") or url),browser
        finally:
            if driver:
                try:driver.quit()
                except:pass


def _matching_city_selects(form, origin: str, destination: str):
    wanted={_norm(origin),_norm(destination)}
    found=[]
    for select in form.find_all("select"):
        labels={_norm(o.get_text(" ",strip=True)) for o in select.find_all("option")}
        if wanted.issubset(labels) and select.get("name"):
            found.append(select)
    return found

def _option_value(select, text: str) -> str|None:
    wanted=_norm(text)
    for option in select.find_all("option"):
        if _norm(option.get_text(" ",strip=True))==wanted:
            return str(option.get("value") if option.get("value") is not None else option.get_text(" ",strip=True))
    return None

def _successful_form_fields(form) -> list[tuple[str,str]]:
    fields=[]
    for tag in form.find_all(["input","select","textarea"]):
        name=tag.get("name")
        if not name or tag.has_attr("disabled"): continue
        if tag.name=="input":
            typ=_norm(tag.get("type") or "text")
            if typ in {"button","reset","file","image"}: continue
            if typ in {"checkbox","radio"} and not tag.has_attr("checked"): continue
            fields.append((name,str(tag.get("value") or "")))
        elif tag.name=="textarea":
            fields.append((name,tag.get_text()))
        else:
            selected=tag.find_all("option",selected=True)
            if not selected:
                first=tag.find("option")
                selected=[first] if first else []
            for option in selected:
                fields.append((name,str(option.get("value") if option.get("value") is not None else option.get_text(" ",strip=True))))
    submit=form.find(["button","input"],attrs={"type":re.compile(r"^submit$",re.I)})
    if submit and submit.get("name"):
        fields.append((str(submit.get("name")),str(submit.get("value") or submit.get_text(" ",strip=True) or "1")))
    return fields

def _replace_field(fields:list[tuple[str,str]], name:str, value:str) -> list[tuple[str,str]]:
    out=[(k,v) for k,v in fields if k!=name]
    out.append((name,value))
    return out

def _fetch_submitted_form_html_http(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Submit the site's native route form directly over HTTP.

    This is the primary path. It mirrors a normal browser form submit but does not
    depend on Select2, JavaScript clicks, WebDriver or a visible browser window.
    """
    session=requests.Session()
    headers={
        "User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36",
        "Accept":"text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language":"ru-RU,ru;q=0.9,en;q=0.6",
        "Cache-Control":"no-cache",
    }
    r=session.get(url,headers=headers,timeout=(10,30),allow_redirects=True)
    r.raise_for_status()
    soup=BeautifulSoup(r.text,"lxml")
    hit=None; city=[]
    for form in soup.find_all("form"):
        c=_matching_city_selects(form,origin,destination)
        if len(c)>=2:
            hit=form; city=c[:2]; break
    if hit is None:
        raise RuntimeError("HTTP: форма маршрута с двумя списками городов не найдена")
    from_value=_option_value(city[0],origin); to_value=_option_value(city[1],destination)
    if from_value is None or to_value is None:
        raise RuntimeError("HTTP: значения городов не найдены в форме")
    fields=_successful_form_fields(hit)
    fields=_replace_field(fields,str(city[0].get("name")),from_value)
    fields=_replace_field(fields,str(city[1].get("name")),to_value)
    action=urljoin(r.url,hit.get("action") or r.url)
    method=_norm(hit.get("method") or "get").upper()
    req_headers={**headers,"Referer":r.url}
    if method=="POST":
        rr=session.post(action,data=fields,headers=req_headers,timeout=(10,35),allow_redirects=True)
    else:
        rr=session.get(action,params=fields,headers=req_headers,timeout=(10,35),allow_redirects=True)
    rr.raise_for_status()
    if not rr.text.strip(): raise RuntimeError("HTTP: сайт вернул пустой ответ")
    return rr.text,rr.url,"HTTP"


def _fetch_submitted_form_html_curl(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Same native-form submission as HTTP path, but through system curl.
    This is important on Windows hosts where Python/OpenSSL is rejected by the carrier TLS endpoint.
    """
    with tempfile.TemporaryDirectory() as td:
        jar=str(Path(td)/"cookies.txt")
        raw=_curl_fetch(url,cookie_jar=jar,timeout=40)
        text=raw.decode("utf-8","replace")
        soup=BeautifulSoup(text,"lxml")
        hit=None; city=[]
        for form in soup.find_all("form"):
            c=_matching_city_selects(form,origin,destination)
            if len(c)>=2: hit=form; city=c[:2]; break
        if hit is None: raise RuntimeError("curl: форма маршрута с двумя списками городов не найдена")
        a=_option_value(city[0],origin); b=_option_value(city[1],destination)
        if a is None or b is None: raise RuntimeError("curl: значения городов не найдены")
        fields=_successful_form_fields(hit)
        fields=_replace_field(fields,str(city[0].get("name")),a); fields=_replace_field(fields,str(city[1].get("name")),b)
        action=urljoin(url,hit.get("action") or url); method=_norm(hit.get("method") or "get").upper()
        out=_curl_fetch(action,method=method,fields=fields,referer=url,cookie_jar=jar,timeout=45)
        return out.decode("utf-8","replace"),action,"curl"

def _fetch_validated_route_table(url: str, origin: str, destination: str):
    """Get a route table with strict route validation.

    Direct HTTP is preferred. A fully headless browser is used only if the site
    requires JavaScript. A failed attempt never changes the saved route pack.
    """
    errors=[]
    for label,fn in (("HTTP",_fetch_submitted_form_html_http),("curl",_fetch_submitted_form_html_curl),("headless",_fetch_submitted_form_html_browser)):
        try:
            html,final_url,transport=fn(url,origin,destination)
            soup=BeautifulSoup(html,"lxml")
            headers,cells,headings=_find_route_table(soup,origin,destination)
            return headers,cells,headings,final_url,transport
        except Exception as exc:
            errors.append(f"{label}: {type(exc).__name__}: {exc}")
            _log(f"route form {url} {label} failed: {type(exc).__name__}: {exc}")
    raise RuntimeError("; ".join(errors))

def _fetch_current_route_table(url: str, origin: str, destination: str):
    """Read an already-rendered CTS tariff table without assuming form controls."""
    raw,transport,final_url=_download_html_bytes(url,referer=url,timeout=22)
    soup=BeautifulSoup(raw.decode("utf-8","replace"),"lxml")
    headers,cells,headings=_find_route_table(soup,origin,destination)
    return headers,cells,headings,final_url,transport

# Compatibility alias used by older tests/instrumentation. It is intentionally
# headless-only; new collectors use _fetch_validated_route_table instead.
def _fetch_submitted_form_html(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    return _fetch_submitted_form_html_browser(url,origin,destination)

def _find_route_table(soup: BeautifulSoup, origin: str, destination: str):
    wanted=_norm(destination)
    if "тарифы уточняйте у наших менеджеров" in _norm(soup.get_text(" ",strip=True)):
        raise UnpublishedTariff("CTSGroup: на выбранное направление сайт предлагает уточнять тариф у менеджера")
    # Strong route validation: a heading must identify the selected origin when present.
    headings=" ".join(x.get_text(" ",strip=True) for x in soup.find_all(["h1","h2","h3","h4"]))
    if not any(re.search(r'из\s+города\s+'+re.escape(_norm(origin))+r'(?:\s|$)', _norm(text.get_text(' ',strip=True))) for text in soup.find_all(['h1','h2','h3','h4'])):
        raise RuntimeError(f"ответ относится к другому городу отправления: {headings[:240]}")
    for table in soup.find_all("table"):
        headers=[x.get_text(" ",strip=True) for x in table.find_all("th")]
        for tr in table.find_all("tr"):
            cells=[x.get_text(" ",strip=True) for x in tr.find_all(["th","td"])]
            if cells and _norm(cells[0])==wanted:
                return headers,cells,headings
    raise UnpublishedTariff(f"CTSGroup: в исходящей таблице нет строки назначения {destination}")


def _cts_parse_xlsx_bytes(raw: bytes, source_url: str, destination: str) -> tuple[list[str],list[str]]:
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    wanted=_norm(destination)
    for ws in wb.worksheets:
        recent_headers=[]
        for row in ws.iter_rows(values_only=True):
            vals=["" if v is None else str(v).strip() for v in row]
            joined=" | ".join(vals)
            if any(k in _norm(joined) for k in ("город получения","цена","до 100","1-100")):
                recent_headers=vals
            first=next((v for v in vals[:4] if v),"")
            if _norm(first)==wanted:
                nums=[_num(v) for v in vals]
                if sum(x is not None for x in nums)>=8:
                    return recent_headers,vals
    raise RuntimeError(f"CTS XLSX: строка {destination} с тарифными колонками не найдена")


def _cts_xlsx_candidates(page_html: str, page_url: str, origin: str) -> list[str]:
    soup=BeautifulSoup(page_html,"lxml"); out=[]
    for a in soup.find_all("a",href=True):
        href=urljoin(page_url,str(a.get("href")))
        if not href.lower().split("?",1)[0].endswith(".xlsx"): continue
        decoded=unquote(href)
        variants=[]
        if _norm(origin)==_norm("Москва"):
            # The page defaults to Moscow, so its published XLSX link is already
            # the correct source for Москва → Санкт-Петербург.
            variants.append(decoded)
        else:
            # For Санкт-Петербург the public page still initially renders the
            # Moscow download link. Try the official SPB filename variants; each
            # candidate is accepted only after the workbook itself validates the
            # destination row and tariff columns.
            for old,new in (("МСК","СПБ"),("мск","спб"),("MSK","SPB"),("Moscow","SPB"),("Москва","СПБ")):
                if old in decoded: variants.append(decoded.replace(old,new))
        for v in variants:
            parts=urlsplit(v); enc_path=quote(unquote(parts.path),safe="/%:@")
            candidate=urlunsplit((parts.scheme,parts.netloc,enc_path,parts.query,parts.fragment))
            if candidate not in out: out.append(candidate)
    return out

def _cts_recent_xlsx_candidates(origin: str, days_back: int = 70) -> list[str]:
    """Generate origin-specific official XLSX filename candidates."""
    today=datetime.now().date(); dates=[]
    monday=today-timedelta(days=today.weekday())
    for i in range(max(2,days_back//7+1)): dates.append(monday-timedelta(days=7*i))
    dates += [today,today.replace(day=1),(today.replace(day=1)-timedelta(days=1)).replace(day=1)]
    seen=[]
    for d in dates:
        if d not in seen: seen.append(d)
    if _norm(origin)==_norm("Москва"):
        prefixes=("МСК прайс","Мск прайс","MSK price")
    else:
        prefixes=("СПБ прайс","СПб прайс","SPB price")
    hosts=("https://cts-group.ru","https://www.cts-group.ru","http://cts-group.ru")
    out=[]
    for d in seen:
        ds=d.strftime("%d.%m.%y")
        for prefix in prefixes:
            path="/doc/"+quote(f"{prefix} {ds}.xlsx",safe="")
            for host in hosts:
                u=host+path
                if u not in out: out.append(u)
    return out

def _collect_cts_from_xlsx(origin: str, destination: str) -> tuple[list[str],list[str],str,str]:
    page_urls=["https://cts-group.ru/prices","https://www.cts-group.ru/prices","http://cts-group.ru/prices"]
    candidates=[]; discovery_errors=[]
    for page_url in page_urls:
        try:
            raw,transport=_download_bytes(page_url,referer=page_url.rsplit('/',1)[0]+'/',timeout=28)
            html=raw.decode("utf-8","replace")
            candidates.extend(_cts_xlsx_candidates(html,page_url,origin))
            if candidates: break
        except Exception as exc:
            discovery_errors.append(f"{page_url}: {exc}")
    candidates += _cts_recent_xlsx_candidates(origin)
    # de-duplicate while preserving newest / page-derived candidates first
    candidates=list(dict.fromkeys(candidates))
    errors=[]
    for url in candidates:
        try:
            data,tr=_download_bytes(url,referer="https://cts-group.ru/prices",timeout=25)
            if not data.startswith(b"PK\x03\x04"): raise RuntimeError("ответ не похож на XLSX")
            headers,cells=_cts_parse_xlsx_bytes(data,url,destination)
            return headers,cells,url,f"{tr}+xlsx"
        except Exception as e:
            # Keep the error summary compact; there can be many date candidates.
            if len(errors)<8: errors.append(f"{unquote(url)}: {type(e).__name__}: {e}")
    detail=" | ".join(errors)
    if discovery_errors: detail=("discovery: "+" | ".join(discovery_errors[:2])+"; "+detail)
    raise RuntimeError("CTS XLSX candidates failed: "+detail)

def _cts_values_from_cells(cells: list[str]) -> tuple[dict[str,dict[str,Any]], float]:
    if len(cells)<14:
        raise RuntimeError(f"CTS: ожидалось >=14 колонок, получено {len(cells)}")
    minimum=_num(cells[2])
    if minimum is None:
        raise RuntimeError("CTS: не найден минимум")
    def rate(w:float):
        idx=13 if w<=100 else 12 if w<=200 else 11 if w<=300 else 10 if w<=500 else 9 if w<=800 else 8 if w<=1000 else 7 if w<=1200 else 6 if w<=1500 else 5 if w<=2000 else 4 if w<=3000 else 3
        return _num(cells[idx])
    vals={}
    for p in COMMON_PROFILES:
        if p.get("is_minimum_profile"):
            vals[p["id"]]={"kind":"exact","price":round(minimum,2)}
            continue
        r=rate(float(p["weight_kg"]))
        if r is not None and r>0:
            vals[p["id"]]={"kind":"exact","price":round(max(minimum,float(p["weight_kg"])*r),2),"rate_per_kg":r,"minimum":minimum}
    if len(vals)<20:
        raise RuntimeError(f"CTS: слишком мало валидных диапазонов: {len(vals)}")
    return vals,float(minimum)

def _cts_origin_url(raw, page_url, origin):
    soup=BeautifulSoup(raw,'lxml')
    form=soup.select_one('form#prices')
    if form is None or str(form.get('method','get')).lower()!='get':
        raise RuntimeError('CTSGroup: форма прайса не распознана')
    control=form.select_one('#price_from')
    field=control.select_one('input[name]') if control else None
    choice=next((n for n in control.select('[data-value]') if _norm(n.get_text(' ',strip=True))==_norm(origin)),None) if control else None
    if field is None or choice is None or not choice.get('data-value'):
        raise UnpublishedTariff('CTSGroup: город отправления не найден в справочнике сайта')
    target=urljoin(page_url,form.get('action') or page_url)
    if urlsplit(target).netloc!=urlsplit(page_url).netloc or urlsplit(target).path.rstrip('/')!='/prices':
        raise RuntimeError('CTSGroup: форма ведёт на другой источник')
    return target+'?'+urlencode({field['name']:choice['data-value']})


def collect_cts(origin="Санкт-Петербург",destination="Москва") -> dict[str,Any]:
    from .online_tariffs import fetch
    url='https://cts-group.ru/prices'
    raw,meta=fetch('CTSgroup',origin,destination,url)
    try:headers,cells,headings=_find_route_table(BeautifulSoup(raw,'lxml'),origin,destination)
    except RuntimeError:
        # Resolve the current origin ID from the actual form; never guess a date,
        # substitute a sibling workbook, or take Moscow's default table for SPB.
        selected=_cts_origin_url(raw,meta['source_url'],origin)
        raw,meta=fetch('CTSgroup',origin,destination,selected,referer=url)
        headers,cells,headings=_find_route_table(BeautifulSoup(raw,'lxml'),origin,destination)
    vals,_=_cts_values_from_cells(cells)
    meta['source_type']='Официальная таблица CTSGroup'
    save_live_update('CTSgroup',origin,destination,vals,meta)
    return {'company':'CTSgroup','ok':True,'rows':len(vals),'message':f'CTSGroup: загружено {len(vals)} диапазонов'}


def _newline_rates(headers:list[str],cells:list[str]):
    # Common official layout: destination, days, minimum, <=250, <=750, <=1250, <=2500, <=5000.
    minimum=_num(cells[2]) if len(cells)>=3 else None
    nums=[_num(x) for x in cells[3:8]] if len(cells)>=8 else []
    if minimum is not None and len(nums)==5 and all(x is not None for x in nums):
        return minimum,list(zip([250.,750.,1250.,2500.,5000.],[float(x) for x in nums]))
    # Header-based fallback.
    hs=[_norm(x) for x in headers]
    rates=[]
    for i,h in enumerate(hs):
        if i>=len(cells):break
        if minimum is None and "мин" in h: minimum=_num(cells[i])
        m=re.search(r"(?:до|<=?)\s*(250|750|1250|2500|5000)",h)
        if m:
            v=_num(cells[i]);
            if v is not None: rates.append((float(m.group(1)),v))
    return minimum,sorted(rates)


def _newline_pdf_values(pdf_bytes: bytes, destination: str) -> tuple[float,list[tuple[float,float]],str]:
    from .tariff_documents import pdf_reader
    from .legacy_backend import _newline_parse_layout_text
    reader=pdf_reader(pdf_bytes)
    routes={}; page_texts=[]
    for page_no,page in enumerate(reader.pages[:80],start=1):
        try:
            try: text=page.extract_text(extraction_mode="layout") or ""
            except TypeError: text=page.extract_text() or ""
        except Exception: continue
        page_texts.append((page_no,text))
        routes.update(_newline_parse_layout_text(text,source_file="NL002-live.pdf",page_no=page_no))

    wanted=_norm(destination)
    route_names=[name for name in routes if _norm(name)==wanted]
    if wanted=="москва":
        route_names += [name for name in routes if _norm(name).startswith("москва-") or _norm(name).startswith("москва ")]
    route_names=list(dict.fromkeys(route_names))

    # Fallback for PDF editions where column spacing/layout extraction changed:
    # find a Moscow row and read RUB tokens directly from the same physical line.
    if not route_names and wanted=="москва":
        rub_re=re.compile(r"(?<!\d)(\d[\d \u00a0]*(?:[,.]\d+)?)\s*руб\.?",re.I)
        for page_no,text in page_texts:
            for raw in text.splitlines():
                if "москва" not in _norm(raw): continue
                vals=[_num(x) for x in rub_re.findall(raw)]
                vals=[float(x) for x in vals if x is not None and x>0]
                if len(vals)>=4:
                    city_match=re.search(r"(Москва(?:[-–— ][А-Яа-яЁёA-Za-z]+)?)",raw,re.I)
                    name=(city_match.group(1).strip() if city_match else "Москва")
                    minimum=vals[0]; limits=[250.,750.,1250.,2500.,5000.,10000.]
                    candidates=[{"mode":"minimum","value":minimum,"low":None,"high":None,"file":"NL002-live.pdf","page":page_no}]
                    for lim,rate in zip(limits,vals[1:]):
                        candidates.append({"mode":"per_kg","value":rate,"low":0,"high":lim,"file":"NL002-live.pdf","page":page_no})
                    routes[name]=candidates; route_names=[name]; break
            if route_names: break

    schedules=[]
    for route_name in route_names:
        candidates=routes.get(route_name) or []
        minimum=next((float(x['value']) for x in candidates if x.get('mode')=='minimum' and x.get('value') is not None),None)
        rates=[]
        for x in candidates:
            if x.get('mode')!='per_kg' or x.get('value') is None: continue
            high=x.get('high')
            # A printed open-ended final column must not be dropped. Missing
            # finite columns remain absent; no rate is extended by guesswork.
            rates.append((float(high) if high is not None else float('inf'),float(x['value'])))
        rates=sorted(rates)
        if minimum is not None and len(rates)>=3:
            # Choose one published Moscow terminal consistently for the whole
            # grid. Lowest validated 100kg terminal tariff is deterministic and
            # avoids mixing terminal rows from profile to profile.
            r100=next((r for lim,r in rates if 100<=lim+1e-9),None)
            total100=max(minimum,100*r100) if r100 is not None else float('inf')
            schedules.append((total100,route_name,minimum,rates))
    if not schedules:
        found=", ".join(sorted(routes)[:12]) or "нет распознанных направлений"
        raise UnpublishedTariff(f"Новая Линия: в PDF нет строки {destination}; примеры распознанных направлений: {found}")
    _,route_name,minimum,rates=min(schedules,key=lambda x:x[0])
    return minimum,rates,route_name

def _confirm_pdf_origin(pdf, origin, company):
    from .tariff_documents import pdf_reader
    reader=pdf_reader(pdf)
    header=' '.join((reader.pages[0].extract_text() or '').split())
    label='Междугородняя доставка' if company=='Новая Линия' else 'Тарифы на перевозку груза'
    pattern=re.compile(re.escape(label)+r'\s+из\s+г\.\s*'+re.escape(origin)+r'(?=\s|$)',re.I)
    if not pattern.search(header):
        raise ValueError(f'{company}: заголовок PDF не подтверждает отправление из {origin}')
    if company=='КИТ' and 'Цены указаны с учетом НДС в российских рублях' not in header:
        raise ValueError('КИТ: валюта прайса не подтверждена')


def _kit_pdf_values(path,origin,destination):
    from . import legacy_backend as lb
    _confirm_pdf_origin(path.read_bytes(),origin,'КИТ')
    index=lb._kit_route_index(path,origin)
    candidates=(index.get('routes') or {}).get(normalize_city(destination),[])
    if not candidates:raise UnpublishedTariff('КИТ: строка назначения отсутствует в свежем PDF')
    minimum=lb._select_strict_candidate(candidates,1,'minimum')
    minimum=float(minimum['value']) if minimum else None
    values={}
    for profile in COMMON_PROFILES:
        weight=profile['weight_kg']
        mode='minimum' if profile.get('is_minimum_profile') else ('fixed' if weight<=35 else 'per_kg')
        item=lb._select_strict_candidate(candidates,weight,mode)
        if not item:continue
        value=float(item['value'])
        total=max(minimum or 0,weight*value) if mode=='per_kg' else value
        values[profile['id']]={'kind':'exact','price':round(total,2),'source_page':item.get('page')}
        if mode=='per_kg':values[profile['id']].update(rate_per_kg=value,minimum=minimum)
    return values


def collect_kit(origin,destination):
    from .online_tariffs import fetch
    ident=carrier_catalogs.kit_origin(origin)
    url='https://tk-kit.ru/rates-new/get-pdf-new?'+urlencode({
        'id':'none','name':'Тарифы из гор. '+origin,'transport_out_code':ident,
        'transport_in_code':'','currency':'RUB','typeF':'1'})
    raw,meta=fetch('КИТ',origin,destination,url,'PDF','https://tk-kit.ru/rates-new')
    from .legacy import DOWNLOAD_DIR
    terminal=carrier_catalogs.kit_pdf_origin(origin,raw)
    values=_kit_pdf_values(DOWNLOAD_DIR/meta['source_file'],terminal,destination)
    meta['origin_terminal']=terminal if terminal!=origin else ''
    return values,meta


def collect_newline(origin="Санкт-Петербург",destination="Москва") -> dict[str,Any]:
    from_id,origin_terminal=carrier_catalogs.newline_origin(origin)
    source_id="NL-"+route_slug(origin,destination)
    url=f"https://tknl.ru/price/dev.php?AJAX=N&RESULT=Y&PDF=Y&SERVICE=DELIVERY&FROM={from_id}&TO%5B%5D=ALL"
    dl=_bounded_source_download({"id":source_id,"company":"Новая Линия","block_title":f"{origin}-{destination}","url":url,"reference_url":"https://tknl.ru/price/","document_format":"PDF"})
    if dl.get("status")!="downloaded":
        raise RuntimeError(f"Новая Линия: {dl.get('error') or 'PDF не загружен'}")
    pdf=Path(str(dl["path"])).read_bytes(); transport=str(dl.get("connector") or "v42")
    if not pdf.startswith(b"%PDF"):
        raise RuntimeError(f"Новая Линия: endpoint вернул не PDF ({len(pdf)} bytes, {pdf[:40]!r})")
    _confirm_pdf_origin(pdf,origin_terminal,"Новая Линия")
    minimum,rates,route_name=_newline_pdf_values(pdf,destination)
    vals={"min":{"kind":"exact","price":round(minimum,2)}}
    for p in COMMON_PROFILES:
        if p.get("is_minimum_profile"): continue
        w=float(p["weight_kg"]); r=next((r for lim,r in rates if w<=lim+1e-9),None)
        if r is not None:
            vals[p["id"]]={"kind":"exact","price":round(max(minimum,w*r),2),"rate_per_kg":r,"minimum":minimum}
    if len(vals)<12: raise RuntimeError(f"Новая Линия: слишком мало валидных диапазонов из PDF: {len(vals)}")
    missing=[p['id'] for p in COMMON_PROFILES if not p.get('is_minimum_profile') and p['id'] not in vals]
    extra_meta={'unavailable_profiles':{},'missing_profile_errors':{}}
    if missing:
        try:
            from .newline_calculator import collect_missing
            extra,unavailable,proof=collect_missing(origin,destination,origin_terminal,route_name,missing)
            vals.update(extra)
            extra_meta.update(unavailable_profiles=unavailable,unavailable_evidence=proof)
        except Exception as exc:
            extra_meta['missing_profile_errors']={pid:str(exc) for pid in missing}
    now=datetime.now().astimezone().isoformat(timespec="seconds")
    save_live_update("Новая Линия",origin,destination,vals,{"source_type":f"Официальный PDF Новая Линия — live {source_id}","source_url":url,"captured_at":now,"transport":transport,"destination_variant":route_name,"origin_terminal":origin_terminal,"origin":origin,"destination":destination,"source_file":dl["file"],"sha256":dl["sha256"],**extra_meta})
    return {"company":"Новая Линия","ok":True,"rows":len(vals),"message":f"Новая Линия: загружено {len(vals)} диапазонов из {source_id}, строка {route_name} ({transport})"}


def _dellin_values_from_fresh_pdf(path: Path, origin: str, destination: str) -> tuple[dict[str,dict[str,Any]],dict[str,Any]]:
    """Parse only the PDF downloaded in this v42 attempt; never a bundled fallback."""
    from . import legacy_backend as lb
    tariffs=lb.parse_dellin_tariffs(path,origin)
    tariff=tariffs.get(normalize_city(destination))
    if not tariff:
        raise RuntimeError("ДЛ: направление отсутствует в свежем официальном PDF")
    vals={}
    for p in COMMON_PROFILES:
        w=float(p["weight_kg"])
        if p.get("is_minimum_profile"):
            price=float(tariff["minimum"])
            vals[p["id"]]={"kind":"exact","price":round(price,2),"minimum":price}
            continue
        if w>max(high for _,high in lb.KG_TIERS):
            continue
        if w<=35:
            idx=0 if w<=1 else 1 if w<=3 else 2 if w<=5 else 3 if w<=15 else 4
            price=float(tariff["fixed"][idx])
            vals[p["id"]]={"kind":"exact","price":round(price,2)}
            continue
        kg_idx=lb.select_tier(w,lb.KG_TIERS)
        if kg_idx is None:
            continue
        rate=float(tariff["kg_rates"][kg_idx]); minimum=float(tariff["minimum"])
        price=max(minimum,w*rate)
        vals[p["id"]]={"kind":"exact","price":round(price,2),"rate_per_kg":rate,"minimum":minimum}
    if not vals:
        raise RuntimeError("ДЛ: свежий PDF получен, но тарифные строки не распознаны")
    return vals,{"source_type":"Официальный PDF ДЛ — live","source_url":"https://www.dellin.ru/pricelist_pdf/",
                 "source_file":path.name,"captured_at":datetime.now().astimezone().isoformat(timespec="seconds")}


def _profile_values_from_fresh_documents(company: str, origin: str, destination: str) -> tuple[dict[str,dict[str,Any]],dict[str,Any]]:
    """Build LIVE rows exclusively from files downloaded in the current batch.

    Legacy parsers may know about bundled official snapshots as a LAST GOOD
    fallback. v42 explicitly rejects those here: a parser result is eligible for
    LIVE only if its source_file is one of the files freshly downloaded for this
    route/source id in this very refresh batch.
    """
    from . import legacy as core
    from . import legacy_backend as lb
    source_ids=lb._source_ids_for_origin(company,origin) or set()
    result=core.load_results()
    statuses=[x for x in (result.get("sources") or []) if str(x.get("id")) in set(source_ids)]
    fresh=[x for x in statuses if x.get("status") in {"parsed","downloaded"} and str(x.get("freshness") or "").lower() not in {"snapshot","cached","fallback"}]
    if not fresh:
        warning=next((x.get("refresh_warning") or x.get("message") for x in statuses if x.get("refresh_warning") or x.get("message")),None)
        raise RuntimeError(warning or "официальный источник не дал свежего распознанного ответа")

    fresh_names=set()
    fresh_paths=[]
    for st in fresh:
        fresh_names.update(Path(str(name)).name for name in (st.get("files") or []) if name)
    for fm in (result.get("files") or []):
        if str(fm.get("source_id") or "") not in source_ids:
            continue
        if str(fm.get("freshness") or "").lower() in {"snapshot","cached","fallback"}:
            continue
        if fm.get("file"):
            fresh_names.add(Path(str(fm["file"])).name)
        raw=str(fm.get("path") or "")
        path=Path(raw) if raw else core.DOWNLOAD_DIR/str(fm.get("file") or "")
        if path.exists():
            fresh_paths.append(path)
    if not fresh_names:
        raise RuntimeError("свежий источник отмечен, но его скачанный файл не зарегистрирован")

    if company=="ДЛ":
        pdf=next((p for p in fresh_paths if p.suffix.lower()==".pdf"),None)
        if not pdf:
            raise RuntimeError("ДЛ: свежий PDF не найден в текущей попытке")
        return _dellin_values_from_fresh_pdf(pdf,origin,destination)

    vals: dict[str,dict[str,Any]]={}
    source_meta: dict[str,Any]={}
    for p in COMMON_PROFILES:
        try:
            row=lb.collected_published_tariff(company,origin,destination,float(p["weight_kg"]),minimum=bool(p.get("is_minimum_profile")))
        except Exception:
            row=None
        if not row or row.get("status") not in {"ok","cached"} or not isinstance(row.get("price"),(int,float)):
            continue
        # Critical anti-fallback guard: a bundled snapshot can remain LAST GOOD,
        # but it cannot be promoted to LIVE just because a fresh page failed to parse.
        source_file=Path(str(row.get("source_file") or "")).name
        if not source_file or source_file not in fresh_names:
            continue
        freshness=str(row.get("freshness") or "").lower()
        if freshness not in {"collected_official_file","official_pdf","collected_official_pdf"}:
            continue
        kind="lower_bound" if row.get("price_is_minimum") and not p.get("is_minimum_profile") else "exact"
        vals[p["id"]]={"kind":kind,"price":round(float(row["price"]),2)}
        if isinstance(row.get("rate_per_kg"),(int,float)): vals[p["id"]]["rate_per_kg"]=float(row["rate_per_kg"])
        if isinstance(row.get("minimum_charge"),(int,float)): vals[p["id"]]["minimum"]=float(row["minimum_charge"])
        if not source_meta:
            source_meta={
                "source_type":row.get("source_type") or f"Официальный источник {company} — live",
                "source_url":row.get("source_url"),
                "source_file":source_file,
            }
    if not vals:
        raise RuntimeError("свежий файл скачан, но ни один диапазон не был строго распознан именно из этого файла; LAST GOOD не повышен до LIVE")
    source_meta["captured_at"]=datetime.now().astimezone().isoformat(timespec="seconds")
    evidence=next((f for f in result.get('files',[]) if f.get('file')==source_meta.get('source_file')),{})
    source_meta.update({'origin':origin,'destination':destination,'transport':evidence.get('connector'),
                        'source_url':evidence.get('final_url') or evidence.get('url') or source_meta.get('source_url'),
                        'sha256':evidence.get('sha256')})
    return vals,source_meta




def _route_page_url(company: str, origin: str, destination: str) -> str:
    pair=(normalize_city(origin),normalize_city(destination))
    urls={
        ("ДЛ","Санкт-Петербург","Москва"): "https://www.dellin.ru/directions/gruzoperevozki-po-rossii/sankt-peterburg-moskva/",
        ("ДЛ","Москва","Санкт-Петербург"): "https://www.dellin.ru/directions/gruzoperevozki-po-rossii/moskva-sankt-peterburg/",
        ("Байкал Сервис","Санкт-Петербург","Москва"): "https://www.baikalsr.ru/city/spb__moscow/",
        ("Байкал Сервис","Москва","Санкт-Петербург"): "https://www.baikalsr.ru/city/moscow__spb/",
    }
    try:return urls[(company,*pair)]
    except KeyError:
        if company=="Байкал Сервис":return carrier_catalogs.baikal_route_url(*pair)
        if company=="ДЛ":return "https://www.dellin.ru/directions/gruzoperevozki-po-rossii/"+city_slug(pair[0])+"-"+city_slug(pair[1])+"/"
        raise ValueError(company)


def _city_pattern(city):
    return city_pattern(city)


def _route_soup(raw, final_url, expected_url, origin, destination):
    expected=urlsplit(expected_url); actual=urlsplit(final_url)
    if (actual.hostname or '').removeprefix('www.') != (expected.hostname or '').removeprefix('www.') or unquote(actual.path).rstrip('/')!=unquote(expected.path).rstrip('/'):
        raise RuntimeError('Источник перенаправил запрос: выбранный маршрут не подтверждён')
    # Let the HTML charset/BOM be decoded before matching Russian city names.
    soup=BeautifulSoup(raw,'lxml')
    headings=[node.get_text(' ',strip=True) for node in soup.find_all(['h1'])]
    if soup.title:headings.append(soup.title.get_text(' ',strip=True))
    route_re=re.compile(r'\b'+_city_pattern(origin)+r'\b.{0,100}?\b'+_city_pattern(destination)+r'\b',re.I)
    if not any(route_re.search(_norm(h)) for h in headings):
        raise RuntimeError(f'Источник ответил, но заголовок не подтверждает маршрут {origin} → {destination}')
    return soup


def _live_route_minimum(company: str, origin: str, destination: str) -> dict[str,Any]:
    """Read a CURRENT lower-bound directly from the carrier's official route page.

    This is intentionally network-only.  No bundled JSON value is accepted here.
    When the exact public calculator is unavailable, a fresh official ``от`` value
    is still useful evidence that the button really reached the carrier today.
    """
    url=_route_page_url(company,origin,destination)
    raw,transport,final_url=_download_html_bytes(url,referer=url,timeout=50)
    if not raw: raise RuntimeError(f"{company}: официальный маршрутный URL вернул пустой ответ")
    soup=_route_soup(raw,final_url,url,origin,destination)
    text=re.sub(r"\s+"," ",soup.get_text(" ",strip=True))
    o=normalize_city(origin); d=normalize_city(destination)
    nd=_norm(d)
    price=None
    if company=="Байкал Сервис":
        # Both inbound ("из Москва") and outbound ("в Москва") cards exist.
        # Only the latter belongs to this page's origin and requested destination.
        pattern=re.compile(r'^в\s+'+_city_pattern(d)+r'\s+от\s*([\d\s\u00a0]+(?:[,.]\d+)?)\s*(?:руб\.?|₽)',re.I)
        for node in soup.find_all('a'):
            match=pattern.search(re.sub(r'\s+',' ',node.get_text(' ',strip=True)))
            if match:
                price=_num(match.group(1))
                if price:break
    elif company=="ДЛ":
        # Dellin route pages explicitly state/publish minimum route prices.
        # Exact route pages normally put the lower bound in title/H1; if that is
        # omitted, their "Популярные направления из <origin>" block is a valid
        # same-origin fallback. We never use a hardcoded number.
        route_zone=" ".join(filter(None,[soup.title.get_text(" ",strip=True) if soup.title else "", (soup.find("h1").get_text(" ",strip=True) if soup.find("h1") else "")]))
        m=re.search(r"(?:стоимость\s+)?от\s*([\d\s\u00a0]+(?:[,.]\d+)?)\s*(?:р\.?|руб\.?|₽)",route_zone,re.I)
        if m: price=_num(m.group(1))
        if price is None:
            # Find the heading and only inspect a bounded following text window.
            lower=text.lower().replace("ё","е")
            marker=re.search(r'популярные направления из\s+'+_city_pattern(o),lower,re.I)
            pos=marker.start() if marker else -1
            zone=text[pos:pos+2400] if pos>=0 else ''
            m=re.search(rf"{re.escape(d)}\s+([\d\s\u00a0]+(?:[,.]\d+)?)\s*₽",zone,re.I)
            if m: price=_num(m.group(1))
    if price is None or price<=0:
        raise RuntimeError(f"{company}: свежая маршрутная страница получена, но цена «от» не распознана")
    now=datetime.now().astimezone().isoformat(timespec="seconds")
    return {
        "status":"ok","price":round(float(price),2),"price_is_minimum":True,
        "source_type":f"Официальная маршрутная страница {company} — LIVE",
        "source_url":final_url,"freshness":"live","browser":transport,"captured_at":now,
        "message":"Свежая официальная цена «от» получена сетевым запросом в текущей попытке; это не точная ставка выбранного веса.",
    }

def _vozovoz_route_url(origin: str, destination: str) -> str:
    pair=(normalize_city(origin),normalize_city(destination))
    if pair==("Санкт-Петербург","Москва"):
        return "https://vozovoz.ru/order/create/sankt-peterburg_moskva/"
    if pair==("Москва","Санкт-Петербург"):
        return "https://vozovoz.ru/order/create/moskva_sankt-peterburg/"
    return "https://vozovoz.ru/order/create/"+city_slug(pair[0])+"_"+city_slug(pair[1])+"/"


def _vozovoz_live_route_minimum(origin: str, destination: str) -> dict[str,Any]:
    url=_vozovoz_route_url(origin,destination)
    raw,transport,final_url=_download_html_bytes(url,referer="https://vozovoz.ru/",timeout=50)
    if not raw:
        raise RuntimeError("Возовоз: маршрутная страница не загружена")
    soup=_route_soup(raw,final_url,url,origin,destination)
    text=re.sub(r"\s+"," ",soup.get_text(" ",strip=True))
    # Parse only the explicit route-cost sentence, not unrelated popular-route cards.
    m=re.search(r'Цена\s+за\s+услугу\s+[«"]?Перевозка\s+между\s+городами[»"]?.{0,220}?от\s+([\d\s\u00a0]+)\s*₽',text,re.I)
    if not m:
        m=re.search(r'Стоимость\s+перевозки.{0,240}?от\s+([\d\s\u00a0]+)\s*₽',text,re.I)
    if not m:
        raise RuntimeError("Возовоз: свежая страница получена, но маршрутный минимум не распознан")
    price=float(re.sub(r"\s+","",m.group(1).replace("\u00a0"," ")))
    return {"status":"ok","price":price,"price_is_minimum":True,"source_type":"Официальная маршрутная страница Возовоза — LIVE","source_url":final_url,"freshness":"live","browser":transport,"captured_at":datetime.now().astimezone().isoformat(timespec="seconds"),"message":"Свежая официальная цена «от»; точный расчёт публичного калькулятора недоступен."}

def _dellin_api_live(appkey: str, origin: str, destination: str, weight: float, volume: float) -> dict[str,Any]:
    """Calculate terminal-to-terminal cost through Dellin's official public API.

    The application key is user-supplied and stored only in runtime/settings.json.
    Terminal ids are resolved fresh from the official terminal-search method so
    they are not hardcoded into the build.
    """
    from . import legacy_backend as lb
    key=str(appkey or '').strip()
    if not key:
        raise RuntimeError('ДЛ API: appkey не настроен')
    o=normalize_city(origin); d=normalize_city(destination)
    terminal_url='https://api.dellin.ru/v1/public/request_terminals.json'
    calc_url='https://api.dellin.ru/v2/calculator.json'
    dims={
        'length':max(0.01,float(volume))**(1/3),
        'width':max(0.01,float(volume))**(1/3),
        'height':max(0.01,float(volume))**(1/3),
        'weight':float(weight),'maxVolume':float(volume),
        'totalVolume':float(volume),'totalWeight':float(weight),
    }
    headers={'Accept':'application/json','Content-Type':'application/json','User-Agent':'TariffComparison/44.0'}
    def resolve(city:str,direction:str)->int:
        payload={'appkey':key,'search':city,'direction':direction,'maxCargoDimensions':dims}
        rr=requests.post(terminal_url,json=payload,headers=headers,timeout=(6,18)); rr.raise_for_status()
        data=rr.json() or {}; terms=data.get('terminals') or []
        if isinstance(terms,dict): terms=terms.get('terminal') or terms.get('terminals') or []
        if not isinstance(terms,list): terms=[]
        exact=[x for x in terms if isinstance(x,dict) and _norm(x.get('city'))==_norm(city)]
        pool=exact
        chosen=next((x for x in pool if x.get('default') is True or str(x.get('default')).lower() in {'1','true','yes'}),None)
        chosen=chosen or (pool[0] if pool else None)
        if not chosen or chosen.get('id') in (None,''):
            raise RuntimeError(f'ДЛ API: терминал для {city} не найден')
        return int(chosen['id'])
    dep=resolve(o,'derival'); arr=resolve(d,'arrival')
    side=max(0.01,float(volume))**(1/3)
    produce=(datetime.now().astimezone()+timedelta(days=1)).date().isoformat()
    payload={
        'appkey':key,
        'delivery':{
            'deliveryType':{'type':'auto'},
            'derival':{'produceDate':produce,'variant':'terminal','terminalID':dep},
            'arrival':{'variant':'terminal','terminalID':arr},
        },
        'cargo':{
            'quantity':1,'length':round(side,4),'width':round(side,4),'height':round(side,4),
            'weight':float(weight),'totalWeight':float(weight),'totalVolume':float(volume),
        },
    }
    rr=requests.post(calc_url,json=payload,headers=headers,timeout=(6,22)); rr.raise_for_status()
    data=rr.json() or {}
    if not isinstance(data,dict) or data.get('errors') or data.get('error') or (isinstance(data.get('metadata'),dict) and data['metadata'].get('status',200)>=400):
        raise RuntimeError('ДЛ API: расчёт отклонён; проверьте appkey и параметры груза')
    from .online_tariffs import number
    total=data.get('data',{}).get('price') if isinstance(data.get('data'),dict) else data.get('price')
    try:price=number(total)
    except ValueError:
        msg=data.get('errors') or data.get('error') or data.get('message') or 'ответ без итоговой стоимости'
        raise RuntimeError(f'ДЛ API: {msg}')
    return {
        'status':'ok','price':round(float(price),2),'price_is_minimum':False,
        'source_type':'Официальный API Деловых Линий — LIVE',
        'source_url':calc_url,'freshness':'live','browser':'API',
        'message':f'LIVE через официальный API; терминалы получены онлайн для {o} → {d}.',
    }


def _collect_live_calculator(company: str, origin: str, destination: str, profile_id: str, timeout: float=50) -> tuple[dict[str,dict[str,Any]],dict[str,Any]]:
    """Get a fresh exact point from an official calculator/API without cache fallback."""
    from . import legacy_backend as lb
    p=PROFILE_BY_ID[profile_id]
    if p.get("is_minimum_profile"):
        raise RuntimeError('Калькулятор не публикует минимум. Выберите конкретный вес для онлайн-расчёта.')
    w=float(p["weight_kg"])
    v=float(lb.matrix_control_volume(w))
    row=None
    if company in {"Werner","Главтрасса"}:
        # Both carriers expose the same documented keyless api_calc shape for
        # Moscow/SPB. Use a bounded v42 request instead of the legacy multi-retry path.
        ids={"Москва":"35","Санкт-Петербург":"36"}; o=normalize_city(origin); d=normalize_city(destination)
        dep,arr=carrier_catalogs.keyless_city_ids(company,o,d) if o not in ids or d not in ids else (ids[o],ids[d])
        if not dep or not arr: raise RuntimeError("город не найден в публичном API-справочнике")
        side=max(0.01,v)**(1/3)
        params=[("method","api_calc"),("responseFormat","json"),("depPoint",dep),("arrPoint",arr),
                ("cargoMest[1]","1"),("cargoKg[1]",f"{w:g}"),("cargoL[1]",f"{side:.4f}"),
                ("cargoW[1]",f"{side:.4f}"),("cargoH[1]",f"{side:.4f}"),("cargoCalculation[1]","1")]
        endpoint="https://wernerus.ru/api/calc/" if company=="Werner" else "https://glavtrassa.ru/api/calc/"
        referer="https://wernerus.ru/clients/prices/" if company=="Werner" else "https://glavtrassa.ru/clients/calc/"
        target=endpoint+'?'+urlencode(params)
        response=download(target,referer=referer,expected='JSON',timeout=timeout,preferred_transport='curl-cffi-tls12')
        data=json.loads(response.content)
        price=data.get('price') if isinstance(data,dict) else None
        if not isinstance(price,(int,float)) or isinstance(price,bool) or price<=0 or data.get('error'):
            raise RuntimeError('API вернул ответ без подтверждённой итоговой стоимости')
        row={"status":"ok","price":round(float(price),2),"source_type":f"Официальный публичный API {company}","source_url":target,"freshness":"live","browser":response.transport}
    elif company=="ДЛ":
        # Prefer the official API when the user supplied an appkey.  Without a
        # key we use the public widget; if its JS/WAF blocks automation, re-read
        # the CURRENT official route page and expose only its fresh lower bound.
        settings=lb.load_settings(); errors=[]
        if settings.get("dellin_appkey"):
            try:
                row=_dellin_api_live(str(settings.get("dellin_appkey")),origin,destination,w,v)
            except Exception as exc:
                errors.append(f"API: {exc}")
        if row is None and not _browser_enabled():
            row=_live_route_minimum("ДЛ",origin,destination)
        if row is None:
            from .public_web import calculate_from_public_site
            calc=calculate_from_public_site("ДЛ",origin,destination,w,v,allow_visible_fallback=False)
            if calc.get("ok"):
                row={**calc,"status":"ok","source_type":calc.get("source_label") or "Официальный калькулятор Деловых Линий","freshness":"live"}
            else:
                errors.append("виджет: "+str(calc.get("message") or "без детализации"))
                row=_live_route_minimum("ДЛ",origin,destination)
                row["message"] = str(row.get("message") or "") + " Точный расчёт недоступен: " + " | ".join(errors[-2:])
    elif company=="Возовоз":
        settings=lb.load_settings()
        if settings.get("vozovoz_api_key"):
            row=lb.calculate_vozovoz(origin,destination,w,v)
            if str(row.get("freshness") or "").lower() != "live":
                raise RuntimeError("Возовоз не вернул production LIVE-ответ")
        elif not _browser_enabled():
            row=_vozovoz_live_route_minimum(origin,destination)
        else:
            # No production API key: use the official route-specific public
            # calculator. The published demo API is deliberately NOT accepted as LIVE.
            from .public_web import calculate_from_public_site
            route_url=_vozovoz_route_url(origin,destination)
            calc=calculate_from_public_site("Возовоз",origin,destination,w,v,allow_visible_fallback=False,url_override=route_url,route_preselected=True)
            if calc.get("ok"):
                row={**calc,"status":"ok","source_type":calc.get("source_label") or "Официальный публичный калькулятор Возовоза","freshness":"live"}
            else:
                row=_vozovoz_live_route_minimum(origin,destination)
    elif company=="Байкал Сервис":
        settings=lb.load_settings()
        if settings.get("baikal_api_key") and settings.get("baikal_api_url"):
            row=lb.calculate_baikal(origin,destination,w,v)
            if str(row.get("freshness") or "").lower()!="live":
                raise RuntimeError("Байкал Сервис API не вернул production LIVE-ответ")
        elif not _browser_enabled():
            row=_live_route_minimum(company,origin,destination)
        else:
            from .public_web import calculate_from_public_site
            errors=[]
            # Route-specific page first: it contains both the calculator and the
            # correct origin/destination context. Then try the generic calculator.
            live_urls=(_route_page_url(company,origin,destination),
                       )
            for live_url in live_urls:
                calc=calculate_from_public_site(company,origin,destination,w,v,allow_visible_fallback=False,url_override=live_url,route_preselected=(live_url==live_urls[0]))
                if calc.get("ok"):
                    row={**calc,"status":"ok","source_type":calc.get("source_label") or "Официальный публичный калькулятор Байкал Сервис","freshness":"live"}
                    break
                errors.append(str(calc.get("message") or live_url))
            if row is None:
                row=_live_route_minimum(company,origin,destination)
                row["message"] = str(row.get("message") or "") + " Точный публичный калькулятор не ответил: " + " | ".join(errors[-2:])
    else:
        raise RuntimeError("онлайн-калькулятор не настроен")
    if not row or row.get("status") not in {"ok","cached"} or not isinstance(row.get("price"),(int,float)):
        raise RuntimeError(str((row or {}).get("message") or "официальный калькулятор не вернул стоимость"))
    kind="lower_bound" if row.get("price_is_minimum") else "exact"
    vals={p["id"]:{"kind":kind,"price":round(float(row["price"]),2)}}
    if isinstance(row.get("rate_per_kg"),(int,float)): vals[p["id"]]["rate_per_kg"]=float(row["rate_per_kg"])
    meta={"source_type":row.get("source_type") or "Официальный онлайн-калькулятор",
          "source_url":row.get("source_url"),"captured_at":datetime.now().astimezone().isoformat(timespec="seconds"),
          "transport":row.get("browser") or "API","origin":origin,"destination":destination,
          "volume_m3":None if kind=='lower_bound' else v,
          "calculation_basis":('Опубликованная цена «от» для направления; точная стоимость выбранного веса не подтверждена'
                               if kind=='lower_bound' else f"Расчёт калькулятора: {w:g} кг, {v:g} м³, терминал → терминал")}
    return vals,meta


def _collect_calculator_grid(company, origin, destination):
    """Expand an exact calculator to all control weights, never a route 'from'."""
    from .network import connection_scope
    values={};errors=[];missing={};meta={}
    deadline=time.monotonic()+240
    with connection_scope():
        for p in COMMON_PROFILES:
            if p.get('is_minimum_profile'):continue
            try:
                if time.monotonic()>=deadline:raise TimeoutError('Время проверки весовой сетки истекло')
                rows,info=_collect_live_calculator(company,origin,destination,p['id'],timeout=min(30,max(1,deadline-time.monotonic())))
                if any(row.get('kind')=='lower_bound' for row in rows.values()):
                    # This page states one route minimum, not 28 exact weights.
                    if not values:
                        bound=next(iter(rows.values()))
                        values['min']={**bound,**info}
                    errors.append('Источник публикует только цену «от». Точная весовая сетка недоступна.')
                    meta=info;break
                for row in rows.values():row.update(info)
                values.update(rows);meta=info
            except Exception as exc:
                detail=str(exc);errors.append(p['id']+': '+detail);missing[p['id']]=detail
                # A missing key, unsupported route, or blocked site must not be
                # hammered 28 times. Retrying this carrier is an explicit action.
                if not values:raise
    if not values:raise ValueError('Калькулятор не подтвердил ни одной цены')
    meta={k:v for k,v in meta.items() if k not in {'volume_m3','calculation_basis','source_url','source_file','sha256'}}
    meta.update({'partial_errors':errors[:5],'missing_profile_errors':missing})
    return values,meta


def _collect_keyless_api_grid(company, origin, destination, profile_id):
    from concurrent.futures import ThreadPoolExecutor, as_completed
    from .network import connection_scope
    probe=PROFILE_BY_ID[profile_id]
    if probe.get('is_minimum_profile'):probe=PROFILE_BY_ID['w100']
    values,meta=_collect_live_calculator(company,origin,destination,probe['id'])
    for value in values.values():value.update(meta)
    deadline=time.monotonic()+75
    profiles=[p for p in COMMON_PROFILES if not p.get('is_minimum_profile') and p['id']!=probe['id']]
    errors=[]
    def batch(profiles):
        collected={};failed=[]
        with connection_scope():
            for profile in profiles:
                try:
                    if time.monotonic()>=deadline:raise TimeoutError('Истекло время проверки весовой сетки')
                    rows,row_meta=_collect_live_calculator(company,origin,destination,profile['id'],timeout=min(20,max(1,deadline-time.monotonic())))
                    for row in rows.values():row.update(row_meta)
                    collected.update(rows)
                except Exception as exc:failed.append(f"{profile['id']}: {exc}")
        return collected,failed
    with ThreadPoolExecutor(max_workers=2) as pool:
        futures=[pool.submit(batch,profiles[i::2]) for i in range(2)]
        for future in as_completed(futures):
            rows,failed=future.result();values.update(rows);errors.extend(failed)
    heavy=[v['price'] for pid,v in values.items() if PROFILE_BY_ID[pid]['weight_kg']>=100]
    if 'w20000' in values and len(heavy)>=3 and len(set(heavy))==1:
        raise ValueError('API не подтвердил тариф маршрута: одна и та же сумма для всех весов от 100 до 20 000 кг. Такой ответ требует проверки перевозчика.')
    # Metadata shared by the company must not replace per-weight request parameters.
    meta={k:v for k,v in meta.items() if k not in {'volume_m3','calculation_basis','source_url','source_file','sha256'}}
    meta.update({'grid_rows':len(values),'partial_errors':errors[:5]})
    return values,meta


def _rail_branch_url(origin):
    city='saint-petersburg' if normalize_city(origin)=='Санкт-Петербург' else city_slug(origin)
    return 'https://www.railcontinent.ru/filials/'+city+'/'


def _collect_rail_export(origin,destination):
    from .online_tariffs import fetch,parse_railcontinent_workbook
    city='С-Петербург' if normalize_city(origin)=='Санкт-Петербург' else normalize_city(origin)
    # The official export resolves its current dated file itself. Never guess a date.
    url='https://www.railcontinent.ru/ajax/calc_export.php?'+urlencode({'city':city})
    raw,meta=fetch('Рейл Континент',origin,destination,url,'XLSX',_rail_branch_url(origin),prefer_ranges=True)
    values=parse_railcontinent_workbook(raw,origin,destination)
    meta['source_type']='Официальный текущий XLSX Рейл Континент — исходящие тарифы'
    return values,meta


def _bounded_source_download(source: dict[str,Any], *, request_read_timeout: int = 12, curl_timeout: int = 16, winhttp_timeout: int = 16) -> dict[str,Any]:
    """Bounded live download for one official source.

    Unlike the legacy collector this path never falls back to a bundled snapshot:
    a network failure must remain FAILED so v42 can keep old data only as LAST GOOD.
    """
    from . import legacy as core
    url=str(source.get("url") or "")
    started=time.monotonic()
    try:
        fmt=source.get('document_format') or ''
        if source.get('prefer_ranges'):
            try:result=download_ranges(url,expected=fmt,block_size=source.get('range_chunk_bytes',8000))
            except DownloadError as partial_exc:
                if 'HTTP 401' in str(partial_exc) or 'HTTP 403' in str(partial_exc):raise
                try:result=download(url,referer=source.get('reference_url') or url,expected=fmt,timeout=20)
                except DownloadError as full_exc:raise DownloadError(f'Загрузка частями: {partial_exc}; целиком: {full_exc}') from full_exc
        else:
            try:result=download(url,referer=source.get('reference_url') or url,expected=fmt,timeout=50)
            except RejectedResponse:raise
            except DownloadError:
                if fmt.upper() not in {'PDF','XLS','XLSX'}:raise
                result=download_ranges(url,expected=fmt,block_size=source.get('range_chunk_bytes',8000))
        body=result.content; final_url=result.url; ct=result.content_type; status_code=result.status_code; transport=result.transport
    except Exception as exc:
        return {"status":"error","url":url,"duration_sec":round(time.monotonic()-started,2),"error":str(exc),"error_info":explain_error(exc)}
    ext=core.guess_extension(final_url or url,ct,body)
    fname=f"{core.slugify(str(source.get('id') or 'SRC'))}_{core.slugify(str(source.get('company') or 'company'))}_{core.slugify(str(source.get('block_title') or 'tariff'))}_{hashlib.sha256(body).hexdigest()[:12]}{ext}"
    path=core.DOWNLOAD_DIR/fname; path.parent.mkdir(parents=True,exist_ok=True); path.write_bytes(body)
    return {"status":"downloaded","status_code":status_code or 200,"url":url,"final_url":final_url,"content_type":ct,
            "file":fname,"path":str(path),"size_bytes":len(body),"extension":ext,
            "duration_sec":round(time.monotonic()-started,2),"error":"","freshness":"live","connector":f"v42_{transport}","sha256":hashlib.sha256(body).hexdigest()}


def _refresh_one_document_source(source: dict[str,Any]) -> dict[str,Any]:
    """Download/extract one official source without touching the shared registry.

    Running these workers in parallel avoids one broken carrier serially blocking
    all other companies. The registry is merged once after every worker returns.
    """
    from . import legacy as core
    status={
        "id":source.get("id"),"company":source.get("company"),"block_key":source.get("block_key"),
        "block_title":source.get("block_title"),"source_type":source.get("source_type", ""),
        "url":source.get("url", ""),"reference_url":source.get("reference_url",source.get("url","")),
        "title":source.get("title", ""),"document_format":source.get("document_format", ""),
        "started_at":datetime.now().astimezone().isoformat(timespec="seconds"),"status":"pending",
        "message":"","files":[],"rows_extracted":0,
    }
    files=[]; rows=[]; errors=[]
    try:
        dl=_bounded_source_download(source)
        status.update({
            "status":dl.get("status","unknown"),"message":dl.get("error","") or "",
            "status_code":dl.get("status_code"),"duration_sec":dl.get("duration_sec"),
            "content_type":dl.get("content_type", ""),"freshness":dl.get("freshness", ""),
            "connector":dl.get("connector",source.get("connector", "")),
        })
        if dl.get("file"):
            dl=dict(dl); dl["source_id"]=source.get("id"); files.append(dl); status["files"].append(dl["file"])
            path=Path(str(dl.get("path") or ""))
            try:
                extracted=core.extract_file(path,source); rows.extend(extracted); status["rows_extracted"]+=len(extracted)
                if extracted: status["status"]="parsed"
            except Exception as exc:
                errors.append({"source_id":source.get("id"),"company":source.get("company"),"stage":"parse","error":str(exc)})
            # v42 intentionally does not crawl secondary links here: every extra
            # download can multiply WAF/TLS delays. Route pages themselves are parsed;
            # carriers with dedicated files already point to those files in sources.json.
        status["finished_at"]=datetime.now().astimezone().isoformat(timespec="seconds")
    except Exception as exc:
        status.update({"status":"error","message":str(exc),"finished_at":datetime.now().astimezone().isoformat(timespec="seconds")})
        errors.append({"source_id":source.get("id"),"company":source.get("company"),"stage":"download","error":str(exc)})
    return {"status":status,"files":files,"rows":rows,"errors":errors}


def _merge_document_chunks(catalog, chunks):
    from . import legacy as core
    from . import legacy_backend as lb
    previous=core.load_results()
    target_ids={str(s.get("id")) for s in catalog}
    result=core.empty_results(); result["collected_at"]=datetime.now().astimezone().isoformat(timespec="seconds")
    result["sources"]=[x for x in (previous.get("sources") or []) if str((x or {}).get("id")) not in target_ids]
    result["rows"]=[x for x in (previous.get("rows") or []) if str((x or {}).get("source_id")) not in target_ids]
    result["files"]=[x for x in (previous.get("files") or []) if str((x or {}).get("source_id")) not in target_ids]
    result["errors"]=[x for x in (previous.get("errors") or []) if str((x or {}).get("source_id")) not in target_ids]
    for chunk in chunks:
        result["sources"].append(chunk["status"]); result["rows"].extend(chunk["rows"]); result["files"].extend(chunk["files"]); result["errors"].extend(chunk["errors"])
    result["summary"]=core.build_summary(result); core.save_results(result)

    # Force geometry-aware parsers to see newly downloaded files immediately.
    try: lb._clear_collected_tariff_cache()
    except Exception: pass
    for name in ("_WERNER_TABLE_CACHE","_PEK_ROUTE_FILE_CACHE","_KIT_PDF_CANDIDATE_CACHE","_NEWLINE_ROUTE_INDEX_CACHE"):
        try: getattr(lb,name).clear()
        except Exception: pass
    statuses=result.get("sources") or []
    return statuses


def _collect_document_batch(companies:list[str], origin:str, destination:str, on_company=None) -> dict[str,dict[str,Any]]:
    """Publish each completed carrier without waiting for a slow unrelated PDF."""
    if not companies:return {}
    from concurrent.futures import ThreadPoolExecutor, as_completed
    from collections import Counter
    from . import legacy as core
    from . import legacy_backend as lb
    selected_ids=set()
    for company in companies:selected_ids.update(lb._source_ids_for_origin(company,origin) or set())
    catalog=[source for source in core.load_sources().get('sources',[])
             if source.get('collect_enabled') is not False and source.get('company') in companies and str(source.get('id')) in selected_ids]
    if not catalog:raise RuntimeError('Не найден официальный источник для выбранного направления')
    pending=Counter(source['company'] for source in catalog)
    chunks={company:[] for company in companies}; output={}
    with ThreadPoolExecutor(max_workers=min(4,len(catalog))) as pool:
        futures={pool.submit(_refresh_one_document_source,source):source for source in catalog}
        for future in as_completed(futures):
            source=futures[future]; company=source['company']
            try:chunk=future.result()
            except Exception as exc:
                chunk={'status':{'id':source.get('id'),'company':company,'status':'error','message':str(exc),'files':[],'rows_extracted':0},'files':[],'rows':[],'errors':[]}
            chunks[company].append(chunk);pending[company]-=1
            if pending[company]==0:
                statuses=_merge_document_chunks([x for x in catalog if x['company']==company],chunks[company])
                output[company]={'statuses':[x for x in statuses if x.get('company')==company]}
                if on_company:on_company(company)
    return output


def collect_selected(companies:list[str], origin:str, destination:str, profile_id:str="w100", on_progress=None, *, full_grid=False, should_stop=None) -> list[dict[str,Any]]:
    """Independent carrier jobs, with immediate per-company results."""
    from concurrent.futures import ThreadPoolExecutor, wait, FIRST_COMPLETED
    from . import online_tariffs
    if profile_id not in PROFILE_BY_ID:
        raise ValueError('Неизвестный весовой профиль')
    companies=list(dict.fromkeys(companies))
    attempts={};should_stop=should_stop or (lambda:False)
    results={}
    docs=[c for c in companies if c=='ПЭК' or (c=='КИТ' and normalize_city(origin) in {'Москва','Санкт-Петербург'})]

    def finish(c, vals=None, meta=None, result=None, error=None):
        aid=attempts[c]
        try:
            if error:raise RuntimeError(error)
            if vals is not None:
                if not vals:raise ValueError('Не получено ни одного тарифного диапазона')
                save_live_update(c,origin,destination,vals,meta,attempt_id=aid)
                result={'company':c,'ok':True,'rows':len(vals),'message':f'{c}: загружено диапазонов: {len(vals)}'}
                if (meta or {}).get('partial_errors'):
                    result.update({'partial':True,'message':f'{c}: загружено {len(vals)} из 28 весов; часть запросов не завершилась'})
            rows=int((result or {}).get('rows') or 0)
            if not rows:raise ValueError('Не получено ни одной тарифной строки')
            partial_error=('Не все веса загружены: '+' | '.join(meta['partial_errors'])) if (meta or {}).get('partial_errors') else None
            finish_live_attempt(c,origin,destination,aid,rows=rows,error=partial_error)
        except Exception as exc:
            message=f'{c}: {type(exc).__name__}: {exc}'
            finish_live_attempt(c,origin,destination,aid,rows=0,error=message)
            result={'company':c,'ok':False,'rows':0,'message':message,'error_info':explain_error(message)}
        _log(f'{origin} → {destination} | {result["message"]}', console_message=(f'{origin} → {destination} | {c}: '+result['error_info']['summary']) if result.get('error_info') else None)
        if on_progress:on_progress(result)
        return result

    def document_job():
        for c in docs:attempts[c]=begin_live_attempt(c,origin,destination,profile_id)
        output={}
        def ready(company):
            try:
                vals,meta=_profile_values_from_fresh_documents(company,origin,destination)
                output[company]=finish(company,vals,meta)
            except Exception as exc:output[company]=finish(company,error=str(exc))
        try:_collect_document_batch(docs,origin,destination,on_company=ready)
        except Exception as exc:
            for company in docs:
                if company not in output:output[company]=finish(company,error=str(exc))
        for company in docs:
            if company not in output:ready(company)
        return list(output.values())

    def network_job(c):
        attempts[c]=begin_live_attempt(c,origin,destination,profile_id)
        try:
            if c in online_tariffs.ADAPTERS:
                if full_grid and c=='Пролайн' and {normalize_city(origin),normalize_city(destination)}!={'Москва','Санкт-Петербург'}:
                    from .proline import collect as proline_grid
                    vals,meta=proline_grid(origin,destination,profile_ids=[p['id'] for p in COMMON_PROFILES if not p.get('is_minimum_profile')])
                else:vals,meta=online_tariffs.collect(c,origin,destination,profile_id)
                return finish(c,vals,meta)
            if c=='КИТ':
                vals,meta=collect_kit(origin,destination)
                return finish(c,vals,meta)
            if c=='CTSgroup':return finish(c,result=collect_cts(origin,destination))
            if c=='Рейл Континент':
                vals,meta=_collect_rail_export(origin,destination)
                return finish(c,vals,meta)
            if c in {'Werner','Главтрасса'}:
                document_error=None
                if c=='Werner':
                    try:
                        from .official_documents import collect as collect_official_document
                        vals,meta=collect_official_document(c,origin,destination)
                        return finish(c,vals,meta)
                    except Exception as exc:document_error=str(exc)
                vals,meta=_collect_keyless_api_grid(c,origin,destination,profile_id)
                if document_error:meta['document_error']=document_error
                return finish(c,vals,meta)
            if c=='Новая Линия':return finish(c,result=collect_newline(origin,destination))
            if c=='Грузопоток':
                online_tariffs.fetch(c,origin,destination,'https://gruzopotok.com/','HTML')
                raise UnpublishedTariff('Грузопоток: автоматический весовой прайс не подтверждён. На странице заказчика стоимость предлагается запросить у менеджера. Загрузите полученный документ через «Загрузить прайс».')
            if c in {'ДЛ','Возовоз'}:
                from .official_documents import collect as collect_official_document
                try:
                    vals,meta=collect_official_document(c,origin,destination)
                    return finish(c,vals,meta)
                except Exception as document_error:
                    try:
                        vals,meta=(_collect_calculator_grid(c,origin,destination) if full_grid else _collect_live_calculator(c,origin,destination,profile_id))
                        meta['document_error']=str(document_error)
                        return finish(c,vals,meta)
                    except Exception as calculator_error:
                        raise RuntimeError(f'Прайс: {document_error}. Калькулятор: {calculator_error}. Можно загрузить PDF/Excel перевозчика кнопкой «Загрузить прайс».') from calculator_error
            vals,meta=(_collect_calculator_grid(c,origin,destination) if full_grid else _collect_live_calculator(c,origin,destination,profile_id))
            return finish(c,vals,meta)
        except Exception as exc:return finish(c,error=str(exc))

    from contextvars import copy_context
    tasks=iter([(network_job,c) for c in companies if c not in docs]+([(document_job,None)] if docs else []))
    with ThreadPoolExecutor(max_workers=7) as pool:
        futures={};exhausted=False
        while futures or not exhausted:
            while len(futures)<7 and not exhausted and not should_stop():
                task=next(tasks,None)
                if task is None:exhausted=True;break
                fn,c=task
                futures[pool.submit(copy_context().run,fn,*(() if c is None else (c,)))]=c
            if not futures:break
            done,_=wait(futures,return_when=FIRST_COMPLETED)
            for future in done:
                c=futures.pop(future)
                try:
                    output=future.result()
                    for result in output if c is None else [output]:results[result['company']]=result
                except Exception as exc:
                    for target in docs if c is None else [c]:results[target]=finish(target,error=str(exc))
    return [results[c] for c in companies if c in results]
