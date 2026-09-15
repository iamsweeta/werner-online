"""Populate the customer's workbook structure exclusively from application data.

The bundled template contains headings and formatting only, never customer
prices. Small parcels and MIN are RUB; heavy columns are published RUB/kg,
matching the reference workbook's two types of price columns.
"""
from __future__ import annotations

import io
import math
import zipfile
import xml.etree.ElementTree as ET
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.cell.cell import MergedCell
from openpyxl.workbook.views import BookView
from openpyxl.workbook.properties import CalcProperties

from . import v42_engine as e
from .tariff_model import tariff_value, is_rate_profile

SHEETS = {'Werner':'WernerNEW', 'Рейл Континент':'РейлКонтинент',
          'Новая Линия':'НоваяЛиния'}
GRAPH_SHEETS = ['WernerNEW','WernerOld','ДЛ','ПЭК','КИТ','Возовоз','Мейджик',
                'Главтрасса','Пролайн','Фортуна','ФастТранс',
                'РейлКонтинент','АТЭК','НоваяЛиния','БСК','ЭкспедицияПлюс','CTSgroup',
                'Енисей','Терминал','Сибтехнотранс','ИмперияАвто','Магистраль',
                'Транзит','АРТК','Аскор','КараванКарго','СолнечныйМагадан','Анкор',
                'ЖелдорАльянс','ЖелдорЭкспресс','МТК авто','МТК жд','Байкал Сервис']
HELPER_SHEETS = {'Сравнение-старыеW-новыеW','новыеW','старыеW','Ценообразование',
                 'склейка','рабочий','Графики','Направления'}


def _text(ws, row, col, value):
    if isinstance(ws.cell(row,col),MergedCell):return ws.cell(row,col)
    cell = ws.cell(row, col)
    if value is None:
        cell.value=None
        return cell
    cell.value=str(value)
    cell.data_type = 's'  # remote errors and filenames never become formulas
    return cell


def _routes(origin, destination, all_loaded):
    current = e.route_pair(origin, destination)
    found = {current}
    if all_loaded:
        found.update(e.ROUTE_CONFIG)
        paths = list((e.RUNTIME_DIR/'routes').glob('*.json'))
        paths += list((e.RUNTIME_DIR/'imports'/'routes').glob('*.json'))
        for path in paths[:1000]:
            data = e._read_json(path, {})
            candidates = [data.get('route', {})] + list(data.get('companies', {}).values())
            for row in candidates:
                o, d = row.get('origin'), row.get('destination')
                if o and d and e.is_supported_route(o,d):found.add(e.route_pair(o,d))
    return [current] + sorted(found-{current})


def _allowed(item, live_only, include_imports):
    return not live_only or item.get('online') or (include_imports and item.get('uploaded'))


def _value(item, profile, live_only, include_imports):
    if item.get('bulk_status')=='pending' or not _allowed(item,live_only,include_imports):return None
    value=tariff_value(item, profile)
    return round(value,4) if value is not None else None



def build_workbook(origin, destination, selected, *, live_only=False, include_imports=True, all_loaded=True,
                   routes_override=None, matrix_provider=None, collection_info=None):
    wb=load_workbook(e.DATA_DIR/'export_template.xlsx')
    if 'Грузопоток' in wb:del wb['Грузопоток']
    if 'Байкал Сервис' not in wb:
        ws=wb.copy_worksheet(wb['ДЛ']);ws.title='Байкал Сервис'
    mapping={SHEETS.get(c,c):c for c in e.COMPANIES}
    routes=list(routes_override) if routes_override is not None else _routes(origin,destination,all_loaded)
    if not routes:raise ValueError('Нет маршрутов для Excel')
    origin,destination=routes[0]
    n=len(routes)+1
    selected=set(selected)
    # One snapshot protects the export from a concurrent import/refresh.
    if matrix_provider is not None:
        data=((route,matrix_provider(*route)) for route in routes)
    else:
        with e.STATE_LOCK:
            snapshot={route:e.matrix(*route,list(e.COMPANIES)) for route in routes}
        data=snapshot.items()
    values={}
    counts={'online':0,'document':0,'missing':0}
    audit=wb.create_sheet('Источники')
    audit.append(['Компания','Откуда','Куда','Диапазон','Стоимость отправки, ₽','Статус',
                  'Получено','Официальный URL','Расчёт','Файл пользователя','Дата документа',
                  'SHA256','Страница / строка','Причина отсутствия / ошибка','Опубликованная ставка, ₽/кг','Минимальная плата, ₽'])
    for route, rows in data:
        audit_groups={}
        for row in rows:
            p=row['profile']
            for item in row['items']:
                c=item['company']
                values[(route,c,p['id'])]=_value(item,p,live_only,include_imports) if c in selected else None
                if c not in selected:continue
                if _allowed(item,live_only,include_imports) and tariff_value(item,p) is not None:
                    counts['document' if item.get('uploaded') else 'online']+=1
                else:counts['missing']+=1
                allowed=_allowed(item,live_only,include_imports)
                status=('Проверено при сборе' if collection_info and item.get('collected_online') else 'Файл пользователя' if item.get('uploaded') else 'LIVE' if item.get('online')
                        else 'LAST GOOD' if item.get('price') is not None else 'Нет данных')
                if collection_info and item.get('availability')=='on_request':status='По запросу'
                if item.get('bulk_status')=='pending':status='Не проверено'
                record=[c,*route,p['label'],item.get('comparison_value') if allowed else None,status,
                              item.get('captured_at'),item.get('source_url'),item.get('calculation_basis'),
                              item.get('original_filename'),item.get('document_date'),item.get('sha256'),
                              str(item.get('source_page') or item.get('source_row') or ''),
                              item.get('refresh_error') or ('Источник даёт сумму отправки; ставка за кг не опубликована' if is_rate_profile(p) and item.get('comparison_value') is not None and tariff_value(item,p) is None else item.get('message') if item.get('price') is None else ''),
                              item.get('published_rate_per_kg') if allowed else None,item.get('minimum_charge') if allowed else None]
                if collection_info:
                    record[6]=record[6] or item.get('checked_at')
                    key=tuple(str(v or '') for i,v in enumerate(record) if i not in (3,4))
                    if key not in audit_groups:audit_groups[key]=record
                    else:
                        audit_groups[key][3]+=', '+p['label']
                        audit_groups[key][4]=None
                else:audit.append(record)
        for record in audit_groups.values():audit.append(record)
    for ws in wb:
        if ws.title in HELPER_SHEETS or ws.title=='Источники':continue
        company=mapping.get(ws.title)
        _text(ws,1,1,ws.title);_text(ws,1,3,None)
        for idx,route in enumerate(routes,start=2):
            ws.cell(idx,1,idx-1);_text(ws,idx,2,route[0]);_text(ws,idx,3,route[1])
            for col,p in enumerate(e.COMMON_PROFILES,start=4):
                val=values.get((route,company,p['id']))
                if isinstance(val,(int,float)):
                    ws.cell(idx,col,val)
                else:_text(ws,idx,col,val)
            # The reference has volume columns. No volume tariffs are invented
            # from the weight grid or copied from the customer's workbook.
            for col in range(34,64):_text(ws,idx,col,None)
        ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:AF{n}'
    glue=wb['склейка']
    for col,value in enumerate(['№','Из','В','Маршрут','Строка'],start=1):_text(glue,1,col,value)
    directions=wb['Направления']
    for r,(o,d) in enumerate(routes,start=2):
        for col,value in enumerate([r-1,o,d,f'{o} → {d}',r],start=1):
            glue.cell(r,col,value)
        directions.cell(r,2,r-1);_text(directions,r,3,o);_text(directions,r,4,d)
        _text(directions,r,5,'Общий онлайн-сбор' if collection_info else 'Выбранный' if r==2 else 'Сохранённый')
    graph=wb['Графики'];work=wb['рабочий']
    _text(graph,16,2,origin);_text(graph,16,6,destination)
    _text(graph,14,2,None);_text(graph,15,2,None)
    _text(graph,18,3,None);graph.column_dimensions['D'].hidden=True
    _text(graph,54,2,None)
    _text(graph,55,2,None)
    if collection_info:
        _text(graph,54,2,None)
        _text(graph,56,2,None)
        report=wb.create_sheet('Сбор')
        for label,value in [
            ('Отчёт','Все компании и маршруты'),('Начало проверки',collection_info['created_at']),
            ('Последний ответ',collection_info['updated_at']),('Создан Excel',datetime.now().astimezone().isoformat(timespec='seconds')),
            ('Маршрутов в плане',collection_info['total_routes']),('Маршрутов проверено',collection_info['completed_routes']),
            ('Маршрутов в этом файле',len(routes)),('Проверок компаний в плане',collection_info['total_checks']),
            ('Проверок выполнено',collection_info['completed_checks']),('Часть',f"{collection_info.get('part',1)} из {collection_info.get('parts',1)}"),
            ('Режим','Текущие данные + прайс-листы' if collection_info.get('mode')=='saved' else 'Обновление онлайн + прайс-листы' if collection_info.get('include_imports') else 'Обновление только онлайн'),
            ('Числовых ячеек из онлайн-ответов',counts['online']),('Числовых ячеек из файлов',counts['document']),
            ('Ячеек без точного значения',counts['missing']),
            ('Завершение','Все запланированные проверки выполнены' if collection_info['status']=='done' else 'Частичная выгрузка: непроверенные маршруты не включены'),
            ('Обновление','Запустите новый общий сбор в приложении, затем скачайте Excel'),
            ('Время цен','Цены проверялись последовательно. Дата каждого источника указана в «Источники». Это не одновременный снимок всех сайтов.'),
            ('Источники','Подтверждённая онлайн-цена имеет приоритет. При её отсутствии используется подтверждённый файл пользователя. Архивные цены не подставляются.' if collection_info.get('include_imports') else 'Только ответы онлайн-сбора. Файлы пользователя и архивные цены не включены.'),
            ('Весовые источники','До 50 кг — сумма отправки; от 100 кг — опубликованная ставка руб/кг. Сумма и минимальная плата хранятся отдельно в «Источниках». Нет ставки — известна только сумма расчёта.'),
            ('Отсутствующие цены','Пустая ячейка: числовой тариф не получен, не опубликован или исключён фильтром. Причина — в Источниках. Это не нулевая стоимость.')]:
            report.append([label,value])
        report.append(['Результат проверки компаний','Количество'])
        labels={'complete':'Все 28 весов','partial':'Неполная онлайн-сетка / по запросу','failed':'Ошибка онлайн-источника','unavailable':'Нет подключённого публичного тарифа','saved':'Обработано без новых сетевых запросов'}
        for status,count in collection_info['outcomes'].items():report.append([labels.get(status,status),count])
        report.column_dimensions['A'].width=38;report.column_dimensions['B'].width=105
        for row in report:
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
            report.row_dimensions[row[0].row].height=44 if len(str(row[1].value or ''))>110 else 28
            row[0].font=Font(name='Arial',size=10,bold=True)
            row[0].fill=PatternFill('solid',fgColor='EDF2FF')
        report.freeze_panes='B2'
    work['C1']="='Графики'!B16&\" → \"&'Графики'!F16"
    work['G1']=f'=IF(COUNTIF(\'склейка\'!D2:D{n},C1)=0,"",MATCH(C1,\'склейка\'!D2:D{n},0)+1)'
    cache={'рабочий':{'C1':f'{origin} → {destination}','G1':2},'Графики':{}}
    for cell,col in [('B16','C'),('F16','D')]:
        dv=DataValidation(type='list',formula1=f"'Направления'!${col}$2:${col}${n}")
        graph.add_data_validation(dv);dv.add(graph[cell])
    for idx,title in enumerate(GRAPH_SHEETS):
        r=19+idx;wr=4+idx;company=mapping.get(title)
        _text(graph,r,2,title);graph.cell(r,3,company in selected)
        _text(work,wr,1,title);work.cell(wr,2,f"='Графики'!C{r}")
        cache['рабочий'][f'B{wr}']=company in selected
        for col,p in enumerate(e.COMMON_PROFILES,start=4):
            from openpyxl.utils import get_column_letter
            letter=get_column_letter(col)
            ref=f"INDEX('{title}'!{letter}$1:{letter}${n},$G$1)"
            work.cell(wr,col,f'=IF($B{wr}=FALSE,"",IF($G$1="","",IF(ISNUMBER({ref}),{ref},"")))')
            graph.cell(r,col+1,f'=IF(ISNUMBER(\'рабочий\'!{letter}{wr}),\'рабочий\'!{letter}{wr},"")')
            cached=wb[title].cell(2,col).value if company in selected else None
            cache['рабочий'][f'{letter}{wr}']=cached
            cache['Графики'][f'{get_column_letter(col+1)}{r}']=cached
        for col in range(35,65):_text(graph,r,col,None)
    graph.freeze_panes='E19';work.freeze_panes='D4'
    # Preserve the reference's old/new tariff tabs as explicit empty history,
    # without reusing a single old price or pricing coefficient.
    for title in ['новыеW','старыеW','Сравнение-старыеW-новыеW','Ценообразование']:
        ws=wb[title]
        _text(ws,7,1,None)
    for title in ['старыеW','Сравнение-старыеW-новыеW','Ценообразование']:
        wb[title].sheet_state='hidden'
    wb.active=wb.sheetnames.index('Графики')
    wb.views=[BookView(activeTab=wb.sheetnames.index('Графики'))]
    wb.calculation=CalcProperties(calcMode='auto',fullCalcOnLoad=True,forceFullCalc=True)
    wb._tariff_formula_cache=cache
    # Native charts follow the selected route through the working sheet.
    for start,end,title,anchor in [(5,13,'Малые отправления, ₽','B1'),(15,33,'Тарифы от 100 кг, ₽/кг','R1')]:
        chart=LineChart();chart.title=title;chart.height=6;chart.width=24
        chart.display_blanks='gap';chart.y_axis.title='₽' if start==5 else '₽/кг'
        for idx,sheet in enumerate(GRAPH_SHEETS):
            if mapping.get(sheet) not in selected:continue
            row=19+idx
            chart.add_data(Reference(graph,min_col=start,max_col=end,min_row=row,max_row=row),from_rows=True)
            chart.series[-1].tx = None
            from openpyxl.chart.series import SeriesLabel
            chart.series[-1].tx=SeriesLabel(strRef=None,v=sheet)
        chart.set_categories(Reference(graph,min_col=start,max_col=end,min_row=18,max_row=18))
        graph.add_chart(chart,anchor)
    for ws in wb:
        ws.sheet_view.showGridLines=False
        ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        for row in ws:
            for cell in row:
                if cell.row>1 and not cell.has_style:
                    cell.font=Font(name='Arial',size=10)
                    cell.alignment=Alignment(vertical='center')
    audit.freeze_panes='D2';audit.auto_filter.ref=audit.dimensions
    for cell in audit[1]:
        cell.font=Font(name='Arial',size=10,bold=True);cell.fill=PatternFill('solid',fgColor='FFF2CC')
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    audit.row_dimensions[1].height=32;audit.column_dimensions['E'].width=24
    for row in audit:
        for cell in row:
            if cell.data_type=='f':cell.data_type='s'
    for col in ['A','B','C','D','F','G']:audit.column_dimensions[col].width=24
    for col,width in [('J',38),('K',20),('L',26),('M',18),('N',65),('G',32)]:audit.column_dimensions[col].width=width
    audit.row_dimensions[1].height=44
    audit.column_dimensions['H'].width=50;audit.column_dimensions['I'].width=70
    if collection_info:
        audit.column_dimensions['D'].width=48
        for row in audit.iter_rows(min_row=2):
            for col in (3,6,8,9,10,11,12,13):row[col].alignment=Alignment(wrap_text=True,vertical='top')
            audit.row_dimensions[row[0].row].height=100 if len(str(row[3].value or ''))>180 else 56
    return wb


def export_bytes(*args,**kwargs):
    wb=build_workbook(*args,**kwargs)
    out=io.BytesIO();wb.save(out)
    # Store the evaluated snapshot as Excel formula caches for viewers that do
    # not recalculate on opening. The formulas remain editable and auto-recalc.
    ns='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    cache={f'xl/worksheets/sheet{wb.sheetnames.index(name)+1}.xml':cells
           for name,cells in wb._tariff_formula_cache.items()}
    result=io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(out.getvalue())) as source,zipfile.ZipFile(result,'w',zipfile.ZIP_DEFLATED) as target:
        for entry in source.infolist():
            raw=source.read(entry.filename)
            if entry.filename in cache:
                root=ET.fromstring(raw)
                for cell in root.iter(ns+'c'):
                    address=cell.get('r')
                    if address not in cache[entry.filename] or cell.find(ns+'f') is None:continue
                    value=cache[entry.filename][address]
                    v=cell.find(ns+'v')
                    if v is None:v=ET.SubElement(cell,ns+'v')
                    if isinstance(value,bool):cell.set('t','b');v.text='1' if value else '0'
                    elif isinstance(value,(int,float)):
                        cell.attrib.pop('t',None);v.text=str(value)
                    else:cell.set('t','str');v.text=str(value or '')
                raw=ET.tostring(root,encoding='utf-8',xml_declaration=True)
            target.writestr(entry,raw)
    return result.getvalue()
