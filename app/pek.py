"""Strict, bounded-memory reader for the current official PEK workbook."""
from __future__ import annotations
import hashlib
import math
import threading
from collections import OrderedDict
from openpyxl import load_workbook
from .cities import normalize_city,UnpublishedTariff
from .online_tariffs import number,norm,values_from_tiers
from .source_conditions import tax_basis

_CACHE=OrderedDict()
_LOCK=threading.Lock()

def parse_route(path,origin,destination):
    origin=normalize_city(origin);destination=normalize_city(destination)
    digest=hashlib.sha256(path.read_bytes()).hexdigest();key=(digest,origin)
    with _LOCK:
        cached=_CACHE.get(key)
        if cached is None:
            cached=_read_origin(path,origin)
            _CACHE[key]=cached
            while len(_CACHE)>4:_CACHE.popitem(last=False)
        else:_CACHE.move_to_end(key)
    routes,group,header,conditions=cached
    found=routes.get(destination,[])
    if not found:raise UnpublishedTariff('ПЭК: нет строки выбранного направления в свежем XLSX')
    if len(found)!=1:raise ValueError('ПЭК: повторяется направление с неоднозначными ценами')
    row_no,row=found[0];tiers=[];fixed=[];active=''
    from .legacy_backend import _strict_weight_range
    minimum=number(row[8]) if row[8] not in ('',None) else None
    for i in range(3,21):
        if group[i]:active=norm(group[i])
        if i==8 or row[i] in ('',None,'-','—'):continue
        limits=_strict_weight_range(str(header[i] or ''))
        if not limits:raise ValueError('ПЭК: не распознан весовой заголовок')
        target=tiers if '1 кг' in active else fixed
        target.append((limits[0] or 0,limits[1] if limits[1] is not None else math.inf,number(row[i])))
    if not tiers or not fixed:raise ValueError('ПЭК: фиксированные и весовые тарифы не подтверждены')
    values=values_from_tiers(tiers,fixed=fixed,minimum=minimum)
    return values,{**conditions,'source_row':row_no,'route_verified':True,
                   'calculation_basis':'Фиксированный тариф малого груза с ограничением объёма по прайсу; далее max(минимальная плата, вес × ₽/кг). Допуслуги не включены.'}

def _read_origin(path,origin):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        if 'Перевозка' not in wb.sheetnames:raise ValueError('ПЭК: нет листа «Перевозка»')
        rows=wb['Перевозка'].iter_rows(max_col=21,values_only=True)
        preamble=[next(rows) for _ in range(21)]
        group,header=preamble[19:21]
        if norm(header[1])!='город отправитель' or norm(header[2])!='город получатель' or any('руб' not in norm(group[i]) for i in (3,8,9)):
            raise ValueError('ПЭК: заголовки направления и валюты изменились')
        labels=' '.join(norm(x) for x in group)
        if '1 кг' not in labels or 'руб' not in labels:raise ValueError('ПЭК: единицы тарифов не подтверждены')
        routes={}
        for row_no,row in enumerate(rows,22):
            if normalize_city(row[1])==origin:
                routes.setdefault(normalize_city(row[2]),[]).append((row_no,row))
        from .tariff_documents import dated
        text='\n'.join(' '.join(str(v or '') for v in r) for r in preamble)
        return routes,group,header,{'tax_basis':tax_basis(' '.join(str(v or '') for v in group)),'document_date':dated(text)}
    except StopIteration as exc:raise ValueError('ПЭК: файл обрывается до заголовков') from exc
    finally:wb.close()
