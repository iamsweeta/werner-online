from __future__ import annotations

import base64, io, json, os, re, subprocess, tempfile, threading, time, traceback

import requests
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any
from urllib.parse import quote, unquote, urlencode, urljoin, urlsplit, urlunsplit

from bs4 import BeautifulSoup

from .v40_engine import COMMON_PROFILES, PROFILE_BY_ID, RUNTIME_DIR, save_live_update

LOG_PATH=RUNTIME_DIR/"v40_collect.log"
_BROWSER_LOCK=threading.Lock()

def _log(msg: str) -> None:
    RUNTIME_DIR.mkdir(parents=True, exist_ok=True)
    stamp=datetime.now().astimezone().isoformat(timespec="seconds")
    with LOG_PATH.open("a",encoding="utf-8") as f: f.write(f"[{stamp}] {msg}\n")
    print(f"[v40] {msg}", flush=True)

def _norm(s: str) -> str:
    return " ".join(str(s or "").lower().replace("ё","е").split())

def _num(s: Any) -> float|None:
    t=str(s or "").replace("\xa0"," ").replace(",",".")
    m=re.search(r"(?<!\d)(\d+(?:[ .]\d{3})*(?:\.\d+)?)",t)
    if not m:return None
    try:return float(m.group(1).replace(" ",""))
    except:return None

def _make_driver():
    from selenium import webdriver
    errors=[]
    try:
        opt=webdriver.EdgeOptions(); opt.add_argument("--headless=new"); opt.add_argument("--disable-gpu"); opt.add_argument("--no-first-run"); opt.add_argument("--disable-extensions"); opt.add_argument("--window-size=1400,1000")
        return webdriver.Edge(options=opt),"Edge"
    except Exception as e: errors.append(f"Edge: {e}")
    try:
        opt=webdriver.ChromeOptions(); opt.add_argument("--headless=new"); opt.add_argument("--disable-gpu"); opt.add_argument("--no-first-run"); opt.add_argument("--disable-extensions"); opt.add_argument("--window-size=1400,1000")
        return webdriver.Chrome(options=opt),"Chrome"
    except Exception as e: errors.append(f"Chrome: {e}")
    raise RuntimeError("; ".join(errors))


def _curl_fetch(url: str, *, method: str = "GET", fields: list[tuple[str,str]] | None = None, referer: str | None = None, cookie_jar: str | None = None, timeout: int = 40) -> bytes:
    exe = "curl.exe" if os.name == "nt" else "curl"
    base=[exe,"-L","--fail","--silent","--show-error","--connect-timeout","10","--max-time",str(timeout),
          "-A","Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36"]
    if os.name == "nt": base += ["--ssl-no-revoke"]
    if referer: base += ["-e",referer]
    if cookie_jar: base += ["-b",cookie_jar,"-c",cookie_jar]
    target=url
    data_args=[]
    if method.upper()=="POST": data_args=["-X","POST","--data",urlencode(fields or [])]
    elif fields:
        sep="&" if "?" in target else "?"; target += sep + urlencode(fields)
    attempts=[[]]
    if os.name == "nt":
        # Some Russian carrier sites terminate Schannel's default handshake.
        # A forced IPv4 / HTTP/1.1 / TLS1.2 retry uses the same trusted cert
        # store without disabling certificate verification.
        attempts.append(["-4","--http1.1","--tlsv1.2"])
    errors=[]
    for extra in attempts:
        cmd=base+extra+data_args+[target]
        try:
            cp=subprocess.run(cmd,stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+5,check=False)
        except Exception as exc:
            errors.append(str(exc)); continue
        if cp.returncode==0 and cp.stdout: return bytes(cp.stdout)
        errors.append((cp.stderr or b"curl failed").decode("utf-8","replace")[-800:])
    raise RuntimeError(" | ".join(errors[-2:]) or "curl failed")

def _powershell_fetch(url: str, *, timeout: int = 45) -> bytes:
    """Windows-native download without interpolating the URL into PowerShell code."""
    with tempfile.TemporaryDirectory(prefix="tariff_ps_") as td:
        out=Path(td)/"response.bin"
        env=os.environ.copy(); env["V40_URL"]=url; env["V40_OUT"]=str(out)
        script=(
            "$ErrorActionPreference='Stop'; $ProgressPreference='SilentlyContinue'; "
            "$h=@{'User-Agent'='Mozilla/5.0 (Windows NT 10.0; Win64; x64)'}; "
            "Invoke-WebRequest -UseBasicParsing -Uri $env:V40_URL -Headers $h -OutFile $env:V40_OUT"
        )
        cp=subprocess.run(["powershell.exe","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-Command",script],
                          stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+10,check=False,env=env)
        if cp.returncode!=0 or not out.exists() or out.stat().st_size==0:
            err=(cp.stderr or cp.stdout or b"PowerShell download failed").decode("utf-8","replace")[-1000:]
            raise RuntimeError(err)
        return out.read_bytes()

def _winhttp_fetch(url: str, *, timeout: int = 45) -> bytes:
    """Last Windows-native HTTPS transport using WinHTTP/COM and normal cert validation."""
    with tempfile.TemporaryDirectory(prefix="tariff_winhttp_") as td:
        out=Path(td)/"response.bin"
        env=os.environ.copy(); env["V40_URL"]=url; env["V40_OUT"]=str(out); env["V40_TIMEOUT_MS"]=str(max(5000,int(timeout*1000)))
        script=(
            "$ErrorActionPreference='Stop'; "
            "$u=$env:V40_URL; $o=$env:V40_OUT; $ms=[int]$env:V40_TIMEOUT_MS; "
            "$w=New-Object -ComObject 'WinHttp.WinHttpRequest.5.1'; "
            "$w.SetTimeouts(10000,10000,$ms,$ms); $w.Open('GET',$u,$false); "
            "$w.SetRequestHeader('User-Agent','Mozilla/5.0 (Windows NT 10.0; Win64; x64)'); $w.Send(); "
            "if($w.Status -lt 200 -or $w.Status -ge 400){throw ('HTTP '+$w.Status)}; "
            "$s=New-Object -ComObject ADODB.Stream; $s.Type=1; $s.Open(); $s.Write($w.ResponseBody); $s.SaveToFile($o,2); $s.Close()"
        )
        cp=subprocess.run(["powershell.exe","-NoProfile","-NonInteractive","-ExecutionPolicy","Bypass","-Command",script],
                          stdout=subprocess.PIPE,stderr=subprocess.PIPE,timeout=timeout+15,check=False,env=env)
        if cp.returncode!=0 or not out.exists() or out.stat().st_size==0:
            err=(cp.stderr or cp.stdout or b"WinHTTP download failed").decode("utf-8","replace")[-1000:]
            raise RuntimeError(err)
        return out.read_bytes()

def _browser_fetch_bytes(url: str, referer: str) -> bytes:
    driver=None
    with _BROWSER_LOCK:
        try:
            driver,browser=_make_driver(); driver.set_page_load_timeout(35); driver.get(referer); time.sleep(.5)
            result=driver.execute_async_script(r"""
              const done=arguments[arguments.length-1], u=arguments[0];
              fetch(u,{credentials:'include',cache:'no-store'}).then(async r=>{
                const b=new Uint8Array(await r.arrayBuffer()); let s='';
                for(let i=0;i<b.length;i+=0x8000)s+=String.fromCharCode(...b.subarray(i,i+0x8000));
                done({ok:r.ok,status:r.status,url:r.url,data:btoa(s)});
              }).catch(e=>done({ok:false,error:String(e)}));
            """,url) or {}
            if not result.get('ok') or not result.get('data'): raise RuntimeError(str(result.get('error') or result.get('status')))
            return base64.b64decode(result['data'])
        finally:
            if driver:
                try: driver.quit()
                except: pass

def _download_bytes(url: str, *, referer: str | None = None, timeout: int = 45) -> tuple[bytes,str]:
    errors=[]
    headers={"User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36","Referer":referer or url}
    try:
        r=requests.get(url,headers=headers,timeout=(10,timeout),allow_redirects=True); r.raise_for_status()
        if r.content: return bytes(r.content),"requests"
    except Exception as e: errors.append(f"requests: {type(e).__name__}: {e}")
    try: return _curl_fetch(url,referer=referer,timeout=timeout),"curl"
    except Exception as e: errors.append(f"curl: {type(e).__name__}: {e}")
    if os.name == "nt":
        try: return _powershell_fetch(url,timeout=timeout),"PowerShell"
        except Exception as e: errors.append(f"powershell: {type(e).__name__}: {e}")
        try: return _winhttp_fetch(url,timeout=timeout),"WinHTTP"
        except Exception as e: errors.append(f"winhttp: {type(e).__name__}: {e}")
    if referer:
        try: return _browser_fetch_bytes(url,referer),"headless-fetch"
        except Exception as e: errors.append(f"headless: {type(e).__name__}: {e}")
    raise RuntimeError("; ".join(errors))

def _fetch_submitted_form_html_browser(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Submit the site's real route form with fetch(), without visible windows.

    The browser is only a same-origin HTTP client here. We do not click custom
    widgets and we never read a result before validating the response route.
    """
    driver=None
    with _BROWSER_LOCK:
        try:
            driver,browser=_make_driver(); driver.set_page_load_timeout(35); driver.get(url); time.sleep(1.0)
            script=r"""
            const done=arguments[arguments.length-1];
            const origin=arguments[0], dest=arguments[1];
            const norm=s=>(s||'').trim().toLowerCase().replace(/ё/g,'е');
            (async()=>{
              try{
                const forms=[...document.forms];
                let hit=null, city=[];
                for(const form of forms){
                  const selects=[...form.querySelectorAll('select')];
                  const candidates=selects.filter(s=>{const opts=[...s.options].map(o=>norm(o.textContent));return opts.includes(norm(origin))&&opts.includes(norm(dest));});
                  if(candidates.length>=2){hit=form;city=candidates;break;}
                }
                if(!hit){ done({ok:false,error:'route form with two city selects not found'}); return; }
                function setText(s,text){ const o=[...s.options].find(x=>norm(x.textContent)===norm(text)); if(!o)return false; s.value=o.value; [...s.options].forEach(x=>x.selected=(x===o)); s.dispatchEvent(new Event('change',{bubbles:true})); return true; }
                if(!setText(city[0],origin)||!setText(city[1],dest)){done({ok:false,error:'city options not set'});return;}
                const fd=new FormData(hit);
                const submit=hit.querySelector('button[type=submit],input[type=submit],button:not([type])');
                if(submit && submit.name && !fd.has(submit.name)) fd.append(submit.name, submit.value||submit.textContent||'1');
                let action=hit.action||location.href, method=(hit.method||'GET').toUpperCase();
                let resp;
                if(method==='GET'){
                  const qs=new URLSearchParams(fd).toString();
                  const u=new URL(action,location.href); u.search=qs; resp=await fetch(u.toString(),{credentials:'include',redirect:'follow'});
                } else {
                  resp=await fetch(action,{method,body:fd,credentials:'include',redirect:'follow'});
                }
                const text=await resp.text(); done({ok:resp.ok,status:resp.status,url:resp.url,text,method,from:city[0].value,to:city[1].value});
              }catch(e){done({ok:false,error:String(e&&e.stack||e)});}
            })();
            """
            result=driver.execute_async_script(script,origin,destination) or {}
            if not result.get("ok") or not result.get("text"):
                raise RuntimeError(str(result.get("error") or f"HTTP {result.get('status')}"))
            return str(result["text"]),str(result.get("url") or url),browser
        finally:
            if driver:
                try:driver.quit()
                except:pass


def _matching_city_selects(form, origin: str, destination: str):
    wanted={_norm(origin),_norm(destination)}
    found=[]
    for select in form.find_all("select"):
        labels={_norm(o.get_text(" ",strip=True)) for o in select.find_all("option")}
        if wanted.issubset(labels) and select.get("name"):
            found.append(select)
    return found

def _option_value(select, text: str) -> str|None:
    wanted=_norm(text)
    for option in select.find_all("option"):
        if _norm(option.get_text(" ",strip=True))==wanted:
            return str(option.get("value") if option.get("value") is not None else option.get_text(" ",strip=True))
    return None

def _successful_form_fields(form) -> list[tuple[str,str]]:
    fields=[]
    for tag in form.find_all(["input","select","textarea"]):
        name=tag.get("name")
        if not name or tag.has_attr("disabled"): continue
        if tag.name=="input":
            typ=_norm(tag.get("type") or "text")
            if typ in {"button","reset","file","image"}: continue
            if typ in {"checkbox","radio"} and not tag.has_attr("checked"): continue
            fields.append((name,str(tag.get("value") or "")))
        elif tag.name=="textarea":
            fields.append((name,tag.get_text()))
        else:
            selected=tag.find_all("option",selected=True)
            if not selected:
                first=tag.find("option")
                selected=[first] if first else []
            for option in selected:
                fields.append((name,str(option.get("value") if option.get("value") is not None else option.get_text(" ",strip=True))))
    submit=form.find(["button","input"],attrs={"type":re.compile(r"^submit$",re.I)})
    if submit and submit.get("name"):
        fields.append((str(submit.get("name")),str(submit.get("value") or submit.get_text(" ",strip=True) or "1")))
    return fields

def _replace_field(fields:list[tuple[str,str]], name:str, value:str) -> list[tuple[str,str]]:
    out=[(k,v) for k,v in fields if k!=name]
    out.append((name,value))
    return out

def _fetch_submitted_form_html_http(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Submit the site's native route form directly over HTTP.

    This is the primary path. It mirrors a normal browser form submit but does not
    depend on Select2, JavaScript clicks, WebDriver or a visible browser window.
    """
    session=requests.Session()
    headers={
        "User-Agent":"Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/140 Safari/537.36",
        "Accept":"text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
        "Accept-Language":"ru-RU,ru;q=0.9,en;q=0.6",
        "Cache-Control":"no-cache",
    }
    r=session.get(url,headers=headers,timeout=(10,30),allow_redirects=True)
    r.raise_for_status()
    soup=BeautifulSoup(r.text,"lxml")
    hit=None; city=[]
    for form in soup.find_all("form"):
        c=_matching_city_selects(form,origin,destination)
        if len(c)>=2:
            hit=form; city=c[:2]; break
    if hit is None:
        raise RuntimeError("HTTP: форма маршрута с двумя списками городов не найдена")
    from_value=_option_value(city[0],origin); to_value=_option_value(city[1],destination)
    if from_value is None or to_value is None:
        raise RuntimeError("HTTP: значения городов не найдены в форме")
    fields=_successful_form_fields(hit)
    fields=_replace_field(fields,str(city[0].get("name")),from_value)
    fields=_replace_field(fields,str(city[1].get("name")),to_value)
    action=urljoin(r.url,hit.get("action") or r.url)
    method=_norm(hit.get("method") or "get").upper()
    req_headers={**headers,"Referer":r.url}
    if method=="POST":
        rr=session.post(action,data=fields,headers=req_headers,timeout=(10,35),allow_redirects=True)
    else:
        rr=session.get(action,params=fields,headers=req_headers,timeout=(10,35),allow_redirects=True)
    rr.raise_for_status()
    if not rr.text.strip(): raise RuntimeError("HTTP: сайт вернул пустой ответ")
    return rr.text,rr.url,"HTTP"


def _fetch_submitted_form_html_curl(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    """Same native-form submission as HTTP path, but through system curl.
    This is important on Windows hosts where Python/OpenSSL is rejected by the carrier TLS endpoint.
    """
    with tempfile.TemporaryDirectory() as td:
        jar=str(Path(td)/"cookies.txt")
        raw=_curl_fetch(url,cookie_jar=jar,timeout=40)
        text=raw.decode("utf-8","replace")
        soup=BeautifulSoup(text,"lxml")
        hit=None; city=[]
        for form in soup.find_all("form"):
            c=_matching_city_selects(form,origin,destination)
            if len(c)>=2: hit=form; city=c[:2]; break
        if hit is None: raise RuntimeError("curl: форма маршрута с двумя списками городов не найдена")
        a=_option_value(city[0],origin); b=_option_value(city[1],destination)
        if a is None or b is None: raise RuntimeError("curl: значения городов не найдены")
        fields=_successful_form_fields(hit)
        fields=_replace_field(fields,str(city[0].get("name")),a); fields=_replace_field(fields,str(city[1].get("name")),b)
        action=urljoin(url,hit.get("action") or url); method=_norm(hit.get("method") or "get").upper()
        out=_curl_fetch(action,method=method,fields=fields,referer=url,cookie_jar=jar,timeout=45)
        return out.decode("utf-8","replace"),action,"curl"

def _fetch_validated_route_table(url: str, origin: str, destination: str):
    """Get a route table with strict route validation.

    Direct HTTP is preferred. A fully headless browser is used only if the site
    requires JavaScript. A failed attempt never changes the saved route pack.
    """
    errors=[]
    for label,fn in (("HTTP",_fetch_submitted_form_html_http),("curl",_fetch_submitted_form_html_curl),("headless",_fetch_submitted_form_html_browser)):
        try:
            html,final_url,transport=fn(url,origin,destination)
            soup=BeautifulSoup(html,"lxml")
            headers,cells,headings=_find_route_table(soup,origin,destination)
            return headers,cells,headings,final_url,transport
        except Exception as exc:
            errors.append(f"{label}: {type(exc).__name__}: {exc}")
            _log(f"route form {url} {label} failed: {type(exc).__name__}: {exc}")
    raise RuntimeError("; ".join(errors))

# Compatibility alias used by older tests/instrumentation. It is intentionally
# headless-only; new collectors use _fetch_validated_route_table instead.
def _fetch_submitted_form_html(url: str, origin: str, destination: str) -> tuple[str,str,str]:
    return _fetch_submitted_form_html_browser(url,origin,destination)

def _find_route_table(soup: BeautifulSoup, origin: str, destination: str):
    wanted=_norm(destination)
    # Strong route validation: a heading must identify the selected origin when present.
    headings=" ".join(x.get_text(" ",strip=True) for x in soup.find_all(["h1","h2","h3","h4"]))
    if "из города" in _norm(headings) and _norm(origin) not in _norm(headings):
        raise RuntimeError(f"ответ относится к другому городу отправления: {headings[:240]}")
    for table in soup.find_all("table"):
        headers=[x.get_text(" ",strip=True) for x in table.find_all("th")]
        for tr in table.find_all("tr"):
            cells=[x.get_text(" ",strip=True) for x in tr.find_all(["th","td"])]
            if cells and _norm(cells[0])==wanted:
                return headers,cells,headings
    raise RuntimeError(f"в ответе нет строки назначения {destination}")


def _cts_parse_xlsx_bytes(raw: bytes, source_url: str) -> tuple[list[str],list[str]]:
    from openpyxl import load_workbook
    wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True)
    wanted=_norm("Москва")
    for ws in wb.worksheets:
        recent_headers=[]
        for row in ws.iter_rows(values_only=True):
            vals=["" if v is None else str(v).strip() for v in row]
            joined=" | ".join(vals)
            if any(k in _norm(joined) for k in ("город получения","цена","до 100","1-100")):
                recent_headers=vals
            first=next((v for v in vals[:4] if v),"")
            if _norm(first)==wanted:
                nums=[_num(v) for v in vals]
                if sum(x is not None for x in nums)>=8:
                    return recent_headers,vals
    raise RuntimeError("CTS XLSX: строка Москва с тарифными колонками не найдена")


def _cts_xlsx_candidates(page_html: str, page_url: str) -> list[str]:
    soup=BeautifulSoup(page_html,"lxml"); out=[]
    for a in soup.find_all("a",href=True):
        href=urljoin(page_url,str(a.get("href")))
        if not href.lower().split("?",1)[0].endswith(".xlsx"): continue
        decoded=unquote(href)
        variants=[]
        for old,new in (("МСК","СПБ"),("мск","спб"),("MSK","SPB"),("Moscow","SPB"),("Москва","СПБ")):
            if old in decoded: variants.append(decoded.replace(old,new))
        for v in variants:
            parts=urlsplit(v); enc_path=quote(unquote(parts.path),safe="/%:@")
            candidate=urlunsplit((parts.scheme,parts.netloc,enc_path,parts.query,parts.fragment))
            if candidate not in out: out.append(candidate)
    return out

def _cts_recent_xlsx_candidates(days_back: int = 70) -> list[str]:
    """Generate validated SPB filename candidates even when /prices itself is TLS-blocked.

    CTS' official Moscow download currently follows the `МСК прайс DD.MM.YY.xlsx`
    pattern.  We try the analogous SPB names for recent publication dates, but
    every downloaded workbook is still rejected unless it contains a Moscow row
    and the expected tariff columns.
    """
    today=datetime.now().date()
    dates=[]
    # Weekly Monday editions plus month starts/current day cover the naming used
    # by recent official files without hard-coding a single date forever.
    monday=today-timedelta(days=today.weekday())
    for i in range(max(2,days_back//7+1)): dates.append(monday-timedelta(days=7*i))
    dates += [today, today.replace(day=1)]
    prev=(today.replace(day=1)-timedelta(days=1)).replace(day=1); dates.append(prev)
    seen_dates=[]
    for d in dates:
        if d not in seen_dates: seen_dates.append(d)
    out=[]
    prefixes=("СПБ прайс","СПб прайс","SPB price")
    hosts=("https://cts-group.ru","https://www.cts-group.ru","http://cts-group.ru")
    for d in seen_dates:
        ds=d.strftime("%d.%m.%y")
        for prefix in prefixes:
            filename=f"{prefix} {ds}.xlsx"
            path="/doc/"+quote(filename,safe="")
            for host in hosts:
                u=host+path
                if u not in out: out.append(u)
    return out

def _collect_cts_from_xlsx(origin: str, destination: str) -> tuple[list[str],list[str],str,str]:
    page_urls=["https://cts-group.ru/prices","https://www.cts-group.ru/prices","http://cts-group.ru/prices"]
    candidates=[]; discovery_errors=[]
    for page_url in page_urls:
        try:
            raw,transport=_download_bytes(page_url,referer=page_url.rsplit('/',1)[0]+'/',timeout=28)
            html=raw.decode("utf-8","replace")
            candidates.extend(_cts_xlsx_candidates(html,page_url))
            if candidates: break
        except Exception as exc:
            discovery_errors.append(f"{page_url}: {exc}")
    candidates += _cts_recent_xlsx_candidates()
    # de-duplicate while preserving newest / page-derived candidates first
    candidates=list(dict.fromkeys(candidates))
    errors=[]
    for url in candidates:
        try:
            data,tr=_download_bytes(url,referer="https://cts-group.ru/prices",timeout=25)
            if not data.startswith(b"PK\x03\x04"): raise RuntimeError("ответ не похож на XLSX")
            headers,cells=_cts_parse_xlsx_bytes(data,url)
            return headers,cells,url,f"{tr}+xlsx"
        except Exception as e:
            # Keep the error summary compact; there can be many date candidates.
            if len(errors)<8: errors.append(f"{unquote(url)}: {type(e).__name__}: {e}")
    detail=" | ".join(errors)
    if discovery_errors: detail=("discovery: "+" | ".join(discovery_errors[:2])+"; "+detail)
    raise RuntimeError("CTS XLSX candidates failed: "+detail)

def collect_cts(origin="Санкт-Петербург",destination="Москва") -> dict[str,Any]:
    url="https://cts-group.ru/prices"
    try:
        headers,cells,headings,final_url,browser=_fetch_validated_route_table(url,origin,destination)
    except Exception as route_exc:
        _log(f"CTS route-table failed, trying official XLSX: {type(route_exc).__name__}: {route_exc}")
        headers,cells,final_url,browser=_collect_cts_from_xlsx(origin,destination)
        headings=f"Официальный XLSX CTS: {origin} → {destination}"
    if len(cells)<14: raise RuntimeError(f"CTS: ожидалось >=14 колонок, получено {len(cells)}")
    minimum=_num(cells[2]);
    if minimum is None: raise RuntimeError("CTS: не найден минимум")
    def rate(w:float):
        idx=13 if w<=100 else 12 if w<=200 else 11 if w<=300 else 10 if w<=500 else 9 if w<=800 else 8 if w<=1000 else 7 if w<=1200 else 6 if w<=1500 else 5 if w<=2000 else 4 if w<=3000 else 3
        return _num(cells[idx])
    vals={}
    for p in COMMON_PROFILES:
        if p.get("is_minimum_profile"):
            vals[p["id"]]={"kind":"exact","price":round(minimum,2)}; continue
        r=rate(float(p["weight_kg"]));
        if r is not None and r>0: vals[p["id"]]={"kind":"exact","price":round(max(minimum,float(p["weight_kg"])*r),2),"rate_per_kg":r,"minimum":minimum}
    if len(vals)<20: raise RuntimeError(f"CTS: слишком мало валидных диапазонов: {len(vals)}")
    now=datetime.now().astimezone().isoformat(timespec="seconds")
    save_live_update("CTSgroup",vals,{"source_type":"Официальная тарифная таблица CTS Group — live","source_url":final_url,"captured_at":now,"browser":browser})
    return {"company":"CTSgroup","ok":True,"rows":len(vals),"message":f"CTS: загружено {len(vals)} диапазонов"}

def _newline_rates(headers:list[str],cells:list[str]):
    # Common official layout: destination, days, minimum, <=250, <=750, <=1250, <=2500, <=5000.
    minimum=_num(cells[2]) if len(cells)>=3 else None
    nums=[_num(x) for x in cells[3:8]] if len(cells)>=8 else []
    if minimum is not None and len(nums)==5 and all(x is not None for x in nums):
        return minimum,list(zip([250.,750.,1250.,2500.,5000.],[float(x) for x in nums]))
    # Header-based fallback.
    hs=[_norm(x) for x in headers]
    rates=[]
    for i,h in enumerate(hs):
        if i>=len(cells):break
        if minimum is None and "мин" in h: minimum=_num(cells[i])
        m=re.search(r"(?:до|<=?)\s*(250|750|1250|2500|5000)",h)
        if m:
            v=_num(cells[i]);
            if v is not None: rates.append((float(m.group(1)),v))
    return minimum,sorted(rates)


def _newline_pdf_values(pdf_bytes: bytes, destination: str) -> tuple[float,list[tuple[float,float]],str]:
    from pypdf import PdfReader
    from .legacy_backend import _newline_parse_layout_text
    reader=PdfReader(io.BytesIO(pdf_bytes))
    routes={}; page_texts=[]
    for page_no,page in enumerate(reader.pages[:80],start=1):
        try:
            try: text=page.extract_text(extraction_mode="layout") or ""
            except TypeError: text=page.extract_text() or ""
        except Exception: continue
        page_texts.append((page_no,text))
        routes.update(_newline_parse_layout_text(text,source_file="NL002-live.pdf",page_no=page_no))

    wanted=_norm(destination)
    route_names=[name for name in routes if _norm(name)==wanted]
    if wanted=="москва":
        route_names += [name for name in routes if _norm(name).startswith("москва-") or _norm(name).startswith("москва ")]
    route_names=list(dict.fromkeys(route_names))

    # Fallback for PDF editions where column spacing/layout extraction changed:
    # find a Moscow row and read RUB tokens directly from the same physical line.
    if not route_names and wanted=="москва":
        rub_re=re.compile(r"(?<!\d)(\d[\d \u00a0]*(?:[,.]\d+)?)\s*руб\.?",re.I)
        for page_no,text in page_texts:
            for raw in text.splitlines():
                if "москва" not in _norm(raw): continue
                vals=[_num(x) for x in rub_re.findall(raw)]
                vals=[float(x) for x in vals if x is not None and x>0]
                if len(vals)>=4:
                    city_match=re.search(r"(Москва(?:[-–— ][А-Яа-яЁёA-Za-z]+)?)",raw,re.I)
                    name=(city_match.group(1).strip() if city_match else "Москва")
                    minimum=vals[0]; limits=[250.,750.,1250.,2500.,5000.,10000.]
                    candidates=[{"mode":"minimum","value":minimum,"low":None,"high":None,"file":"NL002-live.pdf","page":page_no}]
                    for lim,rate in zip(limits,vals[1:]):
                        candidates.append({"mode":"per_kg","value":rate,"low":0,"high":lim,"file":"NL002-live.pdf","page":page_no})
                    routes[name]=candidates; route_names=[name]; break
            if route_names: break

    schedules=[]
    for route_name in route_names:
        candidates=routes.get(route_name) or []
        minimum=next((float(x['value']) for x in candidates if x.get('mode')=='minimum' and x.get('value') is not None),None)
        rates=[]
        for x in candidates:
            if x.get('mode')!='per_kg' or x.get('value') is None: continue
            high=x.get('high')
            if high is not None: rates.append((float(high),float(x['value'])))
        rates=sorted(rates)
        if minimum is not None and len(rates)>=3:
            # Choose one published Moscow terminal consistently for the whole
            # grid. Lowest validated 100kg terminal tariff is deterministic and
            # avoids mixing terminal rows from profile to profile.
            r100=next((r for lim,r in rates if 100<=lim+1e-9),None)
            total100=max(minimum,100*r100) if r100 is not None else float('inf')
            schedules.append((total100,route_name,minimum,rates))
    if not schedules:
        found=", ".join(sorted(routes)[:12]) or "нет распознанных направлений"
        raise RuntimeError(f"Новая Линия: PDF получен, но строка {destination} не распознана; распознано: {found}")
    _,route_name,minimum,rates=min(schedules,key=lambda x:x[0])
    return minimum,rates,route_name

def collect_newline(origin="Санкт-Петербург",destination="Москва") -> dict[str,Any]:
    # Official direction-specific PDF endpoint. FROM=175 is Saint Petersburg.
    # This avoids the JS-driven /price/ page entirely.
    if _norm(origin)!=_norm("Санкт-Петербург") or _norm(destination)!=_norm("Москва"):
        raise RuntimeError("Новая Линия: этот адаптер предназначен для Санкт-Петербург → Москва")
    url="https://tknl.ru/price/dev.php?AJAX=N&RESULT=Y&PDF=Y&SERVICE=DELIVERY&FROM=175&TO%5B%5D=ALL"
    pdf,transport=_download_bytes(url,referer="https://tknl.ru/price/",timeout=50)
    if not pdf.startswith(b"%PDF"):
        raise RuntimeError(f"Новая Линия: endpoint вернул не PDF ({len(pdf)} bytes, {pdf[:40]!r})")
    minimum,rates,route_name=_newline_pdf_values(pdf,destination)
    vals={"min":{"kind":"exact","price":round(minimum,2)}}
    max_lim=max(l for l,_ in rates)
    for p in COMMON_PROFILES:
        if p.get("is_minimum_profile"): continue
        w=float(p["weight_kg"]); r=next((r for lim,r in rates if w<=lim+1e-9),None)
        if r is not None:
            vals[p["id"]]={"kind":"exact","price":round(max(minimum,w*r),2),"rate_per_kg":r,"minimum":minimum}
    if len(vals)<12: raise RuntimeError(f"Новая Линия: слишком мало валидных диапазонов из PDF: {len(vals)}")
    now=datetime.now().astimezone().isoformat(timespec="seconds")
    save_live_update("Новая Линия",vals,{"source_type":"Официальный PDF Новая Линия — live NL002","source_url":url,"captured_at":now,"transport":transport,"destination_variant":route_name})
    return {"company":"Новая Линия","ok":True,"rows":len(vals),"message":f"Новая Линия: загружено {len(vals)} диапазонов из NL002, строка {route_name} ({transport})"}

def collect_selected(companies:list[str]) -> list[dict[str,Any]]:
    out=[]
    collectors={"CTSgroup":collect_cts,"Новая Линия":collect_newline}
    for company in companies:
        fn=collectors.get(company)
        if not fn: continue
        try:
            _log(f"{company}: start")
            result=fn(); out.append(result); _log(f"{company}: OK {result}")
        except Exception as exc:
            msg=f"{company}: {type(exc).__name__}: {exc}"; _log(msg); out.append({"company":company,"ok":False,"rows":0,"message":msg})
    return out
