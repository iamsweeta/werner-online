from __future__ import annotations

import json, threading, uuid
from datetime import datetime
from pathlib import Path
from typing import Any

from fastapi import FastAPI, HTTPException, Request, UploadFile, File, Form
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel,StrictStr,StrictInt,StrictFloat

from .v42_engine import (
    BASE_DIR, RUNTIME_DIR, COMPANIES, COMPANY_LABELS, COMMON_PROFILES, PROFILE_BY_ID,
    normalize_city, is_supported_route, paired_destination, quote, matrix, coverage,
    live_path_for,
    live_company_state, age_seconds, LIVE_TTL_SECONDS,
)
from .v42_collectors import collect_selected, LOG_PATH

from .cities import city_names, main_cities, MAIN_ORIGINS
from .tariff_model import tariff_value, tariff_unit, is_rate_profile

VERSION="61.0"
PORT=8423
STATIC_DIR=BASE_DIR/"static"
SETTINGS_PATH=RUNTIME_DIR/"settings.json"
RUNTIME_DIR.mkdir(parents=True,exist_ok=True)

app=FastAPI(title="Tariff Comparison — cities and routes",version=VERSION)
app.mount("/static",StaticFiles(directory=str(STATIC_DIR)),name="static")
@app.middleware('http')
async def no_cache_ui_and_api(request:Request,call_next):
    response=await call_next(request)
    response.headers['Cache-Control']='no-store, no-cache, must-revalidate, max-age=0'
    return response

ORIGINS=city_names()
COLLECT_LOCK=threading.RLock(); COLLECT_JOBS:dict[str,dict[str,Any]]={}; EXACT_REVISION=0

from .bulk_refresh import BulkManager, BusyError, route_plan
BULK=BulkManager(guard=COLLECT_LOCK,route_busy=lambda:any(j.get('status') in {'queued','running'} for j in COLLECT_JOBS.values()))

def _now()->str:return datetime.now().astimezone().isoformat(timespec="seconds")
def _companies(raw:str|None)->list[str]:
    if not raw:return list(COMPANIES)
    requested=[x.strip() for x in raw.split(",") if x.strip()]
    return [x for x in COMPANIES if x in requested]
def _route_key(o:str,d:str)->str:return f"{normalize_city(o)}|{normalize_city(d)}"
def _validate_route(o:str,d:str)->tuple[str,str]:
    o=normalize_city(o); d=normalize_city(d)
    if not is_supported_route(o,d):
        raise HTTPException(400,"Выберите два разных города из списка")
    return o,d

def _settings()->dict[str,Any]:
    try:return json.loads(SETTINGS_PATH.read_text(encoding="utf-8"))
    except:return {}
def _save_settings(data:dict[str,Any]):
    old=_settings(); old.update({k:v for k,v in data.items() if v not in (None,"")})
    SETTINGS_PATH.write_text(json.dumps(old,ensure_ascii=False,indent=2),encoding="utf-8")

@app.get("/")
def root():return FileResponse(STATIC_DIR/"index.html",headers={"Cache-Control":"no-store"})

@app.get("/health")
def health():
    import hashlib
    return {"ok":True,"version":VERSION,"installation_id":hashlib.sha256(str(BASE_DIR.resolve()).encode()).hexdigest()[:16],"engine":"v51_verified_sources","cities_count":len(ORIGINS),"route_selection":"any_distinct_catalog_cities","port_hint":PORT,"collect_log":str(LOG_PATH)}

@app.get("/api/options")
def options(origin:str="Санкт-Петербург",destination:str|None=None,catalog:str="main"):
    if catalog not in {"main","all"}:raise HTTPException(400,"Неизвестный каталог")
    origin=normalize_city(origin)
    if origin not in ORIGINS:raise HTTPException(400,"Выберите город отправления из списка")
    dest=normalize_city(destination) if destination else paired_destination(origin)
    if dest==origin:dest=paired_destination(origin)
    _validate_route(origin,dest)
    integrations=[coverage(origin,dest,c) for c in COMPANIES]
    main=catalog=="main" and origin in MAIN_ORIGINS and dest in main_cities()
    return {"version":VERSION,"catalog":"main" if main else "all","all_origins":ORIGINS,"main_cities":main_cities(),"origins":MAIN_ORIGINS if main else ORIGINS,"destinations":[c for c in (main_cities() if main else ORIGINS) if c!=origin],"selected_origin":origin,"selected_destination":dest,"paired_destination":dest,
            "profiles":COMMON_PROFILES,"companies":[{"id":c,"label":COMPANY_LABELS[c]} for c in COMPANIES],
            "integrations":integrations,"integration_status":{x["id"]:x["coverage_label"] for x in integrations},
            "import_guide":json.loads((BASE_DIR/'data/customer_source_guide.json').read_text(encoding='utf-8')),
            "online_policy":{"live_adapters":list(COMPANIES),"fallback":"LAST GOOD — сохранённая цена, актуальность не подтверждена; LIVE только после текущего успешного обновления"}}

@app.get("/api/compare")
def compare(origin:str,destination:str,profile:str="w100",companies:str|None=None):
    origin,destination=_validate_route(origin,destination)
    if profile not in PROFILE_BY_ID:raise HTTPException(400,"Неизвестный диапазон")
    selected=_companies(companies); p=PROFILE_BY_ID[profile]; items=[quote(c,origin,destination,profile) for c in selected]
    exact=sum(1 for x in items if x.get("comparison_value") is not None and not x.get("price_is_minimum"))
    lower=sum(1 for x in items if x.get("price_is_minimum") and x.get("price") is not None)
    online_exact=sum(1 for x in items if x.get("online") and x.get("comparison_value") is not None)
    online_lower=sum(1 for x in items if x.get("online") and x.get("price_is_minimum") and x.get("price") is not None)
    online_total=online_exact+online_lower
    last_good=sum(1 for x in items if not x.get("online") and not x.get('uploaded') and (x.get("comparison_value") is not None or (x.get("price_is_minimum") and x.get("price") is not None)))
    return {"version":VERSION,"calculated_at":_now(),"origin":origin,"destination":destination,"profile_id":profile,
            "range_weight":p["range_weight"],"comparison_unit":tariff_unit(p),"items":items,"exact_count":exact,
            "lower_bound_count":lower,"online_count":online_total,"online_exact_count":online_exact,
            "online_lower_bound_count":online_lower,"last_good_count":last_good,"imported_count":sum(bool(x.get('uploaded')) for x in items),"missing_count":len(items)-exact-lower}


@app.post('/api/import/preview')
def import_preview(company:str=Form(...),origin:str=Form(...),destination:str=Form(...),file:UploadFile=File(...)):
    from .document_imports import preview
    from .tariff_documents import MAX_BYTES
    origin,destination=_validate_route(origin,destination)
    try:
        return preview(file.file.read(MAX_BYTES+1),file.filename or '',company,origin,destination)
    except Exception as exc:
        raise HTTPException(400,str(exc)[:700]) from exc
    finally:file.file.close()


@app.post('/api/import/jobs',status_code=202)
def import_job_start(company:str=Form(...),origin:str=Form(...),destination:str=Form(...),file:UploadFile=File(...),job_id:str|None=Form(None)):
    from .route_import_jobs import start
    from .tariff_documents import MAX_BYTES
    origin,destination=_validate_route(origin,destination)
    try:return start(file.file.read(MAX_BYTES+1),file.filename or '',company,origin,destination,job_id)
    except ValueError as exc:raise HTTPException(400,str(exc)) from exc
    finally:file.file.close()


@app.get('/api/import/jobs/{job_id}')
def import_job_status(job_id:str):
    from .route_import_jobs import status
    try:return status(job_id)
    except ValueError as exc:raise HTTPException(400,str(exc)) from exc
    except FileNotFoundError as exc:raise HTTPException(404,str(exc)) from exc


class ImportCommit(BaseModel):
    token:str
    resolutions:dict[str,int]|None=None


@app.post('/api/import/commit')
def import_commit(body:ImportCommit):
    from .document_imports import commit
    try:return commit(body.token)
    except (ValueError,KeyError,FileNotFoundError) as exc:raise HTTPException(400,str(exc)) from exc


@app.get('/api/imports')
def imports_list(origin:str,destination:str):
    from .document_imports import pack
    origin,destination=_validate_route(origin,destination)
    return {'companies':pack(origin,destination).get('companies',{})}


@app.delete('/api/import')
def import_remove(company:str,origin:str,destination:str):
    from .document_imports import remove
    origin,destination=_validate_route(origin,destination)
    try:return remove(company,origin,destination)
    except ValueError as exc:raise HTTPException(400,str(exc)) from exc


@app.get('/api/import-file/{filename}')
def import_file(filename:str):
    from .document_imports import source_file as imported_source
    try:return FileResponse(imported_source(filename),filename=filename)
    except FileNotFoundError as exc:raise HTTPException(404,'Загруженный файл не найден') from exc


@app.get('/api/import/template')
def import_template(company:str,origin:str,destination:str):
    from io import BytesIO
    from openpyxl import Workbook
    from .tariff_documents import GENERIC_HEADER
    origin,destination=_validate_route(origin,destination)
    if company not in COMPANIES:raise HTTPException(400,'Выберите компанию')
    wb=Workbook();ws=wb.active;ws.title='Тарифы';ws.append(GENERIC_HEADER)
    for p in COMMON_PROFILES:
        ws.append([company,origin,destination,'МИН' if p.get('is_minimum_profile') else p['weight_kg'],None])
    for col in ('A','B','C'):ws.column_dimensions[col].width=26
    ws.column_dimensions['D'].width=14;ws.column_dimensions['E'].width=20;ws.freeze_panes='A2'
    help=wb.create_sheet('Инструкция')
    help.append(['Внесите реальные цены за отправку (рубли) в столбец E. Удалите строки, для которых цены нет.'])
    help.append(['МИН — только опубликованный минимум. Вес — контрольная точка, не ставка за кг.'])
    help.append(['Ставку ₽/кг предварительно умножьте на вес с учётом минимальной стоимости перевозчика.'])
    help.append(['Компания и направление должны совпадать с выбранными при импорте. Цены-примеры отсутствуют.'])
    help.column_dimensions['A'].width=115
    out=BytesIO();wb.save(out)
    return Response(out.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="tariff_import_template.xlsx"'})

@app.get("/api/profile-matrix")
def profile_matrix(origin:str,destination:str,companies:str|None=None):
    origin,destination=_validate_route(origin,destination)
    return {"version":VERSION,"origin":origin,"destination":destination,"profiles":matrix(origin,destination,_companies(companies))}

class CollectRequest(BaseModel):
    companies:list[str]|None=None
    origin:str="Санкт-Петербург"
    destination:str="Москва"
    profile:str|None=None
    strict_only:bool=True
    background:bool=True
    force:bool=True

def _run_collect(key:str,selected:list[str],origin:str,destination:str):
    global EXACT_REVISION
    job=COLLECT_JOBS[key]; job.update({"status":"running","started_at":_now(),"message":"Проверяю онлайн-источники для выбранного направления…"})
    try:
        targets=list(selected)
        job["targets"]=targets; job["progress_rows"]=0
        def progress(result):
            with COLLECT_LOCK:
                by={x['company']:x for x in job.get('results',[])}
                by[result['company']]=result
                job['results']=list(by.values())
                job['completed_companies']=len(by)
                job['progress_rows']=sum(int(x.get('rows') or 0) for x in by.values())
                job['progress_revision']=int(job.get('progress_revision',0))+1
                job['message']=result.get('message','')
        results=collect_selected(targets,origin,destination,job.get("profile") or "w100",on_progress=progress,should_stop=lambda:job.get("stop_requested",False))
        job["results"]=results; job["progress_rows"]=sum(int(x.get("rows") or 0) for x in results)
        success=[x for x in results if x.get("ok")]
        with COLLECT_LOCK:
            EXACT_REVISION+=len(success); job["exact_revision"]=EXACT_REVISION
        unavailable=[x for x in results if (x.get("error_info") or {}).get("code")=="route_unpublished"]
        failed=[x for x in results if not x.get("ok") and x not in unavailable]
        job["unavailable_count"]=len(unavailable)
        job["success_companies"]=[x.get("company") for x in success]
        job["failed_companies"]=[x.get("company") for x in failed]
        job["success_count"]=len(success); job["failed_count"]=len(failed)
        job['partial_count']=sum(bool(x.get('partial')) for x in results)
        if failed:
            job["message"]="Онлайн-проверка завершена. Успешные ответы сохранены, при сбоях оставлен последний подтверждённый прайс. " + " | ".join(x.get("message","") for x in failed)
        else:job["message"]="Онлайн-проверка завершена: доступные live-источники обновлены."
        job["status"]="paused" if job.get("stop_requested") else "done"; job["finished_at"]=_now()
        if job.get("stop_requested"):job["message"]="Загрузка остановлена. Полученные цены сохранены."
    except Exception as exc:
        job.update({"status":"error","finished_at":_now(),"message":f"Ошибка обновления: {type(exc).__name__}: {exc}"})

@app.post("/api/collect")
def collect(body:CollectRequest):
    origin,destination=_validate_route(body.origin,body.destination); key=_route_key(origin,destination)
    profile=body.profile or 'w100'
    if profile not in PROFILE_BY_ID:raise HTTPException(400,'Неизвестный диапазон')
    if body.companies is not None and (not body.companies or any(c not in COMPANIES for c in body.companies)):
        raise HTTPException(400,'Укажите хотя бы одну компанию из списка')
    selected=list(dict.fromkeys(body.companies if body.companies is not None else COMPANIES))
    with COLLECT_LOCK:
        bulk_id=BULK.active()
        if bulk_id:
            return {'ok':False,'status':'busy','bulk_job_id':bulk_id,'message':'Выполняется общий сбор всех компаний. Прогресс — в разделе «Большая таблица».'}
        current=COLLECT_JOBS.get(key)
        if current and current.get("status") in {"queued","running"}:
            return {"ok":True,"status":current["status"],"job_id":current.get("job_id"),"created_at":current.get("created_at"),"origin":origin,"destination":destination,"message":"Обновление этого маршрута уже выполняется"}
        # Do not run two 17-carrier batches at once.  The user's v42.2 log showed
        # Москва→СПб still finishing while СПб→Москва had already started; that
        # doubled browser/TLS load and interleaved the diagnostics.
        other=next(((k,j) for k,j in COLLECT_JOBS.items() if k!=key and j.get("status") in {"queued","running"}),None)
        if other:
            _,active=other
            return {"ok":False,"status":"busy","job_id":active.get("job_id"),"active_origin":active.get("origin"),"active_destination":active.get("destination"),
                    "message":f"Сначала завершится активное обновление {active.get('origin')} → {active.get('destination')}. Повторите обновление второго маршрута после завершения."}
        if not body.force:
            from .online_tariffs import ADAPTERS
            calculators={'Werner','Главтрасса','ДЛ','Возовоз','Байкал Сервис','Пролайн'}
            def due(company):
                state=live_company_state(origin,destination,company)
                if quote(company,origin,destination,profile).get('online'):return False
                age=age_seconds(state.get('last_attempt_at'))
                if state.get('attempt_status')=='unavailable':return age>LIVE_TTL_SECONDS
                if state.get('attempt_status')=='partial':return age>600
                if state.get('attempt_status')=='failed':
                    # A failed site must not restart every minute or on every weight change.
                    unsupported='конкретный вес' in str(state.get('last_error') or '')
                    return (unsupported and state.get('requested_profile')!=profile) or age>600
                if company in calculators and state.get('requested_profile')!=profile:return True
                return age>(LIVE_TTL_SECONDS if state.get('attempt_status')=='success' else 60)
            selected=[c for c in selected if due(c)]
            if not selected:
                return {'ok':True,'status':'fresh','origin':origin,'destination':destination,'message':'Автоматическая проверка пока не требуется. После сбоя повтор выполняется через 10 минут; вручную можно повторить сразу.'}
        job_id=f"live-{datetime.now().strftime('%Y%m%d-%H%M%S')}-{uuid.uuid4().hex[:6]}"
        COLLECT_JOBS[key]={"status":"queued","job_id":job_id,"created_at":_now(),"request_received_at":_now(),"origin":origin,"destination":destination,"profile":body.profile or "w100","progress_rows":0,"exact_revision":EXACT_REVISION,"message":"Поставлено в очередь","requested_companies":selected}
    try:
        with LOG_PATH.open("a",encoding="utf-8") as f:f.write(f"[{_now()}] UI COLLECT ACCEPTED {job_id} | {origin} → {destination} | {len(selected)} companies\n")
    except Exception: pass
    threading.Thread(target=_run_collect,args=(key,selected,origin,destination),daemon=True).start()
    return {"ok":True,"status":"queued","job_id":job_id,"created_at":COLLECT_JOBS[key]["created_at"],"origin":origin,"destination":destination,"requested_companies":len(selected),"message":"Запущена реальная онлайн-проверка источников"}

@app.post('/api/collect/stop')
def stop_collect(body:CollectRequest):
    origin,destination=_validate_route(body.origin,body.destination)
    with COLLECT_LOCK:
        job=COLLECT_JOBS.get(_route_key(origin,destination))
        if not job:raise HTTPException(404,'Обновление не найдено')
        if job.get('status') in {'queued','running'}:
            job['stop_requested']=True;job['message']='Останавливаю новые компании. Текущие проверки будут сохранены.'
        return json.loads(json.dumps(job))

@app.get('/api/bulk/history')
def bulk_history():
    with BULK.db() as db:rows=db.execute('SELECT id,status,created_at,config FROM jobs ORDER BY created_at DESC,rowid DESC LIMIT 20').fetchall()
    return {'jobs':[{'job_id':r['id'],'status':r['status'],'created_at':r['created_at'],'scope':json.loads(r['config'])['scope'],'mode':json.loads(r['config']).get('mode','online')} for r in rows]}

@app.get("/api/collect-status")
def collect_status(origin:str,destination:str):
    origin,destination=_validate_route(origin,destination); job=COLLECT_JOBS.get(_route_key(origin,destination))
    if not job:return {"status":"idle","progress_rows":0,"exact_revision":EXACT_REVISION,"message":"Онлайн-обновление ещё не запускалось"}
    with COLLECT_LOCK:return json.loads(json.dumps(job))

@app.get('/api/active-collect')
def active_collect():
    with COLLECT_LOCK:
        bulk_id=BULK.active()
        if bulk_id:return {'status':'running','mode':'bulk','bulk_job_id':bulk_id}
        job=next((j for j in COLLECT_JOBS.values() if j.get('status') in {'queued','running'}),None)
        return json.loads(json.dumps(job)) if job else {'status':'idle'}


class BulkRequest(BaseModel):
    scope:str='reference'
    origins:list[str]|None=None
    destinations:list[str]|None=None
    mode:str='online'
    include_imports:bool=True


def _bulk_call(fn,*args,**kwargs):
    try:return fn(*args,**kwargs)
    except BusyError as exc:raise HTTPException(409,str(exc)) from exc
    except ValueError as exc:raise HTTPException(400,str(exc)) from exc


@app.post('/api/bulk/plan')
def bulk_plan(body:BulkRequest):
    return _bulk_call(route_plan,body.scope,body.origins,body.destinations)


@app.post('/api/bulk')
def bulk_start(body:BulkRequest):
    return _bulk_call(BULK.start,body.scope,body.origins,body.destinations,mode=body.mode,include_imports=body.include_imports)


@app.get('/api/bulk')
def bulk_latest():return BULK.status()


@app.get('/api/bulk/{job_id}')
def bulk_status(job_id:str):return _bulk_call(BULK.status,job_id)


@app.post('/api/bulk/{job_id}/{action}')
def bulk_action(job_id:str,action:str):
    actions={'pause':lambda:BULK.pause(job_id),'resume':lambda:BULK.resume(job_id),
             'retry':lambda:BULK.resume(job_id,retry=True),'export':lambda:BULK.prepare_export(job_id)}
    if action not in actions:raise HTTPException(400,'Неизвестное действие')
    return _bulk_call(actions[action])


@app.get('/api/bulk/{job_id}/download')
def bulk_download(job_id:str):
    file=_bulk_call(BULK.download,job_id)
    return FileResponse(file,filename=('tariffs_all_companies.xlsx' if file.suffix=='.xlsx' else 'tariffs_all_routes.zip'))


@app.get('/api/price-documents')
def documents_list():
    from .price_library import list_files
    return {'files':list_files()}


@app.get('/api/storage')
def storage_info():
    from .storage import summary
    return summary()


@app.get('/api/route-documents')
def documents_for_route(origin:str,destination:str):
    from .price_library import route_documents
    origin,destination=_validate_route(origin,destination)
    return _bulk_call(route_documents,origin,destination)


class RouteDocumentChoice(BaseModel):
    company:str
    origin:str
    destination:str
    document_id:str|None=None


@app.post('/api/route-documents/select')
def document_choice(body:RouteDocumentChoice):
    from .price_library import select_document
    return _bulk_call(select_document,body.company,body.origin,body.destination,body.document_id)


@app.get('/api/price-documents/{document_id}/routes')
def routes_in_document(document_id:str):
    from .price_library import document_routes
    return _bulk_call(document_routes,document_id)


@app.get('/api/storage/backup')
def storage_backup():
    from .storage import backup
    from starlette.background import BackgroundTask
    path=backup()
    return FileResponse(path,filename='tariff_documents_backup.zip',background=BackgroundTask(path.unlink,missing_ok=True))


@app.get('/api/price-documents/template')
def documents_template(company:str,origin:str='',kind:str='interval'):
    from .price_library import template
    content=_bulk_call(template,company,origin or None,kind)
    return Response(content,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                    headers={'Content-Disposition':'attachment; filename=multi_route_price_template.xlsx'})


@app.post('/api/price-documents/preview')
def documents_preview(company:str=Form(...),origin:str=Form(''),document_date:str=Form(''),file:UploadFile|None=File(None),files:list[UploadFile]|None=File(None)):
    from .price_library import start_many
    from .tariff_documents import MAX_BYTES
    selected=([file] if file is not None else [])+(files or [])
    try:
        if not selected or len(selected)>30:raise HTTPException(400,'Выберите от 1 до 30 документов')
        uploads=[];total=0
        for item in selected:
            raw=item.file.read(MAX_BYTES-total+1);total+=len(raw)
            if total>MAX_BYTES:raise HTTPException(400,'Общий размер выбранных файлов — не более 20 МБ')
            uploads.append((item.filename or '',raw))
        return _bulk_call(start_many,uploads,company,origin or None,document_date or None)
    finally:
        for item in selected:item.file.close()


@app.get('/api/price-documents/preview/{token}')
def document_status(token:str):
    from .price_library import status
    return _bulk_call(status,token)


@app.post('/api/price-documents/commit')
def documents_commit(body:ImportCommit):
    from .price_library import commit
    return _bulk_call(commit,body.token,body.resolutions)


@app.delete('/api/price-documents/{ident}')
def documents_remove(ident:str):
    from .price_library import remove
    return _bulk_call(remove,ident)

@app.get('/api/source-file/{filename}')
def source_file(filename:str):
    from .legacy import DOWNLOAD_DIR
    path=(DOWNLOAD_DIR/filename).resolve()
    if path.parent!=DOWNLOAD_DIR.resolve() or not path.is_file():raise HTTPException(404,'Файл источника не найден')
    return FileResponse(path,filename=path.name)

@app.get("/api/diagnostics")
def diagnostics(origin:str="Санкт-Петербург",destination:str="Москва",profile:str="w100"):
    origin,destination=_validate_route(origin,destination)
    if profile not in PROFILE_BY_ID:raise HTTPException(400,"Неизвестный диапазон")
    items=[quote(c,origin,destination,profile) for c in COMPANIES]
    exact=[x["company"] for x in items if x.get("comparison_value") is not None and not x.get("price_is_minimum")]
    lower=[x["company"] for x in items if x.get("price_is_minimum") and x.get("price") is not None]
    missing=[x["company"] for x in items if x.get("comparison_value") is None and not (x.get("price_is_minimum") and x.get("price") is not None)]
    online=[x["company"] for x in items if x.get("online") and (x.get("comparison_value") is not None or (x.get("price_is_minimum") and x.get("price") is not None))]
    online_exact=[x["company"] for x in items if x.get("online") and x.get("comparison_value") is not None]
    online_lower=[x["company"] for x in items if x.get("online") and x.get("price_is_minimum") and x.get("price") is not None]
    try:tail=LOG_PATH.read_text(encoding="utf-8",errors="replace").splitlines()[-40:]
    except Exception:tail=[]
    lp=live_path_for(origin,destination)
    evidence=[{"company":x.get("company_label") or x.get("company"),"online":bool(x.get("online")),"freshness":x.get("freshness"),"data_origin":x.get("data_origin"),"original_filename":x.get("original_filename"),"document_date":x.get("document_date"),"sha256":x.get("sha256"),
               "price":x.get("price"),"comparison_value":x.get("comparison_value"),"price_is_minimum":bool(x.get("price_is_minimum")),
               "captured_at":x.get("captured_at"),"attempt_id":x.get("live_attempt_id"),"transport":x.get("transport"),
               "source_type":x.get("source_type"),"source_url":x.get("source_url"),"refresh_status":x.get("refresh_status"),
               "refresh_error":x.get("refresh_error"),"error_info":x.get("error_info")} for x in items]
    job=COLLECT_JOBS.get(_route_key(origin,destination))
    report={"ok":True,"version":VERSION,"engine":"v51_verified_sources","route":f"{origin} → {destination}","profile":profile,
            "summary":{"exact":len(exact),"lower_bound":len(lower),"missing":len(missing),"online_total":len(online),"online_exact":len(online_exact),"online_lower_bound":len(online_lower)},
            "exact_companies":exact,"lower_bound_companies":lower,"missing_companies":missing,"online_companies":online,"online_exact_companies":online_exact,"online_lower_bound_companies":online_lower,"live_evidence":evidence,
            "collect_job":job,"live_update_exists":lp.exists(),"live_path":str(lp),"collect_log_tail":tail,
            "data_policy":"Архивные цены исходного проекта не проверены на актуальность и никогда не считаются LIVE. Только строки текущей успешной сетевой попытки с подтверждённым направлением становятся LIVE на 30 минут. Ошибка не превращает сохранённые значения в текущие цены."}

    import platform, ssl, importlib.metadata
    try:transport_version=importlib.metadata.version('curl_cffi')
    except importlib.metadata.PackageNotFoundError:transport_version=None
    report['environment']={'python':platform.python_version(),'os':platform.platform(),'ssl':ssl.OPENSSL_VERSION,
                           'curl_cffi':transport_version,'browser_calculators':bool(_settings().get('public_browser_enabled'))}
    # Diagnostic downloads may be sent for support; strip configured secrets from all strings.
    encoded=json.dumps(report,ensure_ascii=False)
    for key in ('dellin_appkey','vozovoz_api_key','baikal_api_key'):
        secret=_settings().get(key)
        if isinstance(secret,str) and secret:encoded=encoded.replace(json.dumps(secret,ensure_ascii=False)[1:-1],'[СКРЫТО]')
    return json.loads(encoded)


@app.get('/api/diagnostics/download')
def diagnostics_download(origin:str='Санкт-Петербург',destination:str='Москва',profile:str='w100'):
    report=diagnostics(origin,destination,profile)
    return Response(json.dumps(report,ensure_ascii=False,indent=2).encode('utf-8'),media_type='application/json',
                    headers={'Content-Disposition':'attachment; filename="tariff_diagnostics_55_0.json"'})


@app.get("/api/settings")
def settings_get():
    cfg=_settings(); return {"dellin_appkey":bool(cfg.get("dellin_appkey")),"vozovoz_api_key":bool(cfg.get("vozovoz_api_key")),"baikal_api_key":bool(cfg.get("baikal_api_key")),"baikal_api_url":cfg.get("baikal_api_url") or "","public_browser_enabled":cfg.get("public_browser_enabled") is True}

@app.post("/api/settings")
async def settings_post(request:Request):
    data=await request.json(); allowed={k:data.get(k) for k in ("dellin_appkey","vozovoz_api_key","baikal_api_key","baikal_api_url") if data.get(k)}
    if 'public_browser_enabled' in data:
        if not isinstance(data['public_browser_enabled'],bool):raise HTTPException(400,'Ожидается переключатель браузерного расчёта')
        allowed['public_browser_enabled']=data['public_browser_enabled']
    _save_settings(allowed); return {"ok":True}

@app.get("/api/export/excel")
def export_excel(origin:str,destination:str,profile:str="w100",companies:str|None=None,view:str="total",live_only:bool=False,include_imports:bool=True,layout:str="customer",all_loaded:bool=False):
    from io import BytesIO
    from openpyxl import Workbook
    if layout not in {'customer','matrix','route'}:raise HTTPException(400,'Неизвестный формат Excel')
    if layout=='route' and view not in {'total','per_kg'}:raise HTTPException(400,'Неизвестная единица измерения')
    if layout == 'customer':
        from .excel_export import export_bytes
        origin,destination=_validate_route(origin,destination)
        content=export_bytes(origin,destination,_companies(companies),live_only=live_only,
                             include_imports=include_imports,all_loaded=all_loaded)
        return Response(content,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        headers={"Content-Disposition":"attachment; filename=tariffs_v510_customer.xlsx"})
    origin,destination=_validate_route(origin,destination); selected=_companies(companies); rows=matrix(origin,destination,selected); per_kg=str(view).lower()=="per_kg"
    wb=Workbook(); ws=wb.active; ws.title="Тарифы"; ws.append(["Маршрут",f"{origin} → {destination}"]); ws.append(["Версия",VERSION]); ws.append(["Режим","Расчётная стоимость: сумма отправки ÷ контрольный вес, ₽/кг; МИН — ₽" if per_kg else "Стоимость отправки, ₽"]); ws.append([])
    ws.append(["Диапазон","Тип","Ед."]+[COMPANY_LABELS[c] for c in selected])
    for row in rows:
        p=row["profile"]; row_per_kg=per_kg and not p.get("is_minimum_profile"); unit="₽/кг" if row_per_kg else "₽"; vals=[]
        for item in row["items"]:
            if live_only and not item.get('online') and not (include_imports and item.get('uploaded')):
                vals.append(None)
            elif item.get("comparison_value") is not None:
                v=item["comparison_value"] / p["weight_kg"] if row_per_kg else item["comparison_value"]
                vals.append(round(float(v),4) if v is not None else None)
            elif item.get("price_is_minimum"):vals.append(None)
            else:vals.append(None)
        ws.append([p["range_weight"],"Минимум" if p.get("is_minimum_profile") else "Расчётная стоимость" if per_kg else "Стоимость отправки",unit]+vals)
    audit=wb.create_sheet('Источники')
    audit.append(['Компания','Диапазон','Источник данных','Получено / импортировано','Официальный URL','Расчёт','Ошибка онлайн-обновления','Имя файла пользователя','Дата в документе','SHA256','Страница / строка','Условия НДС'])
    for row in rows:
        for item in row['items']:
            audit.append([item['company_label'],row['profile']['range_weight'],'Введено вручную' if item.get('manual') else 'Файл пользователя' if item.get('uploaded') else 'LIVE' if item.get('online') else 'LAST GOOD' if item.get('price') is not None else 'Нет данных',
                          item.get('captured_at'),item.get('source_url'),item.get('calculation_basis'),item.get('refresh_error'),item.get('original_filename'),item.get('document_date'),item.get('sha256'),str(item.get('source_page') or item.get('source_row') or ''),item.get('tax_basis') or 'Не определены; см. оригинал'])
    audit.freeze_panes='A2'
    # User-provided file names and remote error strings are text, never Excel formulas.
    for sheet in wb:
        for cells in sheet:
            for cell in cells:
                if cell.data_type=='f':cell.data_type='s'
    if layout=='route':
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        ws.title='Маршрут'
        ws.sheet_view.showGridLines=False
        ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=26;ws.column_dimensions['C'].width=12
        for col in range(4,ws.max_column+1):ws.column_dimensions[get_column_letter(col)].width=19
        for row_no in (1,3):
            ws.merge_cells(start_row=row_no,start_column=2,end_row=row_no,end_column=max(4,ws.max_column))
            ws.cell(row_no,2).alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[row_no].height=42 if row_no==3 else 30
        ws['A1'].font=Font(name='Calibri',size=14,bold=True);ws['B1'].font=Font(name='Calibri',size=14,bold=True)
        for cell in ws[5]:
            cell.fill=PatternFill('solid',fgColor='EDF2FF');cell.font=Font(name='Calibri',bold=True,color='163D81')
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        ws.row_dimensions[5].height=38
        for row in ws.iter_rows(min_row=6):
            ws.row_dimensions[row[0].row].height=32
            for cell in row:
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.column>=4:cell.number_format='#,##0.####'
                if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F5F7FA')
        ws.auto_filter.ref=f'A5:{get_column_letter(ws.max_column)}{ws.max_row}'
        ws.print_title_rows='1:5';ws.sheet_properties.pageSetUpPr.fitToPage=True
        ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth=1;ws.page_setup.fitToHeight=0
        for col in range(1,audit.max_column+1):audit.column_dimensions[get_column_letter(col)].width=26 if col<5 else 48
        for cell in audit[1]:
            cell.fill=PatternFill('solid',fgColor='EDF2FF');cell.font=Font(bold=True)
            cell.alignment=Alignment(wrap_text=True,vertical='center')
        audit.row_dimensions[1].height=42;audit.auto_filter.ref=audit.dimensions
    ws.freeze_panes="D6"; out=BytesIO(); wb.save(out); suffix="perkg" if per_kg else "total"
    disposition=f'attachment; filename=tariffs_v510_{suffix}.xlsx'
    if layout=='route':
        from urllib.parse import quote as urlquote
        stamp=datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
        filename=f'Тарифы_{origin}_{destination}_{stamp}.xlsx'
        disposition=f'attachment; filename="tariffs_route_{stamp}.xlsx"; filename*=UTF-8\'\'{urlquote(filename,safe="")}'
    return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":disposition})


@app.get('/api/export/route')
def export_route(origin:str,destination:str,companies:str|None=None,view:str='total',live_only:bool=False,include_imports:bool=True):
    """Independent file containing only this route; never the customer book."""
    return export_excel(origin,destination,companies=companies,view=view,live_only=live_only,
                        include_imports=include_imports,layout='route',all_loaded=False)


class ManualPriceRequest(BaseModel):
    company:str
    origin:str
    destination:str
    profile:str
    value:StrictStr|StrictFloat|StrictInt|None=None
    unit:str='rub'

@app.post('/api/manual-price')
def manual_price_put(body:ManualPriceRequest):
    from .manual_prices import put
    return _bulk_call(put,body.company,body.origin,body.destination,body.profile,body.value,body.unit)

@app.delete('/api/manual-price')
def manual_price_remove(body:ManualPriceRequest):
    from .manual_prices import remove
    return _bulk_call(remove,body.company,body.origin,body.destination,body.profile)
